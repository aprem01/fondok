"""Deterministic USALI Summary-sheet reconciler for P&L-family workbooks.

FON-41 (Historicals). On the Angler's 2024 full-year detailed P&L the LLM
extractor read ``p_and_l_usali.operating_revenue.food_beverage_revenue`` =
96,528.1 off a hashed department sheet (``D_REST_CON 400000 …`` — the
restaurant-CONCESSION outlet line) while the workbook's own ``Summary``
sheet states Food & Beverage revenue at 2,739,040.71. The Historicals grid
then showed 2024 F&B at $96K against 2023 $2.11M and the T-12 $3.22M, the
YoY variance engine flagged it, and the Year-1 anchor was grounded on it.

A USALI workbook that carries a ``Summary`` sheet is stating its own
department totals, so this module reads that sheet BY LABEL and, for every
canonical USALI total the extractor emitted, replaces an extracted value
that differs from the Summary by more than 5%. Everything here is pure
Python over the parsed grid (or the workbook bytes when available); no LLM.

Scope rules
-----------
* Runs only for P&L-family documents (``T12`` / ``PNL`` / ``PNL_MONTHLY`` /
  ``PNL_YTD``). Never for OM / STR / market documents — the caller gates
  on doc type and :func:`reconcile_extraction` re-checks it.
* Needs a sheet named ``Summary`` (case-insensitive; also ``P&L Summary``,
  ``Summary P&L``, ``USALI Summary``) whose rows carry USALI labels.
* Only fields the Summary STATES are touched. A field whose concept has no
  Summary row, a non-dollar field, a monthly / quarterly slice, or a field
  whose value is within 5% of the Summary is left exactly as extracted.
* The annual column is the rightmost numeric column that equals the sum of
  the 12 monthly columns within 0.5% on the USALI rows; failing that, the
  column labelled ``Total`` / ``YTD`` / ``Annual`` (``YTD`` only when the
  document itself is year-to-date). A document whose own period basis is
  monthly / quarterly is never reconciled — its "total" column would be a
  different period from the extracted month.

Precision
---------
The parser renders workbook floats with ``%g`` (six significant digits), so
the cached grid holds ``2.73904e+06`` for 2,739,040.71 and ``1.34817e+07``
for 13,481,730.29. When the caller can supply the original workbook bytes
the Summary sheet is re-read with openpyxl at full precision; otherwise the
cached grid is used and the result is exact to six significant digits —
still four orders of magnitude closer than the outlet line it replaces.

Provenance
----------
A replaced field keeps its ``field_name`` and ``unit`` and gets:

* ``value``        — the Summary annual total
* ``confidence``   — 0.98 (a stated total read deterministically)
* ``raw_text``     — the Summary row (label, monthly values, annual total)
* ``source_page``  — the Summary sheet's page index
* ``reviewed``     — ``"reconciled"`` (the Data Room shows the field as changed)
* ``reconciled_from`` — ``{field_name, old_value, sheet, row, label,
  old_source_page, old_raw_text}`` so the analyst can see what the LLM read
* ``note``         — one sentence saying the same in plain language

Those extra keys (``reconciled_from`` / ``note`` / ``reviewed``) are also
what keeps the field out of the strict ``fondok_schemas.ExtractionField``
path the citation verifier re-reads: that verifier matches the cited number
against the lossy page text and would otherwise demote a deterministic
Summary total it cannot find verbatim.
"""

from __future__ import annotations

import contextlib
import logging
import math
import re
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from typing import Any

logger = logging.getLogger(__name__)


# ─────────────────────────── gates ───────────────────────────


PNL_FAMILY_DOC_TYPES: frozenset[str] = frozenset({"T12", "PNL", "PNL_MONTHLY", "PNL_YTD"})

#: Accepted Summary sheet names (compared after :func:`_norm_label`).
SUMMARY_SHEET_NAMES: frozenset[str] = frozenset(
    {"summary", "p&l summary", "summary p&l", "usali summary"}
)

#: Period bases (``engines.historical_baseline._period_basis`` vocabulary)
#: on which a Summary "total" column is a DIFFERENT period from the
#: extracted lines, so the reconciler must not run.
_PARTIAL_BASES_NEVER_RECONCILED: frozenset[str] = frozenset({"MONTHLY", "QUARTERLY", "WEEKLY"})

#: Extracted values are compared only when the field is a whole-dollar
#: line. A ``$000`` / ratio / percent unit would make the comparison —
#: and the replacement — meaningless.
_WHOLE_DOLLAR_UNITS: frozenset[str] = frozenset({"", "usd", "$", "dollars", "us$", "usd_whole"})

RECONCILE_TOLERANCE = 0.05  #: replace when |extracted - summary| > 5% of summary
ANNUAL_SUM_TOLERANCE = 0.005  #: annual column must equal Σ months within 0.5%
RECONCILED_CONFIDENCE = 0.98


def canonical_doc_type(value: str | None) -> str:
    """``"pnl-monthly"`` / ``"PNL MONTHLY"`` / ``"PNL_MONTHLY"`` → ``"PNL_MONTHLY"``."""
    if not value:
        return ""
    v = value.strip().upper().replace("-", "_").replace(" ", "_")
    v = re.sub(r"_+", "_", v)
    if v in {"T_12", "T12"}:
        return "T12"
    return v


def is_pnl_family(doc_type: str | None) -> bool:
    return canonical_doc_type(doc_type) in PNL_FAMILY_DOC_TYPES


# ─────────────────────────── label map ───────────────────────────


#: Section headers (rows with a label and no numbers) that re-scope the
#: ambiguous ``Rooms`` / ``Food & Beverage`` / ``Other Operated`` labels.
_SECTION_HEADERS: dict[str, str] = {
    "revenues": "revenue",
    "revenue": "revenue",
    "operating revenues": "revenue",
    "operating revenue": "revenue",
    "departmental expense": "dept_expense",
    "departmental expenses": "dept_expense",
    "operated department expenses": "dept_expense",
    "operated departmental expenses": "dept_expense",
    "undistributed expenses": "undistributed",
    "undistributed operating expenses": "undistributed",
    "non-operating income & expenses": "nonop",
    "non operating income & expenses": "nonop",
    "non-operating expenses": "nonop",
    "management fees": "nonop",  # not a header on every layout; handled as a line too
    # Sections whose rows must never be read as the statement's totals
    # (restatements and arithmetic checks). The "Interest, Depreciation &
    # Amortization" block is deliberately NOT a skip header: its own lines
    # map to no concept, and the EBITDA / FF&E / "EBITDA: Less Replacement
    # Reserve" rows that follow it on the USALI layout must stay readable.
    "adjusted values": "skip",
    "check": "skip",
    "statistics": "skip",
    "kpis": "skip",
}

#: Labels whose meaning depends on the section they sit under.
_SECTION_SCOPED_LABELS: dict[str, dict[str, str]] = {
    "revenue": {
        "rooms": "rooms_revenue",
        "rooms revenue": "rooms_revenue",
        "rooms revenues": "rooms_revenue",
        "total rooms revenue": "rooms_revenue",
        "food & beverage": "fb_revenue",
        "f&b": "fb_revenue",
        "food & beverage revenue": "fb_revenue",
        "food & beverage revenues": "fb_revenue",
        "total food & beverage revenue": "fb_revenue",
        "other operated departments": "other_revenue",
        "other operated department": "other_revenue",
        "other operated": "other_revenue",
        "other operated departments revenue": "other_revenue",
        "other revenue": "other_revenue",
        "other revenues": "other_revenue",
        "miscellaneous income": "misc_revenue",
        "misc income": "misc_revenue",
        "misc. income": "misc_revenue",
        "miscellaneous revenue": "misc_revenue",
    },
    "dept_expense": {
        "rooms": "rooms_dept_expense",
        "rooms expense": "rooms_dept_expense",
        "rooms expenses": "rooms_dept_expense",
        "food & beverage": "fb_dept_expense",
        "f&b": "fb_dept_expense",
        "food & beverage expense": "fb_dept_expense",
        "food & beverage expenses": "fb_dept_expense",
        "other operated departments": "other_dept_expense",
        "other operated department": "other_dept_expense",
        "other operated": "other_dept_expense",
        "other operated departments expense": "other_dept_expense",
    },
    "undistributed": {
        "administrative & general": "administrative_general",
        "a&g": "administrative_general",
        "general & administrative": "administrative_general",
        "information & telecom systems": "information_telecom",
        "information & telecommunications systems": "information_telecom",
        "information & telecommunications": "information_telecom",
        "information & telecom": "information_telecom",
        "it & telecom": "information_telecom",
        "sales & marketing": "sales_marketing",
        "marketing": "sales_marketing",
        "property operation & maintenance": "property_operations",
        "property operations & maintenance": "property_operations",
        "repairs & maintenance": "property_operations",
        "pom": "property_operations",
        "utilities": "utilities",
        "utility costs": "utilities",
    },
    "nonop": {
        "property & other taxes": "property_taxes",
        "property taxes": "property_taxes",
        "real estate taxes": "property_taxes",
        "insurance": "insurance",
        "rent": "rent_expense",
    },
}

#: Labels that mean the same thing under any section.
_GLOBAL_LABELS: dict[str, str] = {
    "total revenues": "total_revenue",
    "total revenue": "total_revenue",
    "total operating revenue": "total_revenue",
    "total operating revenues": "total_revenue",
    "total departmental expenses": "dept_expenses",
    "total departmental expense": "dept_expenses",
    "total operated department expenses": "dept_expenses",
    "total departmental profit": "dept_profit",
    "total departmental income": "dept_profit",
    "total undistributed expenses": "undistributed_expenses",
    "total undistributed operating expenses": "undistributed_expenses",
    "gross operating profit": "gop",
    "gross operating profit (gop)": "gop",
    "gop": "gop",
    "management fees": "mgmt_fee",
    "management fee": "mgmt_fee",
    "base management fee": "mgmt_fee",
    "base management fees": "mgmt_fee",
    "income before non-operating income & expenses": "income_before_nonop",
    "income before non-operating income and expenses": "income_before_nonop",
    "income before non operating income & expenses": "income_before_nonop",
    "income before non-operating": "income_before_nonop",
    "ebitda": "ebitda",
    "ff&e proforma calculation": "ffe_reserve",
    "ff&e reserve": "ffe_reserve",
    "ffe reserve": "ffe_reserve",
    "replacement reserve": "ffe_reserve",
    "reserve for replacement": "ffe_reserve",
    "ebitda: less replacement reserve": "noi",
    "ebitda less replacement reserve": "noi",
    "ebitda less reserve": "noi",
    "net operating income": "noi",
    "noi": "noi",
}

#: Concepts the reconciler is willing to replace. Anything else the
#: extractor emitted (KPIs, dates, property facts) is never touched.
RECONCILABLE_CONCEPTS: frozenset[str] = frozenset(
    set(_GLOBAL_LABELS.values())
    | {c for m in _SECTION_SCOPED_LABELS.values() for c in m.values()}
)

_MONTH_RE = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\b")
_TOTAL_HEADER_RE = re.compile(
    r"^(?:(?:fy|cy)\s*\d{2,4}|\d{4}\s*(?:total|actual)|total\s*\d{4}|\d{4}\.total|"
    r"total|totals|annual|full\s*year|year|12\s*months|twelve\s*months|ttm|t-?12|"
    r"trailing\s*twelve(?:\s*months)?)$"
)
_YTD_HEADER_RE = re.compile(r"^(?:ytd|year\s*to\s*date|\d{4}\s*ytd|ytd\s*\d{4})$")


_EN_DASH = chr(0x2013)
_EM_DASH = chr(0x2014)
_DASH_ONLY: frozenset[str] = frozenset({"-", _EN_DASH, _EM_DASH})


def _norm_label(raw: Any) -> str:
    """Case / whitespace / ``and``-vs-``&`` insensitive label key.

    ``"FF&E Reserve"`` and ``"Food and Beverage"`` normalise to
    ``"ff & e reserve"`` and ``"food & beverage"``; the label maps and the
    Summary sheet-name set below are passed through the same function at
    import so lookups never depend on how a key was typed.
    """
    s = str(raw if raw is not None else "").strip().lower()
    s = s.replace(_EN_DASH, "-").replace(_EM_DASH, "-")
    s = re.sub(r"\s+and\s+", " & ", s)
    s = s.replace(" and ", " & ")
    s = re.sub(r"\s*&\s*", " & ", s)
    s = re.sub(r"\s+", " ", s)
    s = s.rstrip(":").strip()
    return s


# Normalise every lookup key exactly the way sheet labels are normalised.
SUMMARY_SHEET_NAMES = frozenset(_norm_label(n) for n in SUMMARY_SHEET_NAMES)
_SECTION_HEADERS = {_norm_label(k): v for k, v in _SECTION_HEADERS.items()}
_GLOBAL_LABELS = {_norm_label(k): v for k, v in _GLOBAL_LABELS.items()}
_SECTION_SCOPED_LABELS = {
    section: {_norm_label(k): v for k, v in labels.items()}
    for section, labels in _SECTION_SCOPED_LABELS.items()
}


_NUM_CLEAN_RE = re.compile(r"[,$\s]")


def _to_number(raw: Any) -> float | None:
    """Parse a grid cell to a float; ``None`` for blanks, text, errors, percents."""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, int | float):
        return float(raw) if math.isfinite(float(raw)) else None
    s = str(raw).strip()
    if not s or s in _DASH_ONLY or s.startswith("#"):
        return None
    if s.endswith("%"):
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg = True
        s = s[1:-1]
    if s.endswith("-") and len(s) > 1:
        neg = True
        s = s[:-1]
    s = _NUM_CLEAN_RE.sub("", s)
    if s.startswith("-"):
        neg = not neg
        s = s[1:]
    if not s:
        return None
    try:
        v = float(s)
    except ValueError:
        return None
    if not math.isfinite(v):
        return None
    return -v if neg else v


# ─────────────────────────── data shapes ───────────────────────────


@dataclass(frozen=True)
class SummaryRow:
    concept: str
    label: str
    row: int  #: Excel row number (workbook path) or 1-based grid row (cache path)
    annual: float
    monthly: tuple[float | None, ...]
    text: str  #: the row as cited on the reconciled field's ``raw_text``


@dataclass
class SummaryTable:
    sheet_name: str
    page_index: int
    annual_column_rule: str  #: ``"sum_of_months"`` / ``"labelled_total"``
    source: str  #: ``"workbook"`` / ``"parser_cache"``
    rows: dict[str, SummaryRow] = field(default_factory=dict)


@dataclass(frozen=True)
class ReconciledChange:
    field_name: str
    concept: str
    old_value: float
    new_value: float
    sheet: str
    row: int
    label: str


@dataclass
class ReconcileResult:
    fields: list[dict[str, Any]]
    confidence: dict[str, Any]
    changes: list[ReconciledChange] = field(default_factory=list)
    table: SummaryTable | None = None
    skipped_reason: str | None = None

    @property
    def changed(self) -> bool:
        return bool(self.changes)


# ─────────────────────────── sheet discovery ───────────────────────────


def find_summary_page(pages: Iterable[Mapping[str, Any]] | None) -> Mapping[str, Any] | None:
    """The cached parser page whose ``metadata.sheet_name`` is a Summary sheet."""
    if not pages:
        return None
    for p in pages:
        if not isinstance(p, Mapping):
            continue
        meta = p.get("metadata") or {}
        name = meta.get("sheet_name") if isinstance(meta, Mapping) else None
        if isinstance(name, str) and _norm_label(name) in SUMMARY_SHEET_NAMES:
            return p
    return None


def summary_grid_from_workbook(
    file_bytes: bytes, sheet_name: str
) -> tuple[list[list[Any]], list[int]] | None:
    """Re-read ``sheet_name`` from the workbook bytes at full precision.

    Returns ``(grid, excel_row_numbers)`` or ``None`` when openpyxl cannot
    open the bytes / the sheet is absent (``.xls`` via xlrd has no float
    precision problem to begin with — the cache path is exact there).
    """
    try:
        import io

        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover — openpyxl is a hard dependency
        return None
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:  # fall back to the parser cache
        logger.info("summary reconcile: openpyxl could not open workbook (%s)", exc)
        return None
    try:
        target = None
        for ws in wb.worksheets:
            if ws.title == sheet_name or _norm_label(ws.title) == _norm_label(sheet_name):
                target = ws
                break
        if target is None:
            return None
        grid: list[list[Any]] = []
        rows: list[int] = []
        for idx, row in enumerate(target.iter_rows(values_only=True), start=1):
            cells = list(row)
            if not any(c is not None and str(c).strip() for c in cells):
                continue
            grid.append(cells)
            rows.append(idx)
        return grid, rows
    finally:
        with contextlib.suppress(Exception):
            wb.close()


# ─────────────────────────── grid parsing ───────────────────────────


def _is_text(cell: Any) -> bool:
    return isinstance(cell, str) and bool(cell.strip()) and _to_number(cell) is None


def _find_month_header(grid: Sequence[Sequence[Any]]) -> tuple[int, list[int]] | None:
    """``(header_row_idx, month_column_indices)`` for the row with most month cells."""
    best: tuple[int, list[int]] | None = None
    for i, row in enumerate(grid):
        cols = [
            j
            for j, c in enumerate(row)
            if isinstance(c, str) and len(c.strip()) <= 12 and _MONTH_RE.search(c.strip().lower())
        ]
        if len(cols) >= 3 and (best is None or len(cols) > len(best[1])):
            best = (i, cols)
    return best


def _find_labelled_total_column(
    grid: Sequence[Sequence[Any]],
    header_row: int | None,
    month_cols: Sequence[int],
    *,
    allow_ytd: bool,
) -> int | None:
    rows_to_scan: list[int] = []
    if header_row is not None:
        rows_to_scan = [r for r in (header_row - 1, header_row, header_row + 1) if 0 <= r < len(grid)]
    else:
        rows_to_scan = list(range(min(len(grid), 40)))
    after = max(month_cols) if month_cols else -1
    for r in rows_to_scan:
        for j, c in enumerate(grid[r]):
            if j <= after or not isinstance(c, str):
                continue
            lab = _norm_label(c)
            if _TOTAL_HEADER_RE.match(lab) or (allow_ytd and _YTD_HEADER_RE.match(lab)):
                return j
    return None


def _row_label(row: Sequence[Any], first_value_col: int | None) -> str | None:
    """The label cell: rightmost text cell before the first numeric cell."""
    limit = first_value_col if first_value_col is not None else len(row)
    label: str | None = None
    for j, c in enumerate(row):
        if j >= limit:
            break
        if _is_text(c):
            label = str(c).strip()
        elif _to_number(c) is not None and label is not None:
            break
    if label is None and first_value_col is None:
        return None
    return label


def _concept_for_label(label: str, section: str | None) -> str | None:
    lab = _norm_label(label)
    if not lab:
        return None
    if lab in _GLOBAL_LABELS:
        return _GLOBAL_LABELS[lab]
    # Prefix match for the long USALI line names that layouts abbreviate
    # differently ("Income Before Non-Operating Income and Expenses").
    if lab.startswith("income before non"):
        return "income_before_nonop"
    if lab.startswith("ebitda: less replacement") or lab.startswith("ebitda less replacement"):
        return "noi"
    if lab.startswith("total non-operating") or lab.startswith("total non operating"):
        return None  # includes non-op income / rent / other — not a Fondok concept
    if section and section in _SECTION_SCOPED_LABELS:
        return _SECTION_SCOPED_LABELS[section].get(lab)
    return None


def _fmt_cell(v: float | None) -> str:
    if v is None:
        return ""
    return f"{v:,.2f}"


def parse_summary_grid(
    grid: Sequence[Sequence[Any]],
    *,
    sheet_name: str,
    page_index: int,
    row_numbers: Sequence[int] | None = None,
    source: str = "parser_cache",
    doc_basis: str = "FY",
) -> SummaryTable | None:
    """Read a Summary sheet grid by USALI label.

    ``row_numbers[i]`` is the display row for ``grid[i]`` (Excel row when the
    grid came from the workbook; defaults to the 1-based grid index).
    ``doc_basis`` is the document's own period basis (``FY`` / ``T12`` /
    ``YTD`` / ``MONTHLY`` …) and decides whether a ``YTD``-labelled column
    may serve as the total.
    """
    if not grid:
        return None
    rows_disp = list(row_numbers) if row_numbers is not None else list(range(1, len(grid) + 1))
    if len(rows_disp) != len(grid):
        rows_disp = list(range(1, len(grid) + 1))

    header = _find_month_header(grid)
    header_row = header[0] if header else None
    month_cols: list[int] = header[1] if header else []
    if len(month_cols) > 12:
        month_cols = month_cols[:12]
    first_value_col = min(month_cols) if month_cols else None

    # Pass 1 — walk rows, track the section, collect labelled numeric rows.
    section: str | None = None
    labelled: list[tuple[int, str, str, Sequence[Any]]] = []  # (grid_idx, label, concept, row)
    for i, row in enumerate(grid):
        if header_row is not None and i <= header_row:
            continue
        label = _row_label(row, first_value_col)
        if not label:
            continue
        numeric_cells = [j for j, c in enumerate(row) if _to_number(c) is not None]
        lab = _norm_label(label)
        if not numeric_cells:
            # A section header re-scopes the ambiguous labels below it. An
            # unknown header leaves the section as-is (sub-headings inside
            # a section — "Labor & Related" — must not reset it).
            if lab in _SECTION_HEADERS:
                section = _SECTION_HEADERS[lab]
            continue
        if section == "skip":
            continue
        concept = _concept_for_label(label, section)
        if concept is None:
            continue
        labelled.append((i, label, concept, row))

    if not labelled:
        return None

    # Pass 2 — annual column: rightmost column equal to Σ months on ≥ 80% of
    # the USALI rows that carry 12 numeric months; else the labelled total.
    annual_col: int | None = None
    rule: str | None = None
    allow_sum_rule = len(month_cols) == 12 or (doc_basis == "YTD" and 1 <= len(month_cols) <= 12)
    if allow_sum_rule:
        tested: dict[int, list[bool]] = {}
        for _i, _label, _concept, row in labelled:
            months = [_to_number(row[j]) if j < len(row) else None for j in month_cols]
            if any(m is None for m in months):
                continue
            total = sum(m for m in months if m is not None)
            for j, c in enumerate(row):
                if j in month_cols or j <= max(month_cols):
                    continue
                v = _to_number(c)
                if v is None:
                    continue
                tol = max(abs(total) * ANNUAL_SUM_TOLERANCE, 1.0)
                tested.setdefault(j, []).append(abs(v - total) <= tol)
        candidates = [
            j for j, hits in tested.items() if len(hits) >= 3 and sum(hits) / len(hits) >= 0.8
        ]
        if candidates:
            annual_col = max(candidates)
            rule = "sum_of_months"
    if annual_col is None:
        annual_col = _find_labelled_total_column(
            grid, header_row, month_cols, allow_ytd=(doc_basis == "YTD")
        )
        if annual_col is not None:
            rule = "labelled_total"
    if annual_col is None or rule is None:
        return None

    table = SummaryTable(
        sheet_name=sheet_name, page_index=page_index, annual_column_rule=rule, source=source
    )
    for i, label, concept, row in labelled:
        if concept in table.rows:
            continue  # first statement of a concept wins; later blocks are restatements
        annual = _to_number(row[annual_col]) if annual_col < len(row) else None
        if annual is None:
            continue
        annual = round(annual, 2)  # cents; strips binary float noise (…0.7100000004)
        months = tuple(_to_number(row[j]) if j < len(row) else None for j in month_cols)
        parts = [label, *(_fmt_cell(m) for m in months), f"{_fmt_cell(annual)} (annual total)"]
        text = f"{sheet_name} row {rows_disp[i]}: " + "\t".join(parts)
        table.rows[concept] = SummaryRow(
            concept=concept,
            label=label,
            row=rows_disp[i],
            annual=annual,
            monthly=months,
            text=text[:3900],
        )
    return table if table.rows else None


# ─────────────────────────── field reconciliation ───────────────────────────


def _field_concept(field_name: str, doc_type: str | None) -> str | None:
    """Registry concept for an extracted path when it is an annual actual line."""
    from ..ontology import registry as ontology

    reg = ontology.get_registry()
    hit = ontology.concept_for_path(field_name, doc_type=doc_type)
    if hit is None:
        return None
    concept, basis, scope = hit
    if concept not in RECONCILABLE_CONCEPTS or basis != "actual":
        return None
    is_slice, _ = ontology._subordinate_scope(field_name.strip().lower(), reg._subordinate)
    if is_slice or scope not in ("annual", "ttm", "unknown"):
        return None
    return concept


def _is_whole_dollar(unit: Any) -> bool:
    if unit is None:
        return True
    return str(unit).strip().lower() in _WHOLE_DOLLAR_UNITS


def _refresh_confidence(confidence: Mapping[str, Any] | None, fields: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    out: dict[str, Any] = dict(confidence or {})
    by_field = dict(out.get("by_field") or {})
    for f in fields:
        name = f.get("field_name")
        conf = f.get("confidence")
        if name and isinstance(conf, int | float):
            by_field[name] = float(conf)
    out["by_field"] = by_field
    if by_field:
        out["overall"] = sum(by_field.values()) / len(by_field)
    out["low_confidence_fields"] = [n for n, c in by_field.items() if c < 0.85]
    return out


def reconcile_fields(
    fields: Sequence[Mapping[str, Any]],
    confidence: Mapping[str, Any] | None,
    table: SummaryTable,
    *,
    doc_type: str | None,
) -> ReconcileResult:
    """Replace extracted USALI totals that disagree with ``table`` by > 5%."""
    dt = canonical_doc_type(doc_type) or None
    out_fields: list[dict[str, Any]] = []
    changes: list[ReconciledChange] = []
    for f in fields:
        if not isinstance(f, Mapping):
            continue
        fd = dict(f)
        name = fd.get("field_name")
        value = fd.get("value")
        if (
            not isinstance(name, str)
            or not name
            or not isinstance(value, int | float)
            or isinstance(value, bool)
            or not _is_whole_dollar(fd.get("unit"))
        ):
            out_fields.append(fd)
            continue
        concept = _field_concept(name, dt)
        row = table.rows.get(concept) if concept else None
        if row is None:
            out_fields.append(fd)
            continue
        old = float(value)
        if abs(old - row.annual) <= RECONCILE_TOLERANCE * abs(row.annual):
            out_fields.append(fd)
            continue
        old_page = fd.get("source_page")
        old_raw = fd.get("raw_text")
        fd["value"] = row.annual
        fd["confidence"] = RECONCILED_CONFIDENCE
        fd["raw_text"] = row.text
        fd["source_page"] = table.page_index
        fd["reviewed"] = "reconciled"
        fd["reconciled_from"] = {
            "field_name": name,
            "old_value": old,
            "sheet": table.sheet_name,
            "row": row.row,
            "label": row.label,
            "old_source_page": old_page,
            "old_raw_text": (str(old_raw)[:200] if old_raw else None),
        }
        fd["note"] = (
            f"Reconciled to the {table.sheet_name} sheet's '{row.label}' annual total "
            f"(row {row.row}): {row.annual:,.2f}. The extractor read {old:,.2f}"
            + (f" from page {old_page}" if old_page is not None else "")
            + "."
        )
        out_fields.append(fd)
        changes.append(
            ReconciledChange(
                field_name=name,
                concept=row.concept,
                old_value=old,
                new_value=row.annual,
                sheet=table.sheet_name,
                row=row.row,
                label=row.label,
            )
        )

    conf = dict(confidence or {})
    if changes:
        conf = _refresh_confidence(conf, out_fields)
        conf["summary_reconciliation"] = {
            "sheet": table.sheet_name,
            "page": table.page_index,
            "annual_column_rule": table.annual_column_rule,
            "source": table.source,
            "changes": [
                {
                    "field_name": c.field_name,
                    "concept": c.concept,
                    "old_value": c.old_value,
                    "new_value": c.new_value,
                    "row": c.row,
                    "label": c.label,
                }
                for c in changes
            ],
        }
    return ReconcileResult(fields=out_fields, confidence=conf, changes=changes, table=table)


def _doc_period_basis(fields: Sequence[Mapping[str, Any]], doc_type: str | None) -> str:
    """``FY`` / ``T12`` / ``YTD`` / ``MONTHLY`` … via the shared resolver."""
    try:
        from ..engines.historical_baseline import _period_basis

        basis, _partial = _period_basis(list(fields), canonical_doc_type(doc_type) or None)
        return basis
    except Exception:  # never let the basis lookup block the pass
        return "FY"


def reconcile_extraction(
    fields: Sequence[Mapping[str, Any]],
    confidence: Mapping[str, Any] | None,
    *,
    doc_type: str | None,
    extraction_data: Mapping[str, Any] | None,
    file_bytes: bytes | None = None,
) -> ReconcileResult:
    """Top-level entry: gate, locate the Summary sheet, parse it, reconcile.

    Returns the input unchanged (``skipped_reason`` set) when the document
    is not P&L-family, has no Summary sheet, is a monthly / quarterly
    statement, or the sheet has no usable annual column.
    """
    base_fields = [dict(f) for f in fields if isinstance(f, Mapping)]
    base_conf = dict(confidence or {})

    def _skip(reason: str) -> ReconcileResult:
        return ReconcileResult(fields=base_fields, confidence=base_conf, skipped_reason=reason)

    if not is_pnl_family(doc_type):
        return _skip("not_pnl_family")
    if not base_fields:
        return _skip("no_fields")
    pages = (extraction_data or {}).get("pages") if isinstance(extraction_data, Mapping) else None
    page = find_summary_page(pages)
    if page is None:
        return _skip("no_summary_sheet")

    basis = _doc_period_basis(base_fields, doc_type)
    if basis in _PARTIAL_BASES_NEVER_RECONCILED:
        return _skip(f"partial_period:{basis}")

    meta = page.get("metadata") or {}
    sheet_name = str(meta.get("sheet_name") or "Summary")
    try:
        page_index = int(page.get("page_num") or 0)
    except (TypeError, ValueError):
        page_index = 0

    grid: Sequence[Sequence[Any]] | None = None
    row_numbers: Sequence[int] | None = None
    source = "parser_cache"
    if file_bytes:
        wb_grid = summary_grid_from_workbook(file_bytes, sheet_name)
        if wb_grid is not None:
            grid, row_numbers = wb_grid
            source = "workbook"
    if grid is None:
        tables = page.get("tables") or []
        grid = tables[0] if tables and isinstance(tables[0], list) else None
    if not grid:
        return _skip("empty_summary_grid")

    table = parse_summary_grid(
        grid,
        sheet_name=sheet_name,
        page_index=page_index,
        row_numbers=row_numbers,
        source=source,
        doc_basis=basis,
    )
    if table is None:
        return _skip("no_annual_column")
    result = reconcile_fields(base_fields, base_conf, table, doc_type=doc_type)
    if result.changes:
        logger.info(
            "summary reconcile: sheet=%s page=%s rule=%s source=%s replaced %d field(s): %s",
            table.sheet_name,
            table.page_index,
            table.annual_column_rule,
            table.source,
            len(result.changes),
            ", ".join(f"{c.field_name} {c.old_value:,.2f}→{c.new_value:,.2f}" for c in result.changes),
        )
    return result


__all__ = [
    "ANNUAL_SUM_TOLERANCE",
    "PNL_FAMILY_DOC_TYPES",
    "RECONCILABLE_CONCEPTS",
    "RECONCILED_CONFIDENCE",
    "RECONCILE_TOLERANCE",
    "SUMMARY_SHEET_NAMES",
    "ReconcileResult",
    "ReconciledChange",
    "SummaryRow",
    "SummaryTable",
    "canonical_doc_type",
    "find_summary_page",
    "is_pnl_family",
    "parse_summary_grid",
    "reconcile_extraction",
    "reconcile_fields",
    "summary_grid_from_workbook",
]
