"""FON-41 E-013 / E-017 — editable P&L workbooks (export + validated re-import).

Pure openpyxl builders and parsers. No DB access here: the API layer
(``app/api/pl_roundtrip.py``) gathers the CURRENT state (extractions, engine
outputs, assumptions) and hands it in, so every diff below is computed against
what the deal holds right now — never against what the file claims was there.

Historical P&L (``fondok.historicals.v1``)
------------------------------------------
Sheet ``Historical P&L``: rows are USALI line items, one column PAIR per
historical statement — the visible value column, then a hidden ``fondok_id``
column whose cell holds the stable id ``<document_id>::<field_name>`` of the
extracted line that value came from. Row 1 is the period header, row 2 names
the source document (the id column carries its ``document_id``), and the
hidden column A carries the line's concept id. A cell with no id is a line the
statement never published — there is nothing at source to correct.

Future P&L (``fondok.projections.v1``)
--------------------------------------
Sheet ``Future P&L``: engine-computed projection lines (read-only — a changed
value is reported as "computed — edit the assumption instead"), hidden column A
holds the engine path. Each model year is a column GROUP — Amount, % Rev, PAR,
POR — followed by one display-only Exit Year column (the reversion NOI); the
hidden row 2 (``fondok_column``) names every column so the importer finds the
Amount columns by key, not by position. Rows whose column A starts with
``fondok_readonly:`` (Days, Number of Rooms, Available Rooms, Occupied Rooms,
Net Cash Flow, the exit NOI) and every column but Amount are display-only and
never diffed. Sheet ``Assumptions``: every analyst-editable Future P&L
assumption with its stable ``field_overrides`` key, current value, source,
method, whether a note is required, and a Note column for the justification.
Only the Assumptions sheet is importable.

Both workbooks carry a ``README`` sheet whose ``format`` / ``deal_id`` rows the
importer checks first, so a file for another deal (or another format) is
refused outright instead of being half-applied.
"""

from __future__ import annotations

import io
import math
import re
from collections.abc import Callable, Iterator, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .refusals import NOT_AVAILABLE_PREFIX, ReasonCode, refuse, unavailable

HIST_FORMAT = "fondok.historicals.v1"
PROJ_FORMAT = "fondok.projections.v1"
HIST_SHEET = "Historical P&L"
PROJ_SHEET = "Future P&L"
ASSUMPTION_SHEET = "Assumptions"
README_SHEET = "README"
ID_HEADER = "fondok_id"
LINE_HEADER = "fondok_line"
PATH_HEADER = "fondok_path"
#: Column A of the Future P&L sheet's hidden row 2; the row names each column
#: (``amount:<i>`` / ``pct_rev:<i>`` / ``par:<i>`` / ``por:<i>`` / ``exit`` /
#: ``benchmark``). A sheet without it is the earlier positional layout.
COLUMN_KEY_HEADER = "fondok_column"
#: Column-A prefix of a display-only Future P&L row the importer never reads.
READONLY_ROW_PREFIX = "fondok_readonly:"

_HEADER_FILL = PatternFill("solid", fgColor="1A2233")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SUB_FONT = Font(italic=True, color="6B6F76")
_READONLY_FILL = PatternFill("solid", fgColor="F3F2EE")

# ─────────────────────────────── line models ───────────────────────────────

#: USALI line items of the Historical P&L, in statement order. Concept ids
#: from ``app/ontology/concepts.yaml`` — the resolver that picks each
#: statement's line is the registry's, so the export names the SAME extracted
#: field the worksheet and the engines read.
HIST_LINES: tuple[tuple[str, str], ...] = (
    ("occupancy", "Occupancy"),
    ("adr", "ADR"),
    ("revpar", "RevPAR"),
    ("rooms_revenue", "Rooms Revenue"),
    ("fb_revenue", "Food & Beverage Revenue"),
    ("other_revenue", "Other Operated Departments Revenue"),
    ("misc_revenue", "Miscellaneous Income"),
    ("resort_fees", "Resort Fees"),
    ("total_revenue", "Total Operating Revenue"),
    ("rooms_dept_expense", "Rooms Department Expense"),
    ("fb_dept_expense", "Food & Beverage Department Expense"),
    ("other_dept_expense", "Other Operated Departments Expense"),
    ("dept_expenses", "Total Departmental Expenses"),
    ("dept_profit", "Total Departmental Profit"),
    ("administrative_general", "Administrative & General"),
    ("information_telecom", "Information & Telecommunications"),
    ("sales_marketing", "Sales & Marketing"),
    ("property_operations", "Property Operations & Maintenance"),
    ("utilities", "Utilities"),
    ("undistributed_expenses", "Total Undistributed Expenses"),
    ("gop", "Gross Operating Profit"),
    ("mgmt_fee", "Management Fee"),
    ("franchise_royalty_fee", "Franchise Royalty Fee"),
    ("property_taxes", "Property Taxes"),
    ("insurance", "Insurance"),
    ("rent_expense", "Rent (ground lease)"),
    ("ffe_reserve", "FF&E Reserve"),
    ("noi", "Net Operating Income"),
)

#: Future P&L lines: (engine, dotted path inside ``years[i]``, label). Exactly
#: what the engines publish — nothing is derived here.
PROJ_LINES: tuple[tuple[str, str, str], ...] = (
    ("revenue", "occupancy", "Occupancy"),
    ("revenue", "adr", "ADR"),
    ("revenue", "revpar", "RevPAR"),
    ("revenue", "rooms_revenue", "Rooms Revenue"),
    ("fb", "fb_revenue", "Food & Beverage Revenue"),
    ("fb", "other_revenue", "Other Operated Departments Revenue"),
    ("fb", "resort_fees", "Resort Fees"),
    ("expense", "total_revenue", "Total Revenue"),
    ("expense", "dept_expenses.rooms", "Rooms Department Expense"),
    ("expense", "dept_expenses.food_beverage", "F&B Department Expense"),
    ("expense", "dept_expenses.other_operated", "Other Operated Departments Expense"),
    ("expense", "dept_expenses.total", "Total Departmental Expenses"),
    ("expense", "undistributed.administrative_general", "Administrative & General"),
    ("expense", "undistributed.information_telecom", "Information & Telecommunications"),
    ("expense", "undistributed.sales_marketing", "Sales & Marketing"),
    ("expense", "undistributed.property_operations", "Property Operations & Maintenance"),
    ("expense", "undistributed.utilities", "Utilities"),
    ("expense", "undistributed.total", "Total Undistributed Expenses"),
    ("expense", "gop", "Gross Operating Profit"),
    ("expense", "mgmt_fee", "Management Fee"),
    # Sam #9 — only written when the run modelled an incentive fee (see
    # ``OPTIONAL_PROJ_PATHS``); otherwise the row is omitted, never blank-filled.
    ("expense", "incentive_mgmt_fee", "Incentive Management Fee"),
    ("expense", "fixed_charges.property_taxes", "Property Taxes"),
    ("expense", "fixed_charges.insurance", "Insurance"),
    ("expense", "fixed_charges.rent", "Rent"),
    ("expense", "fixed_charges.other_fixed", "Other Fixed Charges"),
    ("expense", "fixed_charges.total", "Total Fixed Charges"),
    ("expense", "noi_institutional", "NOI (before FF&E reserve)"),
    ("expense", "ffe_reserve", "FF&E Reserve"),
)

#: ``PROJ_LINES`` rows written only when the run carries a value for them in
#: at least one year. Sam #9 — the incentive management fee is opt-in, so a
#: deal without one gets no row at all rather than a row of blanks.
OPTIONAL_PROJ_PATHS: frozenset[str] = frozenset({"expense.incentive_mgmt_fee"})

#: FON-97 — the Future P&L sheet's read-only "Benchmark %" column header.
#: The import parser skips any column whose header starts with this word.
BENCHMARK_HEADER = "Benchmark %"

#: FON-97 — ``services.pnl_benchmark_map`` category key → the ``PROJ_LINES``
#: path it benchmarks. Same 13 lines the web's Benchmark column wires; a
#: path missing here has no benchmark and its cell stays blank.
BENCHMARK_LINE_PATHS: dict[str, str] = {
    "rooms": "expense.dept_expenses.rooms",
    "food_beverage": "expense.dept_expenses.food_beverage",
    "other_operated": "expense.dept_expenses.other_operated",
    "administrative_general": "expense.undistributed.administrative_general",
    "information_telecom": "expense.undistributed.information_telecom",
    "sales_marketing": "expense.undistributed.sales_marketing",
    "property_operations": "expense.undistributed.property_operations",
    "utilities": "expense.undistributed.utilities",
    "gop": "expense.gop",
    "mgmt_fee": "expense.mgmt_fee",
    "property_taxes": "expense.fixed_charges.property_taxes",
    "insurance": "expense.fixed_charges.insurance",
    "rent": "expense.fixed_charges.rent",
}


def benchmark_column_values(
    categories: Sequence[Mapping[str, Any] | Any] | None,
) -> dict[str, float]:
    """``{"<engine>.<path>": ratio}`` for the workbook's Benchmark % column.

    ``categories`` is the market-data block's ``pnl_benchmark.categories``
    (dicts or ``PnlBenchmarkCategory`` models). Only categories that carry a
    ratio contribute — a line the report priced per room but gave no ratio
    for is left blank, never derived here.
    """
    out: dict[str, float] = {}
    for cat in categories or ():
        get = cat.get if isinstance(cat, Mapping) else (lambda k, _c=cat: getattr(_c, k, None))
        key = get("key")
        ratio = to_number(get("ratio"))
        path = BENCHMARK_LINE_PATHS.get(str(key)) if key is not None else None
        if path is not None and ratio is not None:
            out[path] = ratio
    return out


@dataclass(frozen=True)
class AssumptionSpec:
    """One analyst-editable Future P&L assumption (a ``field_overrides`` key)."""

    key: str
    label: str
    unit: str
    method: str
    integer: bool = False
    #: Non-empty = a text choice key (E-016 projection method); the Value cell
    #: must be one of these codes, compared case-insensitively.
    choices: tuple[str, ...] = ()


#: Every assumption the Future P&L lets an analyst edit (driver cells, the
#: Assumptions panel and the projection period). The keys are the SAME
#: ``field_overrides`` keys ``ProjectionsSection.tsx`` writes; the method text
#: restates how the projection applies each one, from that panel's own copy.
ASSUMPTIONS: tuple[AssumptionSpec, ...] = (
    AssumptionSpec(
        "starting_occupancy", "Base-year occupancy", "fraction (0.762 = 76.2%)",
        "Year-1 starting point; later years follow the occupancy path",
    ),
    AssumptionSpec(
        "starting_adr", "Base-year ADR", "USD",
        "Year-1 starting point; later years grow from it",
    ),
    AssumptionSpec(
        "revpar_growth", "RevPAR growth", "fraction per year (0.045 = 4.5%)",
        "Compounds annually; drives ADR growth with the occupancy path held",
    ),
    AssumptionSpec(
        "expense_growth", "Dept. expense inflation", "fraction per year",
        "Compounds annually on departmental expenses",
    ),
    AssumptionSpec(
        "other_expense_growth", "Other expense inflation", "fraction per year",
        "Compounds annually on undistributed expenses; blank = follows dept. inflation",
    ),
    AssumptionSpec(
        "resort_fee_per_night", "Resort fee", "USD per night",
        "Charged per night, scaled by the capture rates below",
    ),
    AssumptionSpec(
        "resort_fee_capture_y1", "Resort fee capture — Base year (Year 1)", "fraction",
        "Applies to the Base year (Year 1) column",
    ),
    AssumptionSpec(
        "resort_fee_capture_y2", "Resort fee capture — Year 2", "fraction",
        "Applies to the Year 2 column",
    ),
    AssumptionSpec(
        "resort_fee_capture_y3", "Resort fee capture — Year 3+", "fraction",
        "Applies to Year 3 and every later year",
    ),
    AssumptionSpec(
        "mgmt_fee_pct", "Management fee", "fraction of total revenue",
        "Applied to each year's total revenue",
    ),
    AssumptionSpec(
        "incentive_mgmt_fee_pct_of_gop", "Incentive management fee",
        "fraction of GOP after the base fee (0.10 = 10%)",
        "Each year: max(0, GOP − base management fee) × this rate, deducted "
        "before NOI; blank = no incentive fee",
    ),
    AssumptionSpec(
        "hold_years", "Projection period (hold)", "years",
        "Number of modelled years", integer=True,
    ),
    AssumptionSpec(
        "stabilization_year", "Stabilization year", "model year (1-based)",
        "Selects which projection year the stabilized figures are read from; moves no return",
        integer=True,
    ),
)
ASSUMPTION_BY_KEY: dict[str, AssumptionSpec] = {a.key: a for a in ASSUMPTIONS}

#: E-016 — per-line projection methods (``projection_methods.<line>.method`` /
#: ``.value``), the same keys the Future P&L method chip writes. Line names
#: and order mirror ``engines/expense.py`` ``PROJECTION_METHOD_LINES``.
PROJECTION_METHOD_CHOICES: tuple[str, ...] = ("growth", "pct_revenue", "por", "par")
PROJECTION_METHOD_LINE_LABELS: tuple[tuple[str, str], ...] = (
    ("rooms_dept_expense", "Rooms expense"),
    ("fb_dept_expense", "F&B expense"),
    ("other_dept_expense", "Other operated dept. expense"),
    ("administrative_general", "Administrative & General"),
    ("information_telecom", "Information & Telecom"),
    ("sales_marketing", "Sales & Marketing"),
    ("property_operations", "Property Operations & Maintenance"),
    ("utilities", "Utilities"),
)


def projection_method_key(line: str, field: str) -> str:
    return f"projection_methods.{line}.{field}"


def split_projection_method_key(key: str) -> tuple[str, str] | None:
    parts = key.split(".")
    if len(parts) == 3 and parts[0] == "projection_methods":
        return parts[1], parts[2]
    return None


PROJECTION_METHOD_ASSUMPTIONS: tuple[AssumptionSpec, ...] = tuple(
    spec
    for line, label in PROJECTION_METHOD_LINE_LABELS
    for spec in (
        AssumptionSpec(
            projection_method_key(line, "method"), f"{label} — projection method",
            "growth | pct_revenue | por | par",
            "growth = Year-1 grown each year; pct_revenue = share of revenue; "
            "por = $ per occupied room; par = $ per available room",
            choices=PROJECTION_METHOD_CHOICES,
        ),
        AssumptionSpec(
            projection_method_key(line, "value"), f"{label} — method value",
            "growth / pct_revenue: fraction (0.03 = 3%); por / par: USD per room",
            "Driver for the method above; changing the method without a value "
            "holds the line's own Year-1 ratio",
        ),
    )
)
ASSUMPTIONS = ASSUMPTIONS + PROJECTION_METHOD_ASSUMPTIONS
ASSUMPTION_BY_KEY = {a.key: a for a in ASSUMPTIONS}

#: FON-95 part 2 — Index Analysis penetration targets BY MODEL YEAR. One sheet
#: row per metric per model year (``index_mpi_target_by_year.3`` = the MPI
#: target for Year 3, index points, 100 = parity with the selected benchmark);
#: the import folds the rows back into the ``index_*_target_by_year`` JSON
#: lists the engine runner reads. Blank = no by-year target for that year (the
#: all-years ``index_mpi_target`` / ``index_ari_target`` shows where set, but
#: drives Year 1 only unless the analyst overrode that metric's market growth).
#: The row count follows the deal's ``hold_years``, so these are generated per
#: export rather than listed in ``ASSUMPTIONS``.
BY_YEAR_TARGET_LISTS: tuple[tuple[str, str, str], ...] = (
    ("index_mpi_target_by_year", "Occupancy penetration target (MPI)", "occupancy"),
    ("index_ari_target_by_year", "ADR penetration target (ARI)", "ADR"),
)
BY_YEAR_TARGET_MAX_YEARS = 20
_BY_YEAR_TARGET_KEY_RE = re.compile(
    r"^(?P<list>index_(?:mpi|ari)_target_by_year)\.(?P<year>[1-9]\d?)$"
)


def by_year_target_key(list_key: str, year: int) -> str:
    return f"{list_key}.{year}"


def split_by_year_target_key(key: str) -> tuple[str, int] | None:
    """``index_mpi_target_by_year.3`` → ``("index_mpi_target_by_year", 3)``;
    ``None`` for any other key (or a year past the model cap)."""
    m = _BY_YEAR_TARGET_KEY_RE.match(key or "")
    if not m:
        return None
    year = int(m.group("year"))
    if not 1 <= year <= BY_YEAR_TARGET_MAX_YEARS:
        return None
    return m.group("list"), year


def by_year_target_spec(list_key: str, year: int) -> AssumptionSpec:
    label, noun = next(
        ((lbl, n) for k, lbl, n in BY_YEAR_TARGET_LISTS if k == list_key), (list_key, "")
    )
    year_label = "Base year (Year 1)" if year == 1 else f"Year {year}"
    return AssumptionSpec(
        by_year_target_key(list_key, year),
        f"{label} — {year_label}",
        "index points 25–250 (100 = parity with the selected benchmark) or a ratio 0.25–2.5",
        f"Index Analysis target by year: {year_label} {noun} = benchmark {noun} for that year × target "
        f"÷ 100 while \"Use STR rates in the model\" is on; blank = the all-years target "
        f"(drives Year 1 only unless a market {noun} growth override is set); "
        f"a value outside both bands is refused on import, never applied",
    )


def by_year_target_assumptions(hold_years: int) -> tuple[AssumptionSpec, ...]:
    """The by-year target rows for a deal projecting ``hold_years`` years."""
    hold = max(1, min(int(hold_years or 1), BY_YEAR_TARGET_MAX_YEARS))
    return tuple(
        by_year_target_spec(list_key, year)
        for list_key, _label, _noun in BY_YEAR_TARGET_LISTS
        for year in range(1, hold + 1)
    )


#: Sam, 2026-10-15 — the explicit renovation displacement SCHEDULE, one sheet
#: row per period per field (``renovation_displacement_schedule.2.end_month`` =
#: the second period's last month). The import folds the rows back into the
#: ``renovation_displacement_schedule`` JSON list the engine runner reads
#: (``[{start_month, end_month, rooms_offline_pct, rate_discount_pct, note}]``).
#: The export lists every saved period plus ONE blank period (to add the next
#: one), so the row count follows the deal and these are generated per export.
#: Periods are months from the acquisition close, 1-based and inclusive; the
#: engine weighs each month 1/12 of its model year (``engines/revenue.py``).
SCHEDULE_KEY = "renovation_displacement_schedule"
SCHEDULE_MAX_PERIODS = 24
SCHEDULE_FIELDS: tuple[tuple[str, str, str, str, bool], ...] = (
    (
        "start_month", "First month",
        "model month, 1-based from the acquisition close (1 = first operating month)",
        "Start of the renovation period, inclusive; a model year is 12 months from close",
        True,
    ),
    (
        "end_month", "Last month",
        "model month, 1-based from the acquisition close, ≥ first month",
        "End of the renovation period, inclusive",
        True,
    ),
    (
        "rooms_offline_pct", "Rooms offline",
        "fraction of inventory (0.25 = 25% of rooms out of service)",
        "Each month in the period: occupancy × (1 − rooms offline); the year's factor is the "
        "12-month average",
        False,
    ),
    (
        "rate_discount_pct", "Rate discount",
        "fraction of ADR (0.10 = 10% off on the rooms still selling)",
        "Each month in the period: ADR × (1 − rate discount) on the open rooms; the year's "
        "factor is weighted by open rooms. The schedule replaces the Y1 displacement "
        "percentages — never applied on top",
        False,
    ),
)
_SCHEDULE_KEY_RE = re.compile(
    r"^renovation_displacement_schedule\.(?P<period>[1-9]\d?)\."
    r"(?P<field>start_month|end_month|rooms_offline_pct|rate_discount_pct)$"
)


def schedule_period_key(period: int, field: str) -> str:
    return f"{SCHEDULE_KEY}.{period}.{field}"


def split_schedule_period_key(key: str) -> tuple[int, str] | None:
    """``renovation_displacement_schedule.2.end_month`` → ``(2, "end_month")``;
    ``None`` for any other key (or a period past the cap)."""
    m = _SCHEDULE_KEY_RE.match(key or "")
    if not m:
        return None
    period = int(m.group("period"))
    if not 1 <= period <= SCHEDULE_MAX_PERIODS:
        return None
    return period, m.group("field")


def schedule_period_spec(period: int, field: str) -> AssumptionSpec:
    label, unit, method, integer = next(
        ((lbl, u, mth, integ) for f, lbl, u, mth, integ in SCHEDULE_FIELDS if f == field),
        (field, "", "", False),
    )
    return AssumptionSpec(
        schedule_period_key(period, field),
        f"Renovation period {period} — {label}",
        unit,
        f"Renovation displacement schedule: {method}. A new period needs both months; "
        f"a renovation budget alone never moves operations",
        integer=integer,
    )


def schedule_assumptions(saved_periods: int) -> tuple[AssumptionSpec, ...]:
    """The schedule rows for a deal with ``saved_periods`` periods on file:
    every saved period plus one blank period to fill in."""
    n = max(0, min(int(saved_periods or 0), SCHEDULE_MAX_PERIODS - 1)) + 1
    return tuple(
        schedule_period_spec(period, field)
        for period in range(1, n + 1)
        for field, *_rest in SCHEDULE_FIELDS
    )


def assumption_spec(key: str) -> AssumptionSpec | None:
    """The spec behind an Assumptions-sheet key: a fixed ``ASSUMPTIONS`` entry,
    a by-year penetration target row or a renovation-schedule period row;
    ``None`` = not an editable key."""
    spec = ASSUMPTION_BY_KEY.get(key)
    if spec is not None:
        return spec
    by = split_by_year_target_key(key)
    if by is not None:
        return by_year_target_spec(*by)
    sp = split_schedule_period_key(key)
    return schedule_period_spec(*sp) if sp is not None else None


ASSUMPTION_COLUMNS: tuple[str, ...] = (
    "Key", "Assumption", "Value", "Unit", "Source", "Method",
    "Note required", "Current note", "Note (your justification)",
)


# ─────────────────────────────── helpers ───────────────────────────────


class WorkbookFormatError(ValueError):
    """The upload is not a Fondok workbook of the expected format / deal."""


def to_number(v: Any) -> float | None:
    """A finite number from a stored extraction / engine value, else None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return f if math.isfinite(f) else None
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("$", "")
        if not s:
            return None
        try:
            f = float(s)
        except ValueError:
            return None
        return f if math.isfinite(f) else None
    return None


def parse_cell_number(raw: Any) -> tuple[float | None, bool]:
    """``(value, ok)`` for an UPLOADED cell. Blank → ``(None, True)``.

    Accepts numbers and plain numeric text (thousands separators and a leading
    ``$`` are tolerated). A ``%`` suffix, words, or anything else is refused
    (``ok=False``) — values round-trip in the units the export wrote, so the
    importer never guesses whether "73.8%" meant 0.738 or 73.8.
    """
    if raw is None:
        return None, True
    if isinstance(raw, bool):
        return None, False
    if isinstance(raw, (int, float)):
        f = float(raw)
        return (f, True) if math.isfinite(f) else (None, False)
    if isinstance(raw, str):
        s = raw.strip()
        if not s:
            return None, True
        s2 = s.replace(",", "").replace("$", "").strip()
        try:
            f = float(s2)
        except ValueError:
            return None, False
        return (f, True) if math.isfinite(f) else (None, False)
    return None, False


def same_number(a: float | None, b: float | None) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))


def hist_cell_id(document_id: str, field_name: str) -> str:
    return f"{document_id}::{field_name}"


def split_hist_cell_id(cell_id: str) -> tuple[str, str] | None:
    if not isinstance(cell_id, str) or "::" not in cell_id:
        return None
    doc, fname = cell_id.split("::", 1)
    doc, fname = doc.strip(), fname.strip()
    if not doc or not fname:
        return None
    return doc, fname


def proj_cell_key(engine: str, path: str, year_index: int) -> str:
    """The Future P&L cell's stable key — also its comment ``cell_key`` suffix."""
    return f"{engine}.years[{year_index}].{path}"


def _style_header(ws: Any, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _readme(wb: Workbook, *, fmt: str, deal_id: str, deal_name: str | None,
            lines: Sequence[str]) -> None:
    ws = wb.create_sheet(README_SHEET)
    ws["A1"] = "Fondok editable workbook"
    ws["A1"].font = Font(bold=True, size=13)
    meta = (
        ("format", fmt),
        ("deal_id", deal_id),
        ("deal", deal_name or ""),
        ("exported_at", datetime.now(UTC).isoformat(timespec="seconds")),
    )
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    start = 3 + len(meta) + 1
    for i, line in enumerate(lines):
        ws.cell(row=start + i, column=1, value=line)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 60


def _check_readme(wb: Any, *, fmt: str, deal_id: str) -> None:
    if README_SHEET not in wb.sheetnames:
        raise WorkbookFormatError(
            "This is not a Fondok export (no README sheet). Export the workbook "
            "from Fondok, edit it, and import that file."
        )
    ws = wb[README_SHEET]
    meta: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=40, max_col=2, values_only=True):
        if row and isinstance(row[0], str) and row[1] is not None:
            meta.setdefault(row[0].strip(), str(row[1]).strip())
    if meta.get("format") != fmt:
        raise WorkbookFormatError(
            f"Wrong workbook format {meta.get('format')!r} — expected {fmt!r}."
        )
    if meta.get("deal_id", "").lower() != str(deal_id).lower():
        raise WorkbookFormatError(
            "This workbook was exported from a different deal — import refused."
        )


def _load(data: bytes) -> Any:
    try:
        return load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise WorkbookFormatError("The file could not be read as an .xlsx workbook.") from exc


# ─────────────────────────── historical P&L ───────────────────────────


@dataclass(frozen=True)
class HistColumn:
    document_id: str
    label: str
    filename: str | None = None


@dataclass(frozen=True)
class HistCell:
    field_name: str
    value: float | None


def build_historicals_workbook(
    *,
    deal_id: str,
    deal_name: str | None,
    columns: Sequence[HistColumn],
    cells: Mapping[tuple[str, str], HistCell],
    lines: Sequence[tuple[str, str]] = HIST_LINES,
) -> bytes:
    """Render the Historical P&L workbook. ``cells`` is keyed ``(line_id, document_id)``."""
    wb = Workbook()
    ws = wb.active
    ws.title = HIST_SHEET
    ws.cell(row=1, column=1, value=LINE_HEADER)
    ws.cell(row=1, column=2, value="Line item")
    ws.cell(row=2, column=1, value="")
    ws.cell(row=2, column=2, value="Source document")
    for j, col in enumerate(columns):
        vc = 3 + 2 * j
        ws.cell(row=1, column=vc, value=col.label)
        ws.cell(row=1, column=vc + 1, value=ID_HEADER)
        ws.cell(row=2, column=vc, value=col.filename or "")
        ws.cell(row=2, column=vc + 1, value=col.document_id)
    ncols = 2 + 2 * len(columns)
    _style_header(ws, 1, ncols)
    for c in range(1, ncols + 1):
        ws.cell(row=2, column=c).font = _SUB_FONT
    for i, (line_id, label) in enumerate(lines):
        r = 3 + i
        ws.cell(row=r, column=1, value=line_id)
        ws.cell(row=r, column=2, value=label)
        for j, col in enumerate(columns):
            vc = 3 + 2 * j
            cell = cells.get((line_id, col.document_id))
            if cell is None:
                continue
            if cell.value is not None:
                ws.cell(row=r, column=vc, value=cell.value).number_format = "#,##0.####"
            ws.cell(row=r, column=vc + 1, value=hist_cell_id(col.document_id, cell.field_name))
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 38
    for j in range(len(columns)):
        vc = 3 + 2 * j
        ws.column_dimensions[get_column_letter(vc)].width = 16
        ws.column_dimensions[get_column_letter(vc + 1)].hidden = True
    ws.freeze_panes = "C3"
    _readme(
        wb,
        fmt=HIST_FORMAT,
        deal_id=deal_id,
        deal_name=deal_name,
        lines=(
            "Historical P&L — one column per extracted statement, rows are USALI line items.",
            "Edit values in the visible year columns, save as .xlsx, then use Import on the "
            "Historical P&L.",
            "Each value has a hidden 'fondok_id' column beside it holding the stable id of the "
            "extracted line (document_id::field_name). Do not edit, move or delete hidden columns.",
            "Values are exported exactly as extracted (no unit conversion) — type plain numbers in "
            "the same units; '%' and text are rejected.",
            "A blank cell with no id means the statement has no such line: there is nothing at "
            "source to correct, so a value typed there is reported as a mapping error.",
            "Import shows a preview (changed values old → new, mapping errors, non-numeric cells). "
            "Nothing changes until you Apply; each change is saved as an analyst correction at "
            "source, exactly like editing the cell in the Data Room review.",
            "OM-embedded prior years are not included — correct those in the Data Room.",
        ),
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class HistCurrent:
    """What the deal holds right now, for import validation."""

    #: ``cell_id → current numeric value`` (None for a non-numeric / empty value)
    values: dict[str, float | None] = field(default_factory=dict)
    #: ``cell_id → raw stored value`` (shown as the old value when non-numeric)
    raw: dict[str, Any] = field(default_factory=dict)
    #: document ids on the deal that the export may reference
    documents: set[str] = field(default_factory=set)


def parse_historicals_import(
    data: bytes, *, deal_id: str, current: HistCurrent
) -> dict[str, Any]:
    """Diff an uploaded Historical P&L workbook against ``current``.

    Returns ``{changes, mapping_errors, non_numeric, unchanged}``. Never writes.
    """
    wb = _load(data)
    _check_readme(wb, fmt=HIST_FORMAT, deal_id=deal_id)
    if HIST_SHEET not in wb.sheetnames:
        raise WorkbookFormatError(f"Sheet {HIST_SHEET!r} is missing.")
    ws = wb[HIST_SHEET]
    max_col = ws.max_column
    header = [ws.cell(row=1, column=c).value for c in range(1, max_col + 1)]
    pairs: list[tuple[int, int, str]] = []  # (value col, id col, period label)
    for c in range(2, max_col + 1):
        if header[c - 1] == ID_HEADER:
            label = header[c - 2]
            pairs.append((c - 1, c, str(label) if label is not None else f"column {c - 1}"))
    if not pairs:
        raise WorkbookFormatError(
            "No 'fondok_id' columns found — the hidden id columns were removed or renamed."
        )

    changes: list[dict[str, Any]] = []
    mapping_errors: list[dict[str, Any]] = []
    non_numeric: list[dict[str, Any]] = []
    unchanged = 0
    seen: dict[str, float | None] = {}

    for r in range(3, ws.max_row + 1):
        line_id = ws.cell(row=r, column=1).value
        label = ws.cell(row=r, column=2).value
        line_label = str(label) if label is not None else str(line_id or f"row {r}")
        for vc, ic, period in pairs:
            ref = f"{get_column_letter(vc)}{r}"
            raw_val = ws.cell(row=r, column=vc).value
            raw_id = ws.cell(row=r, column=ic).value
            base = {"cell_ref": ref, "line_id": line_id, "line_label": line_label,
                    "period_label": period}
            id_txt = str(raw_id).strip() if raw_id is not None else ""
            if not id_txt:
                if raw_val not in (None, ""):
                    mapping_errors.append({
                        **base, "cell_id": None, "reason": "no_source_line",
                        "detail": "This statement has no extracted line here, so there is "
                                  "nothing at source to correct.",
                    })
                continue
            parts = split_hist_cell_id(id_txt)
            if parts is None or parts[0] not in current.documents or id_txt not in current.raw:
                mapping_errors.append({
                    **base, "cell_id": id_txt, "reason": "unknown_id",
                    "detail": "The hidden id does not match an extracted line on this deal "
                              "(edited id, removed document or re-extracted statement).",
                })
                continue
            new_val, ok = parse_cell_number(raw_val)
            if not ok:
                non_numeric.append({**base, "cell_id": id_txt, "raw": str(raw_val)})
                continue
            old_val = current.values.get(id_txt)
            if new_val is None:
                if old_val is None:
                    unchanged += 1
                else:
                    mapping_errors.append({
                        **base, "cell_id": id_txt, "reason": "cleared",
                        "detail": "A cleared cell cannot be imported — reject the field in "
                                  "the Data Room review instead.",
                    })
                continue
            if id_txt in seen:
                if not same_number(seen[id_txt], new_val):
                    mapping_errors.append({
                        **base, "cell_id": id_txt, "reason": "duplicate_id",
                        "detail": "The same extracted line appears twice with different values.",
                    })
                continue
            seen[id_txt] = new_val
            if same_number(old_val, new_val):
                unchanged += 1
                continue
            doc_id, fname = parts
            changes.append({
                **base,
                "cell_id": id_txt,
                "document_id": doc_id,
                "field_name": fname,
                "old_value": old_val if old_val is not None else current.raw.get(id_txt),
                "new_value": new_val,
            })
    # A later duplicate can invalidate an earlier change of the same id.
    dup_ids = {e["cell_id"] for e in mapping_errors if e["reason"] == "duplicate_id"}
    changes = [c for c in changes if c["cell_id"] not in dup_ids]
    return {
        "format": HIST_FORMAT,
        "changes": changes,
        "mapping_errors": mapping_errors,
        "non_numeric": non_numeric,
        "unchanged": unchanged,
    }


# ─────────────────────────── projections ───────────────────────────


def _get_path(d: Any, path: str) -> Any:
    cur = d
    for seg in path.split("."):
        if not isinstance(cur, dict):
            return None
        cur = cur.get(seg)
    return cur


def projection_values(
    engine_outputs: Mapping[str, Any],
) -> tuple[int, dict[str, float | None]]:
    """``(year_count, {proj_cell_key: value})`` from ``{engine: outputs-dict}``."""
    years_by_engine: dict[str, list[Any]] = {}
    for eng in {e for e, _p, _l in PROJ_LINES}:
        out = engine_outputs.get(eng)
        yrs = out.get("years") if isinstance(out, dict) else None
        years_by_engine[eng] = yrs if isinstance(yrs, list) else []
    n = len(years_by_engine.get("revenue") or []) or max(
        (len(v) for v in years_by_engine.values()), default=0
    )
    values: dict[str, float | None] = {}
    for eng, path, _label in PROJ_LINES:
        yrs = years_by_engine.get(eng) or []
        for i in range(n):
            y = yrs[i] if i < len(yrs) else None
            values[proj_cell_key(eng, path, i)] = to_number(_get_path(y, path))
    return n, values


# ───────────────────── Future P&L parity (Task 15) ─────────────────────
#
# What the retired browser builder (``ProjectionsSection.onExport``) carried
# beyond the engine lines: the room-count rows, the per-year % Rev / PAR / POR
# columns, Net Cash Flow and the Exit Year NOI. Every figure is read from an
# engine output; a figure the run did not publish is a reason, never a 0.

#: The browser builder's Exit Year row label (``FORWARD_NOI_LABEL`` in the
#: retired web code). Used when the returns engine capped the forward 12 months;
#: any other basis uses :data:`EXIT_NOI_LABEL`. The engine's own period label
#: (``exit_noi_period_label``) is appended either way.
FORWARD_NOI_LABEL = "Forward 12-Month Cash NOI (after FF&E reserve)"
EXIT_NOI_LABEL = "Exit NOI (after FF&E reserve)"
EXIT_COLUMN_LABEL = "Exit Year"
NET_CASH_FLOW_LABEL = "Net Cash Flow"
#: Shown only when the cash flow's NOI differs from the expense engine's in at
#: least one year — the deal carries ``noi_override_by_year`` (an analyst pin),
#: which the returns / cash-flow engines run on and the expense lines do not.
CASH_FLOW_NOI_PIN_LABEL = "NOI used by the cash flow (analyst pin)"
#: Comment on a Net Cash Flow cell built on that pinned NOI.
CASH_FLOW_PIN_NOTE = (
    "This year's Net Cash Flow is built on the analyst-pinned NOI "
    "(noi_override_by_year), not the NOI lines above: see the row "
    f"'{CASH_FLOW_NOI_PIN_LABEL}'. Pinned NOI before the FF&E reserve less FF&E "
    "Reserve = Net Cash Flow."
)
#: The cash flow and the expense engine agree to the cent on an unpinned year.
_PIN_TOLERANCE = 0.01

#: The room-count rows, in the browser builder's order: (row id, label).
OPERATING_ROWS: tuple[tuple[str, str], ...] = (
    ("days", "Days"),
    ("keys", "Number of Rooms"),
    ("available_rooms", "Available Rooms"),
    ("occupied_rooms", "Occupied Rooms"),
)

#: ``PROJ_LINES`` paths that are rates, not dollars — % Rev / PAR / POR do not
#: apply to them (a dash, not a number).
NON_RATIO_PATHS: frozenset[str] = frozenset(
    {"revenue.occupancy", "revenue.adr", "revenue.revpar"}
)

#: % Rev denominator for each departmental expense: that department's own
#: revenue (USALI — the same basis the Benchmark % column states). Every other
#: dollar line is over Total Revenue (``expense.total_revenue``).
TOTAL_REVENUE_PATH = "expense.total_revenue"
PCT_REV_BASIS: dict[str, str] = {
    "expense.dept_expenses.rooms": "revenue.rooms_revenue",
    "expense.dept_expenses.food_beverage": "fb.fb_revenue",
    "expense.dept_expenses.other_operated": "fb.other_revenue",
}

#: The cash-flow engine's unlevered rows whose sum is the year's Net Cash Flow
#: (NOI before the reserve, then the reserve as a negative outflow).
_CF_NOI_ROW = "NOI (before FF&E reserve)"
_CF_FFE_ROW = "FF&E Reserve"


@dataclass(frozen=True)
class Sourced:
    """An engine-read figure, or why there is none (``value`` is then None)."""

    value: float | None
    reason: str | None = None
    code: ReasonCode = ReasonCode.NO_SOURCE


@dataclass(frozen=True)
class ProjectionExtras:
    """The parity rows' inputs for ``year_count`` model years."""

    engines: frozenset[str]
    operating: Mapping[str, tuple[Sourced, ...]]
    net_cash_flow: tuple[Sourced, ...]
    #: The cash flow's own NOI before the FF&E reserve, per year, where it
    #: differs from the expense engine's (an analyst pin); None elsewhere.
    cash_flow_noi_pin: tuple[float | None, ...]
    exit_noi: Sourced
    exit_noi_label: str
    exit_calendar_year: int | None


def _engine_missing(engine: str) -> Sourced:
    return Sourced(None, f"the {engine} engine has no result on this run",
                   ReasonCode.ENGINE_SKIPPED)


def _trace_inputs(trace: Any) -> dict[str, float | None]:
    ins = trace.get("inputs") if isinstance(trace, Mapping) else None
    out: dict[str, float | None] = {}
    for x in ins if isinstance(ins, list) else ():
        if isinstance(x, Mapping) and isinstance(x.get("name"), str):
            out[x["name"]] = to_number(x.get("value"))
    return out


def projection_extras(
    engine_outputs: Mapping[str, Any], year_count: int
) -> ProjectionExtras:
    """Read the parity rows from ``{engine: outputs-dict}``. Nothing is derived
    here from anything but engine outputs.

    * Days, Number of Rooms and Occupied Rooms — the revenue engine's
      ``years[i].occupied_rooms`` trace (its value and its ``keys`` / ``days``
      inputs); Available Rooms — the ``rooms_available`` input of its
      ``years[i].revpar`` trace. The engine's day count is the one its
      Available Rooms and RevPAR are built on, so the rows foot.
    * Net Cash Flow — the cash-flow engine's unlevered ``NOI (before FF&E
      reserve)`` + ``FF&E Reserve`` rows for operating year ``i + 1``.
    * Exit NOI — the returns engine's ``terminal_noi`` and its basis label.
    """
    engines = frozenset(k for k, v in engine_outputs.items() if isinstance(v, Mapping))
    rev = engine_outputs.get("revenue") if "revenue" in engines else None
    prov = rev.get("provenance") if isinstance(rev, Mapping) else None

    operating: dict[str, list[Sourced]] = {rid: [] for rid, _l in OPERATING_ROWS}
    for i in range(year_count):
        if rev is None:
            for rid in operating:
                operating[rid].append(_engine_missing("revenue"))
            continue
        occ_trace = prov.get(f"years[{i}].occupied_rooms") if isinstance(prov, Mapping) else None
        par_trace = prov.get(f"years[{i}].revpar") if isinstance(prov, Mapping) else None
        occ_in = _trace_inputs(occ_trace)
        no_occ = "the revenue run published no occupied-rooms trace for this year"
        operating["days"].append(_sourced(occ_in.get("days"), no_occ))
        operating["keys"].append(_sourced(occ_in.get("keys"), no_occ))
        operating["occupied_rooms"].append(_sourced(
            to_number(occ_trace.get("value")) if isinstance(occ_trace, Mapping) else None,
            no_occ,
        ))
        operating["available_rooms"].append(_sourced(
            _trace_inputs(par_trace).get("rooms_available"),
            "the revenue run published no available-rooms figure (RevPAR trace) for this year",
        ))

    ncf: list[Sourced] = []
    pins: list[float | None] = []
    exp = engine_outputs.get("expense") if "expense" in engines else None
    exp_years = exp.get("years") if isinstance(exp, Mapping) else None
    cf = engine_outputs.get("cash_flow") if "cash_flow" in engines else None
    lines: dict[str, list[Any]] = {}
    if isinstance(cf, Mapping):
        for line in cf.get("unlevered") or ():
            if isinstance(line, Mapping) and isinstance(line.get("values"), list):
                lines.setdefault(str(line.get("label")), line["values"])
    for i in range(year_count):
        pins.append(None)
        if cf is None:
            ncf.append(_engine_missing("cash-flow"))
            continue
        noi_v = _at(lines.get(_CF_NOI_ROW), i + 1)
        ffe_v = _at(lines.get(_CF_FFE_ROW), i + 1)
        if noi_v is None or ffe_v is None:
            missing = _CF_NOI_ROW if noi_v is None else _CF_FFE_ROW
            ncf.append(Sourced(
                None, f"the cash-flow statement has no {missing} for this year"
            ))
        else:
            ncf.append(Sourced(noi_v + ffe_v))
            # Pinned year: the cash flow's after-reserve NOI is not the
            # expense engine's. Without an expense NOI there is nothing on
            # the sheet for it to contradict, so nothing is flagged.
            ey = exp_years[i] if isinstance(exp_years, list) and i < len(exp_years) else None
            exp_noi = to_number(ey.get("noi")) if isinstance(ey, Mapping) else None
            if exp_noi is not None and abs((noi_v + ffe_v) - exp_noi) > _PIN_TOLERANCE:
                pins[i] = noi_v

    ret = engine_outputs.get("returns") if "returns" in engines else None
    basis = ret.get("exit_noi_basis") if isinstance(ret, Mapping) else None
    period = ret.get("exit_noi_period_label") if isinstance(ret, Mapping) else None
    label = FORWARD_NOI_LABEL if basis in (None, "forward_12m") else EXIT_NOI_LABEL
    if isinstance(period, str) and period.strip():
        label = f"{label} — {period.strip()}"
    if ret is None:
        exit_noi = Sourced(None, "the returns engine has no result on this run",
                           ReasonCode.ENGINE_SKIPPED)
    else:
        exit_noi = _sourced(
            to_number(ret.get("terminal_noi")),
            "the returns engine published no terminal NOI on this run",
        )

    cal = rev.get("projection_calendar_years") if isinstance(rev, Mapping) else None
    last = cal[year_count - 1] if isinstance(cal, list) and 0 < year_count <= len(cal) else None
    return ProjectionExtras(
        engines=engines,
        operating={k: tuple(v) for k, v in operating.items()},
        net_cash_flow=tuple(ncf),
        cash_flow_noi_pin=tuple(pins),
        exit_noi=exit_noi,
        exit_noi_label=label,
        exit_calendar_year=last + 1 if isinstance(last, int) else None,
    )


def _sourced(value: float | None, reason: str) -> Sourced:
    return Sourced(value) if value is not None else Sourced(None, reason)


def _at(values: list[Any] | None, idx: int) -> float | None:
    if not isinstance(values, list) or not 0 <= idx < len(values):
        return None
    return to_number(values[idx])


def _proj_line_label(eng: str, path: str) -> str:
    return next((lbl for e, p, lbl in PROJ_LINES if e == eng and p == path), path)


def _ratio(
    amount: float | None, amount_label: str, denom: float | None, denom_label: str,
) -> float | str:
    """``amount ÷ denom`` or the refusal sentence naming the missing input."""
    if amount is None:
        return unavailable(ReasonCode.NO_SOURCE, f"no {amount_label} for this year")
    if denom is None:
        return unavailable(ReasonCode.NO_SOURCE, f"no {denom_label} for this year")
    if denom == 0:
        return unavailable(ReasonCode.NOT_APPLICABLE, f"{denom_label} is zero for this year")
    if denom < 0:
        return unavailable(ReasonCode.NOT_APPLICABLE, f"{denom_label} is negative for this year")
    return amount / denom


def _not_applicable(what: str) -> str:
    return refuse(ReasonCode.NOT_APPLICABLE, what)


def _write(ws: Any, row: int, col: int, value: float | str | None, fmt: str) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.fill = _READONLY_FILL
    if isinstance(value, (int, float)):
        c.number_format = fmt


_AMOUNT_FMT = "#,##0.####"
_COUNT_FMT = "#,##0"
_PCT_FMT = "0.0%"
_PER_ROOM_FMT = "#,##0.00"


def build_projections_workbook(
    *,
    deal_id: str,
    deal_name: str | None,
    year_headers: Sequence[str],
    values: Mapping[str, float | None],
    assumptions: Sequence[Mapping[str, Any]],
    benchmark: Mapping[str, float] | None = None,
    extras: ProjectionExtras | None = None,
) -> bytes:
    """Render the Future P&L + Assumptions workbook.

    ``assumptions`` rows carry ``key, label, value, unit, source, method,
    note_required, current_note`` — already resolved by the caller.

    FON-97 — ``benchmark`` (``benchmark_column_values``) adds a read-only
    "Benchmark %" column AFTER the year columns: the CBRE / HotStats comp-set
    ratio for each mapped line, as a fraction on the same basis the web's
    "% Rev" cell uses. ``None`` (no benchmark on the deal) = no column at all.

    Task 15 — ``extras`` (:func:`projection_extras`) feeds the parity rows the
    retired browser builder carried: Days / Number of Rooms / Available Rooms /
    Occupied Rooms first, Net Cash Flow and the Exit Year NOI last, plus a
    % Rev / PAR / POR column beside every year's Amount. ``None`` = no engine
    output was supplied, so every one of those cells states it is not
    available. All of them are display-only and ignored on import.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = PROJ_SHEET
    n = len(year_headers)
    if extras is None:
        extras = projection_extras({}, n)

    def amount_col(i: int) -> int:
        return 3 + 4 * i

    exit_col = amount_col(n) if n else None
    bm_col = (
        (exit_col + 1 if exit_col is not None else 3) if benchmark is not None else None
    )
    last_col = bm_col or exit_col or 2

    ws.cell(row=1, column=1, value=PATH_HEADER)
    ws.cell(row=1, column=2, value="Line item (computed — read-only)")
    ws.cell(row=2, column=1, value=COLUMN_KEY_HEADER)
    for i, h in enumerate(year_headers):
        for off, (key, suffix) in enumerate(
            (("amount", "Amount"), ("pct_rev", "% Rev"), ("par", "PAR"), ("por", "POR"))
        ):
            ws.cell(row=1, column=amount_col(i) + off, value=f"{h} {suffix}")
            ws.cell(row=2, column=amount_col(i) + off, value=f"{key}:{i}")
    if exit_col is not None:
        cy = extras.exit_calendar_year
        ws.cell(row=1, column=exit_col,
                value=f"{EXIT_COLUMN_LABEL} · {cy}" if cy is not None else EXIT_COLUMN_LABEL)
        ws.cell(row=2, column=exit_col, value="exit")
    if bm_col is not None:
        ws.cell(row=1, column=bm_col, value=BENCHMARK_HEADER)
        ws.cell(row=2, column=bm_col, value="benchmark")
    _style_header(ws, 1, last_col)
    ws.row_dimensions[2].hidden = True

    proj_lines = [
        (eng, path, label)
        for eng, path, label in PROJ_LINES
        if f"{eng}.{path}" not in OPTIONAL_PROJ_PATHS
        or any(values.get(proj_cell_key(eng, path, i)) is not None for i in range(n))
    ]
    available = extras.operating.get("available_rooms", ())
    occupied = extras.operating.get("occupied_rooms", ())

    def _op(seq: Sequence[Sourced], i: int) -> float | None:
        return seq[i].value if i < len(seq) else None

    def ratio_cells(r: int, i: int, amount: float | None, label: str, pct_basis: str) -> None:
        eng, sub = pct_basis.split(".", 1)
        denom = values.get(proj_cell_key(eng, sub, i))
        c = amount_col(i)
        _write(ws, r, c + 1, _ratio(amount, label, denom, _proj_line_label(eng, sub)), _PCT_FMT)
        _write(ws, r, c + 2, _ratio(amount, label, _op(available, i), "Available Rooms"),
               _PER_ROOM_FMT)
        _write(ws, r, c + 3, _ratio(amount, label, _op(occupied, i), "Occupied Rooms"),
               _PER_ROOM_FMT)

    def dash_ratios(r: int, i: int, why: str) -> None:
        for off in (1, 2, 3):
            _write(ws, r, amount_col(i) + off, _not_applicable(why), _AMOUNT_FMT)

    def exit_dash(r: int) -> None:
        if exit_col is not None:
            _write(ws, r, exit_col, _not_applicable("the Exit Year column carries the exit NOI only"),
                   _AMOUNT_FMT)
        if bm_col is not None and ws.cell(row=r, column=1).value.startswith(READONLY_ROW_PREFIX):
            _write(ws, r, bm_col,
                   _not_applicable("the benchmark covers P&L expense lines only"), _AMOUNT_FMT)

    r = 3
    for rid, label in OPERATING_ROWS:
        ws.cell(row=r, column=1, value=f"{READONLY_ROW_PREFIX}{rid}")
        ws.cell(row=r, column=2, value=label)
        seq = extras.operating.get(rid, ())
        for i in range(n):
            src = seq[i] if i < len(seq) else Sourced(None, "not in the run")
            _write(ws, r, amount_col(i),
                   src.value if src.value is not None
                   else unavailable(src.code, src.reason or "not in the run", rid),
                   _COUNT_FMT)
            dash_ratios(r, i, f"{label} is a count, not an amount")
        exit_dash(r)
        r += 1

    for eng, path, label in proj_lines:
        ws.cell(row=r, column=1, value=f"{eng}.{path}")
        ws.cell(row=r, column=2, value=label)
        dotted = f"{eng}.{path}"
        for i in range(n):
            v = values.get(proj_cell_key(eng, path, i))
            _write(ws, r, amount_col(i),
                   v if v is not None else unavailable(
                       ReasonCode.NO_SOURCE,
                       f"the {eng} engine published no {label} for this year", dotted),
                   _AMOUNT_FMT)
            if dotted in NON_RATIO_PATHS:
                dash_ratios(r, i, f"{label} is a rate, not an amount")
            else:
                ratio_cells(r, i, v, label, PCT_REV_BASIS.get(dotted, TOTAL_REVENUE_PATH))
        exit_dash(r)
        if bm_col is not None and benchmark is not None:
            bv = benchmark.get(dotted)
            bc = ws.cell(row=r, column=bm_col, value=bv if bv is not None else refuse(
                ReasonCode.NO_SOURCE, f"the benchmark publishes no ratio for {label}", dotted))
            bc.fill = _READONLY_FILL
            if bv is not None:
                bc.number_format = "0.0%"
        r += 1

    pins = extras.cash_flow_noi_pin
    pinned = [i < len(pins) and pins[i] is not None for i in range(n)]
    if any(pinned):
        ws.cell(row=r, column=1, value=f"{READONLY_ROW_PREFIX}cash_flow_noi_pin")
        ws.cell(row=r, column=2, value=CASH_FLOW_NOI_PIN_LABEL)
        for i in range(n):
            if pinned[i]:
                _write(ws, r, amount_col(i), pins[i], _AMOUNT_FMT)
                ratio_cells(r, i, pins[i], CASH_FLOW_NOI_PIN_LABEL, TOTAL_REVENUE_PATH)
            else:
                _write(ws, r, amount_col(i), _not_applicable(
                    "no pin this year: the cash flow uses the NOI above"), _AMOUNT_FMT)
                dash_ratios(r, i, "no pin this year")
        exit_dash(r)
        r += 1

    ws.cell(row=r, column=1, value=f"{READONLY_ROW_PREFIX}net_cash_flow")
    ws.cell(row=r, column=2, value=NET_CASH_FLOW_LABEL)
    for i in range(n):
        src = extras.net_cash_flow[i] if i < len(extras.net_cash_flow) else Sourced(
            None, "the cash-flow statement has no figure for this year")
        amt = src.value
        _write(ws, r, amount_col(i),
               amt if amt is not None
               else unavailable(src.code, src.reason or "not in the run", "net_cash_flow"),
               _AMOUNT_FMT)
        if pinned[i]:
            ws.cell(row=r, column=amount_col(i)).comment = Comment(CASH_FLOW_PIN_NOTE, "Fondok")
        ratio_cells(r, i, amt, NET_CASH_FLOW_LABEL, TOTAL_REVENUE_PATH)
    exit_dash(r)
    r += 1

    ws.cell(row=r, column=1, value=f"{READONLY_ROW_PREFIX}exit_noi")
    ws.cell(row=r, column=2, value=extras.exit_noi_label)
    for i in range(n):
        _write(ws, r, amount_col(i),
               _not_applicable("the exit NOI is not a modelled year's figure"), _AMOUNT_FMT)
        dash_ratios(r, i, "the exit NOI is not a modelled year's figure")
    if exit_col is not None:
        en = extras.exit_noi
        _write(ws, r, exit_col,
               en.value if en.value is not None
               else unavailable(en.code, en.reason or "not in the run", "terminal_noi"),
               _AMOUNT_FMT)
    if bm_col is not None:
        _write(ws, r, bm_col,
               _not_applicable("the benchmark covers P&L expense lines only"), _AMOUNT_FMT)
    r += 1

    if not year_headers:
        ws.cell(row=r + 1, column=2,
                value="No model run yet — run the model to populate the projection.")
    ws.column_dimensions["A"].hidden = True
    ws.column_dimensions["B"].width = 44
    for i in range(n):
        ws.column_dimensions[get_column_letter(amount_col(i))].width = 16
        for off in (1, 2, 3):
            ws.column_dimensions[get_column_letter(amount_col(i) + off)].width = 11
    if exit_col is not None:
        ws.column_dimensions[get_column_letter(exit_col)].width = 16
    if bm_col is not None:
        ws.column_dimensions[get_column_letter(bm_col)].width = 14
    ws.freeze_panes = "C3"

    wa = wb.create_sheet(ASSUMPTION_SHEET)
    for c, h in enumerate(ASSUMPTION_COLUMNS, start=1):
        wa.cell(row=1, column=c, value=h)
    _style_header(wa, 1, len(ASSUMPTION_COLUMNS))
    for r_i, a in enumerate(assumptions):
        r = 2 + r_i
        wa.cell(row=r, column=1, value=a["key"])
        wa.cell(row=r, column=2, value=a.get("label"))
        wa.cell(row=r, column=3, value=a.get("value"))
        wa.cell(row=r, column=4, value=a.get("unit"))
        wa.cell(row=r, column=5, value=a.get("source") or "—")
        wa.cell(row=r, column=6, value=a.get("method"))
        wa.cell(row=r, column=7, value="Yes" if a.get("note_required") else "No")
        wa.cell(row=r, column=8, value=a.get("current_note") or "")
        for c in (1, 2, 4, 5, 6, 7, 8):
            wa.cell(row=r, column=c).fill = _READONLY_FILL
    for c, w in zip(range(1, 10), (26, 36, 14, 26, 22, 52, 12, 30, 40), strict=True):
        wa.column_dimensions[get_column_letter(c)].width = w
    wa.freeze_panes = "C2"
    _readme(
        wb,
        fmt=PROJ_FORMAT,
        deal_id=deal_id,
        deal_name=deal_name,
        lines=(
            "Future P&L — the engine projection by model year, plus every analyst-editable "
            "assumption that drives it.",
            "Only the Assumptions sheet is importable: edit the Value column, add a Note for every "
            "row whose 'Note required' is Yes, save as .xlsx, then use Import on the Future P&L.",
            "Key is the stable field_overrides key — do not edit it.",
            "Values use the units in the Unit column (fractions, not percents: 0.045 = 4.5%).",
            "Future P&L cells are computed by the engines; a change there is reported as "
            "'computed — edit the assumption instead' and is never applied.",
            "Each model year is four columns: Amount, % Rev (share of Total Revenue; a "
            "departmental expense is over its own department's revenue), PAR (per available "
            "room) and POR (per occupied room) — all computed from the row's Amount and that "
            "year's revenue / room counts. Days, Number of Rooms, Available Rooms and Occupied "
            "Rooms come from the revenue engine; Net Cash Flow (NOI after the FF&E reserve) from "
            "the cash-flow engine; the Exit Year column holds only the NOI the reversion is "
            "valued on, from the returns engine, with its basis in the row label.",
            "Analyst NOI pin: when the deal pins NOI by year (noi_override_by_year), the "
            "returns and cash-flow engines run on the pinned NOI while the P&L lines above "
            "show the operating projection. In any year where they differ, the row "
            f"'{CASH_FLOW_NOI_PIN_LABEL}' shows the cash flow's NOI before the FF&E reserve, "
            "and that year's Net Cash Flow cell carries a comment: pinned NOI less FF&E Reserve "
            "= Net Cash Flow. The row is read-only and ignored on import.",
            "A cell reading 'not available: <reason>' is a figure the model run did not "
            "produce; '—' means the column does not apply to that row. Neither is ever a 0. "
            "The % Rev / PAR / POR / Exit Year columns and those rows are read-only and "
            "ignored on import.",
            "Benchmark % (present only when the deal carries a CBRE / HotStats P&L benchmark) "
            "is the comp-set ratio for the line on the same basis as the model's % Rev — "
            "departmental expense over that department's revenue, everything else over total "
            "revenue. Read-only and ignored on import.",
            "Import shows a preview of changed assumptions (old → new) and any rejected rows. "
            "Nothing changes until you Apply; applied values are saved as analyst overrides, "
            "with your note, through the same save path as the Assumptions panel.",
        ),
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _year_amount_cells(wp: Any) -> Iterator[tuple[int, str, Any, int, int, Any]]:
    """``(row, engine path, label, column, year index, period label)`` for every
    engine-line Amount cell of an uploaded Future P&L sheet.

    Current layout: the hidden row 2 (``fondok_column``) names the columns —
    only ``amount:<i>`` columns are model-year values; % Rev / PAR / POR, the
    Exit Year and the Benchmark % columns are never read, and neither is a
    ``fondok_readonly:`` row. A sheet without that row is the earlier
    positional layout (one column per year from C, Benchmark % after them).
    """
    if wp.cell(row=2, column=1).value == COLUMN_KEY_HEADER:
        cols: list[tuple[int, int, Any]] = []
        for c in range(3, wp.max_column + 1):
            key = wp.cell(row=2, column=c).value
            if isinstance(key, str) and key.startswith("amount:"):
                try:
                    idx = int(key.split(":", 1)[1])
                except ValueError:
                    continue
                head = wp.cell(row=1, column=c).value
                if isinstance(head, str) and head.endswith(" Amount"):
                    head = head[: -len(" Amount")]
                cols.append((c, idx, head))
        first_row = 3
    else:
        cols = []
        for c in range(3, wp.max_column + 1):
            head = wp.cell(row=1, column=c).value
            # FON-97 — the Benchmark % column is not a year: never diffed,
            # never reported, whatever was typed into it.
            if isinstance(head, str) and head.strip().lower().startswith("benchmark"):
                continue
            cols.append((c, c - 3, head))
        first_row = 2
    for r in range(first_row, wp.max_row + 1):
        path = wp.cell(row=r, column=1).value
        if (
            not isinstance(path, str)
            or path.startswith(READONLY_ROW_PREFIX)
            or "." not in path
        ):
            continue
        label = wp.cell(row=r, column=2).value
        for c, idx, head in cols:
            yield r, path, label, c, idx, head


def parse_projections_import(
    data: bytes,
    *,
    deal_id: str,
    current_values: Mapping[str, float | None],
    current_assumptions: Mapping[str, Any],
    needs_note: Callable[[str], bool],
) -> dict[str, Any]:
    """Diff an uploaded Future P&L workbook. Never writes.

    Returns ``{changes, rejected, mapping_errors, non_numeric, computed_edits,
    unchanged}``; only ``changes`` are eligible for Apply.
    """
    wb = _load(data)
    _check_readme(wb, fmt=PROJ_FORMAT, deal_id=deal_id)
    if ASSUMPTION_SHEET not in wb.sheetnames:
        raise WorkbookFormatError(f"Sheet {ASSUMPTION_SHEET!r} is missing.")

    changes: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    mapping_errors: list[dict[str, Any]] = []
    non_numeric: list[dict[str, Any]] = []
    computed_edits: list[dict[str, Any]] = []
    unchanged = 0

    wa = wb[ASSUMPTION_SHEET]
    header = [wa.cell(row=1, column=c).value for c in range(1, wa.max_column + 1)]

    def _col(name: str, fallback: int) -> int:
        for i, h in enumerate(header, start=1):
            if isinstance(h, str) and h.strip().lower().startswith(name.lower()):
                return i
        return fallback

    c_key, c_val, c_note = _col("Key", 1), _col("Value", 3), _col("Note (", 9)
    seen: set[str] = set()
    for r in range(2, wa.max_row + 1):
        raw_key = wa.cell(row=r, column=c_key).value
        key = str(raw_key).strip() if raw_key is not None else ""
        raw_val = wa.cell(row=r, column=c_val).value
        raw_note = wa.cell(row=r, column=c_note).value
        note = str(raw_note).strip() if raw_note is not None else ""
        ref = f"{get_column_letter(c_val)}{r}"
        if not key:
            if raw_val not in (None, ""):
                mapping_errors.append({"cell_ref": ref, "key": None, "reason": "missing_key",
                                       "detail": "A value with no Key cannot be mapped."})
            continue
        spec = assumption_spec(key)
        if spec is None:
            mapping_errors.append({
                "cell_ref": ref, "key": key, "reason": "unknown_key",
                "detail": "Not an editable Future P&L assumption key.",
            })
            continue
        if key in seen:
            mapping_errors.append({"cell_ref": ref, "key": key, "reason": "duplicate_key",
                                   "detail": "The key appears more than once."})
            continue
        seen.add(key)
        if spec.choices:
            new_txt = str(raw_val).strip().lower() if raw_val not in (None, "") else None
            old_txt = current_assumptions.get(key)
            old_txt = str(old_txt).strip().lower() if old_txt not in (None, "") else None
            if new_txt is not None and new_txt not in spec.choices:
                non_numeric.append({"cell_ref": ref, "key": key, "label": spec.label,
                                    "raw": str(raw_val),
                                    "detail": "Must be one of: " + ", ".join(spec.choices) + "."})
                continue
            if new_txt is None and old_txt is not None:
                rejected.append({
                    "cell_ref": ref, "key": key, "label": spec.label, "old_value": old_txt,
                    "new_value": None, "reason": "cleared",
                    "detail": "Clearing is not imported — use Reset on the Future P&L.",
                })
                continue
            if new_txt == old_txt:
                unchanged += 1
                continue
            required = needs_note(key)
            row = {"cell_ref": ref, "key": key, "label": spec.label, "old_value": old_txt,
                   "new_value": new_txt, "note": note, "note_required": required}
            if required and not note:
                rejected.append({**row, "reason": "note_required",
                                 "detail": "This assumption moves the model — add a Note."})
                continue
            changes.append(row)
            continue
        new_val, ok = parse_cell_number(raw_val)
        if ok and new_val is not None and spec.integer and not float(new_val).is_integer():
            ok = False
        if not ok:
            non_numeric.append({"cell_ref": ref, "key": key, "label": spec.label,
                                "raw": str(raw_val),
                                "detail": "Must be a whole number." if spec.integer else None})
            continue
        old_val = to_number(current_assumptions.get(key))
        if new_val is None:
            if old_val is None:
                unchanged += 1
            else:
                rejected.append({
                    "cell_ref": ref, "key": key, "label": spec.label, "old_value": old_val,
                    "new_value": None, "reason": "cleared",
                    "detail": "Clearing is not imported — use Reset on the Future P&L.",
                })
            continue
        if spec.integer:
            new_val = float(int(new_val))
        if same_number(old_val, new_val):
            unchanged += 1
            continue
        required = needs_note(key)
        row = {"cell_ref": ref, "key": key, "label": spec.label, "old_value": old_val,
               "new_value": int(new_val) if spec.integer else new_val,
               "note": note, "note_required": required}
        if required and not note:
            rejected.append({**row, "reason": "note_required",
                             "detail": "This assumption moves the model — add a Note."})
            continue
        changes.append(row)

    if PROJ_SHEET in wb.sheetnames:
        wp = wb[PROJ_SHEET]
        for r, path, label, col, year_index, period in _year_amount_cells(wp):
            eng, sub = path.split(".", 1)
            ck = proj_cell_key(eng, sub, year_index)
            if ck not in current_values:
                continue
            raw = wp.cell(row=r, column=col).value
            old_val = current_values.get(ck)
            # The export wrote "not available: <reason>" where the engine
            # published nothing; left as written, it is no edit.
            if old_val is None and isinstance(raw, str) and raw.strip().startswith(
                NOT_AVAILABLE_PREFIX
            ):
                continue
            new_val, ok = parse_cell_number(raw)
            if ok and same_number(old_val, new_val):
                continue
            computed_edits.append({
                "cell_ref": f"{get_column_letter(col)}{r}",
                "cell_key": ck, "line_label": label, "year_index": year_index,
                "period_label": period, "old_value": old_val,
                "new_value": new_val if ok else str(raw),
                "reason": "computed",
                "detail": "Computed by the engines — edit the assumption instead.",
            })

    return {
        "format": PROJ_FORMAT,
        "changes": changes,
        "rejected": rejected,
        "mapping_errors": mapping_errors,
        "non_numeric": non_numeric,
        "computed_edits": computed_edits,
        "unchanged": unchanged,
    }


__all__ = [
    "ASSUMPTIONS",
    "ASSUMPTION_BY_KEY",
    "BENCHMARK_HEADER",
    "BENCHMARK_LINE_PATHS",
    "COLUMN_KEY_HEADER",
    "EXIT_COLUMN_LABEL",
    "EXIT_NOI_LABEL",
    "FORWARD_NOI_LABEL",
    "NET_CASH_FLOW_LABEL",
    "CASH_FLOW_NOI_PIN_LABEL",
    "OPERATING_ROWS",
    "READONLY_ROW_PREFIX",
    "ProjectionExtras",
    "Sourced",
    "projection_extras",


    "BY_YEAR_TARGET_LISTS",
    "SCHEDULE_FIELDS",
    "SCHEDULE_KEY",
    "schedule_assumptions",
    "schedule_period_key",
    "split_schedule_period_key",
    "assumption_spec",
    "by_year_target_assumptions",
    "by_year_target_key",
    "split_by_year_target_key",
    "HIST_FORMAT",
    "HIST_LINES",
    "PROJ_FORMAT",
    "OPTIONAL_PROJ_PATHS",
    "PROJ_LINES",
    "HistCell",
    "HistColumn",
    "HistCurrent",
    "WorkbookFormatError",
    "benchmark_column_values",
    "build_historicals_workbook",
    "build_projections_workbook",
    "hist_cell_id",
    "parse_historicals_import",
    "parse_projections_import",
    "proj_cell_key",
    "projection_values",
    "split_hist_cell_id",
    "to_number",
]
