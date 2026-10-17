"""Run-to-run comparison of persisted engine outputs (regression baselines).

Sam (2026-10-16): "preserve the current Kimpton model outputs so we can
distinguish expected changes from unintended calculation discrepancies during
regression testing."

``engine_outputs`` is append-only — every run INSERTs one row per engine and
only deal deletion removes them — so every earlier run on a live deal is
already a preserved baseline. This module:

* lists a deal's runs (one entry per ``run_id``),
* loads one run's outputs keyed by engine,
* flattens the headline outputs Sam checks into a stable, ordered field list
  (``extract_headline_fields``), and
* diffs two such lists field by field (``diff_fields``).

Read-only by construction: nothing here computes, alters or persists an
engine value. Every SQL statement carries the tenant predicate.

The same extractor/differ is used by ``scripts/compare_golden_baseline.py``
against the committed golden baseline, so the live comparison and the repo
comparison can never disagree about what a "headline field" is.
"""

from __future__ import annotations

import io
import json
import math
from collections.abc import Iterable, Mapping
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

#: Two values within this tolerance are "unchanged". Engine outputs are
#: deterministic, so anything beyond float-formatting noise is a real move.
REL_TOL = 1e-9
ABS_TOL = 1e-6

#: Display order of the engine groups in the diff.
ENGINE_ORDER: tuple[str, ...] = (
    "revenue",
    "expense",
    "capital",
    "debt",
    "returns",
    "partnership",
)

# ─────────────────────────── field extraction ───────────────────────────


def _get_path(obj: Any, path: str) -> Any:
    """Resolve a provenance-style path (``years[0].dept_expenses.total``)."""
    cur = obj
    for part in path.split("."):
        if cur is None:
            return None
        name, _, rest = part.partition("[")
        if name:
            if not isinstance(cur, Mapping):
                return None
            cur = cur.get(name)
        while rest:
            idx_s, _, rest = rest.partition("]")
            rest = rest[1:] if rest.startswith("[") else rest
            if not isinstance(cur, list):
                return None
            try:
                idx = int(idx_s)
            except ValueError:
                return None
            if idx >= len(cur):
                return None
            cur = cur[idx]
    return cur


def _provenance_for(outputs: Mapping[str, Any], prov_key: str | None) -> dict[str, Any] | None:
    """Compact the engine's per-field provenance entry, when it carries one."""
    if not prov_key:
        return None
    prov = outputs.get("provenance")
    if not isinstance(prov, Mapping):
        return None
    entry = prov.get(prov_key)
    if not isinstance(entry, Mapping):
        return None
    inputs = []
    for inp in entry.get("inputs") or []:
        if not isinstance(inp, Mapping):
            continue
        inputs.append(
            {
                "name": inp.get("name"),
                "value": inp.get("value"),
                "source": inp.get("source"),
                "assumption_key": inp.get("assumption_key"),
            }
        )
    return {
        "state": entry.get("state"),
        "source": entry.get("source"),
        "formula": entry.get("formula"),
        "note": entry.get("note"),
        "reason": entry.get("reason"),
        "assumption_key": entry.get("assumption_key"),
        "inputs": inputs,
    }


# (path suffix inside a year row, label, unit)
_REVENUE_YEAR_FIELDS: tuple[tuple[str, str, str], ...] = (
    ("occupancy", "Occupancy", "pct"),
    ("adr", "ADR", "usd"),
    ("revpar", "RevPAR", "usd"),
    ("rooms_revenue", "Rooms revenue", "usd"),
    ("total_revenue", "Total revenue", "usd"),
)

_EXPENSE_YEAR_FIELDS: tuple[tuple[str, str, str], ...] = (
    ("total_revenue", "Total revenue", "usd"),
    ("dept_expenses.rooms", "Dept. expenses · Rooms", "usd"),
    ("dept_expenses.food_beverage", "Dept. expenses · F&B", "usd"),
    ("dept_expenses.other_operated", "Dept. expenses · Other operated", "usd"),
    ("dept_expenses.total", "Departmental expenses (total)", "usd"),
    ("undistributed.administrative_general", "Undistributed · A&G", "usd"),
    ("undistributed.sales_marketing", "Undistributed · Sales & marketing", "usd"),
    ("undistributed.property_operations", "Undistributed · Property ops & maintenance", "usd"),
    ("undistributed.utilities", "Undistributed · Utilities", "usd"),
    ("undistributed.information_telecom", "Undistributed · Info & telecom", "usd"),
    ("undistributed.total", "Undistributed expenses (total)", "usd"),
    ("gop", "GOP", "usd"),
    ("mgmt_fee", "Management fee", "usd"),
    ("incentive_mgmt_fee", "Incentive management fee", "usd"),
    ("fixed_charges.total", "Fixed charges (total)", "usd"),
    ("noi_institutional", "NOI (before FF&E reserve)", "usd"),
    ("ffe_reserve", "FF&E reserve", "usd"),
    ("noi", "Cash NOI (after FF&E reserve)", "usd"),
)

_STABILIZATION_FIELDS: tuple[tuple[str, str, str], ...] = (
    ("stabilized_year", "Stabilized year", "year"),
    ("source", "Stabilization basis", "text"),
    ("stabilized_occupancy", "Stabilized occupancy", "pct"),
    ("stabilized_adr", "Stabilized ADR", "usd"),
    ("stabilized_revenue", "Stabilized revenue", "usd"),
    ("stabilized_noi_before_reserve", "Stabilized NOI (before FF&E reserve)", "usd"),
    ("stabilized_cash_noi", "Stabilized Cash NOI (after FF&E reserve)", "usd"),
    ("stabilized_noi_margin", "Stabilized NOI margin", "pct"),
)

# (engine, path, provenance key or None, label, unit, group)
_SCALAR_FIELDS: tuple[tuple[str, str, str | None, str, str, str], ...] = (
    ("capital", "total_capital", "total_capital", "Total capitalization", "usd", "Capital"),
    ("capital", "equity_amount", "equity_amount", "Equity", "usd", "Capital"),
    ("capital", "debt_amount", "debt_amount", "Debt (sources)", "usd", "Capital"),
    ("capital", "renovation_total_usd", None, "Renovation (total)", "usd", "Capital"),
    ("capital", "renovation_base_usd", None, "Renovation (base)", "usd", "Capital"),
    ("capital", "renovation_contingency_usd", None, "Renovation contingency", "usd", "Capital"),
    ("capital", "senior_loan_fee_usd", "senior_loan_fee_usd", "Senior loan origination fee", "usd", "Capital"),
    ("capital", "ltc", "ltc", "LTC", "pct", "Capital"),
    ("capital", "price_per_key", "price_per_key", "Price per key", "usd", "Capital"),
    ("debt", "loan_amount", None, "Loan amount", "usd", "Debt"),
    ("debt", "interest_rate", None, "Interest rate", "pct", "Debt"),
    ("debt", "annual_debt_service", None, "Annual debt service", "usd", "Debt"),
    ("debt", "year_one_dscr", "year_one_dscr", "Year-1 DSCR", "x", "Debt"),
    ("debt", "stabilized_dscr", "stabilized_dscr", "Stabilized DSCR", "x", "Debt"),
    ("debt", "avg_dscr", None, "Average DSCR", "x", "Debt"),
    ("debt", "year_one_debt_yield", "year_one_debt_yield", "Year-1 debt yield", "pct", "Debt"),
    ("returns", "levered_irr", "levered_irr", "Levered IRR", "pct", "Returns"),
    ("returns", "unlevered_irr", "unlevered_irr", "Unlevered IRR", "pct", "Returns"),
    ("returns", "equity_multiple", "equity_multiple", "Equity multiple", "x", "Returns"),
    ("returns", "gross_sale_price", "gross_sale_price", "Exit value (gross sale price)", "usd", "Returns"),
    ("returns", "net_proceeds", "net_proceeds", "Net sale proceeds", "usd", "Returns"),
    ("returns", "terminal_noi", "terminal_noi", "Exit NOI", "usd", "Returns"),
    ("returns", "exit_cap_rate", None, "Exit cap rate", "pct", "Returns"),
    ("returns", "hold_years", None, "Hold (years)", "year", "Returns"),
    ("returns", "peak_equity", None, "Peak equity", "usd", "Returns"),
    ("returns", "year_one_coc", "year_one_coc", "Year-1 cash-on-cash", "pct", "Returns"),
    ("partnership", "lp.irr", "lp.irr", "LP IRR", "pct", "Partnership"),
    ("partnership", "gp.irr", "gp.irr", "GP IRR", "pct", "Partnership"),
    ("partnership", "lp.equity_multiple", "lp.equity_multiple", "LP equity multiple", "x", "Partnership"),
    ("partnership", "gp.equity_multiple", "gp.equity_multiple", "GP equity multiple", "x", "Partnership"),
    ("partnership", "lp.distributions", "lp.distributions", "LP distributions", "usd", "Partnership"),
    ("partnership", "gp.distributions", "gp.distributions", "GP distributions", "usd", "Partnership"),
    ("partnership", "promote_amount", "promote_amount", "GP promote", "usd", "Partnership"),
)


def _year_label(row: Any, idx: int) -> str:
    yr = row.get("year") if isinstance(row, Mapping) else None
    return f"Year {yr}" if isinstance(yr, int) else f"Year {idx + 1}"


def _field(
    engine: str,
    path: str,
    label: str,
    unit: str,
    group: str,
    value: Any,
    provenance: dict[str, Any] | None,
) -> dict[str, Any]:
    return {
        "key": f"{engine}.{path}",
        "engine": engine,
        "group": group,
        "label": label,
        "unit": unit,
        "value": value,
        "provenance": provenance,
    }


def extract_headline_fields(engines: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Flatten one run's headline outputs into an ordered field list.

    ``engines`` maps engine name → that engine's ``outputs`` dict (``None``
    or missing when the engine did not complete in the run). Absent values
    stay ``None`` — never a fabricated 0.
    """
    out: list[dict[str, Any]] = []

    rev = engines.get("revenue") or {}
    for i, row in enumerate(rev.get("years") or []):
        ylab = _year_label(row, i)
        for sub, label, unit in _REVENUE_YEAR_FIELDS:
            path = f"years[{i}].{sub}"
            out.append(
                _field("revenue", path, f"{ylab} · {label}", unit, "Revenue",
                       _get_path(rev, path), _provenance_for(rev, path))
            )

    exp = engines.get("expense") or {}
    for i, row in enumerate(exp.get("years") or []):
        ylab = _year_label(row, i)
        for sub, label, unit in _EXPENSE_YEAR_FIELDS:
            path = f"years[{i}].{sub}"
            out.append(
                _field("expense", path, f"{ylab} · {label}", unit, "Expenses & NOI",
                       _get_path(exp, path), _provenance_for(exp, path))
            )
    stab = exp.get("stabilization")
    if isinstance(stab, Mapping):
        for sub, label, unit in _STABILIZATION_FIELDS:
            path = f"stabilization.{sub}"
            out.append(
                _field("expense", path, label, unit, "Stabilization",
                       _get_path(exp, path), None)
            )

    cap = engines.get("capital") or {}
    if cap:
        # Purchase price lives in the Uses list (label "Purchase Price").
        price = None
        price_prov_key = "purchase_price"
        for use in cap.get("uses") or []:
            if isinstance(use, Mapping) and use.get("label") == "Purchase Price":
                price = use.get("amount")
                break
        prov = _provenance_for(cap, price_prov_key)
        out.append(
            _field("capital", "purchase_price", "Purchase price", "usd", "Capital", price, prov)
        )

    for engine, path, prov_key, label, unit, group in _SCALAR_FIELDS:
        src = engines.get(engine) or {}
        if not src:
            continue
        out.append(
            _field(engine, path, label, unit, group, _get_path(src, path),
                   _provenance_for(src, prov_key))
        )
        # DSCR / debt service by year sit right after the debt scalars.
        if engine == "debt" and path == "year_one_debt_yield":
            for i, row in enumerate(src.get("schedule") or []):
                ylab = _year_label(row, i)
                for sub, slabel, sunit in (
                    ("dscr", "DSCR", "x"),
                    ("debt_service", "Debt service", "usd"),
                ):
                    spath = f"schedule[{i}].{sub}"
                    out.append(
                        _field("debt", spath, f"{ylab} · {slabel}", sunit, "Debt",
                               _get_path(src, spath), _provenance_for(src, spath))
                    )

    order = {e: i for i, e in enumerate(ENGINE_ORDER)}
    out.sort(key=lambda f: order.get(f["engine"], len(order)))  # stable
    return out


# ─────────────────────────────── diffing ────────────────────────────────


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def _same(a: Any, b: Any) -> bool:
    if _is_number(a) and _is_number(b):
        return math.isclose(float(a), float(b), rel_tol=REL_TOL, abs_tol=ABS_TOL)
    return a == b


def diff_fields(
    base: Iterable[Mapping[str, Any]],
    head: Iterable[Mapping[str, Any]],
    *,
    include_unchanged: bool = False,
) -> list[dict[str, Any]]:
    """Field-level diff of two ``extract_headline_fields`` lists.

    Each row carries ``base_value``/``head_value``, ``delta_abs`` (head −
    base) and ``delta_pct`` ((head − base) / |base|). Either delta is ``None``
    with a ``delta_reason`` when it cannot be computed (a side missing, a
    non-numeric value, or a zero base for the %). ``status`` is one of
    ``changed`` / ``added`` / ``removed`` / ``unchanged``. Unchanged rows are
    omitted unless ``include_unchanged``.
    """
    base_by_key = {f["key"]: f for f in base}
    head_list = list(head)
    head_by_key = {f["key"]: f for f in head_list}

    keys: list[str] = list(base_by_key)
    keys += [k for k in head_by_key if k not in base_by_key]

    rows: list[dict[str, Any]] = []
    for key in keys:
        b = base_by_key.get(key)
        h = head_by_key.get(key)
        meta = h or b
        assert meta is not None
        bv = b.get("value") if b else None
        hv = h.get("value") if h else None
        if bv is None and hv is None:
            status = "unchanged"
        elif bv is None:
            status = "added"
        elif hv is None:
            status = "removed"
        elif _same(bv, hv):
            status = "unchanged"
        else:
            status = "changed"
        if status == "unchanged" and not include_unchanged:
            continue

        delta_abs: float | None = None
        delta_pct: float | None = None
        reason: str | None = None
        if _is_number(bv) and _is_number(hv):
            delta_abs = float(hv) - float(bv)
            if float(bv) != 0.0:
                delta_pct = delta_abs / abs(float(bv))
            elif delta_abs != 0.0:
                reason = "base is 0 — % change undefined"
            else:
                delta_pct = 0.0
        elif bv is None or hv is None:
            reason = (
                "not in baseline run" if bv is None and hv is not None
                else "not in head run" if hv is None and bv is not None
                else "not in either run"
            )
        else:
            reason = "non-numeric value"

        rows.append(
            {
                "key": key,
                "engine": meta["engine"],
                "group": meta["group"],
                "label": meta["label"],
                "unit": meta["unit"],
                "status": status,
                "base_value": bv,
                "head_value": hv,
                "delta_abs": delta_abs,
                "delta_pct": delta_pct,
                "delta_reason": reason,
                "base_provenance": b.get("provenance") if b else None,
                "head_provenance": h.get("provenance") if h else None,
            }
        )
    return rows


# ───────────────────────────── DB readers ───────────────────────────────


def _decode(value: Any) -> Any:
    if value is None or isinstance(value, (dict, list)):
        return value
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value
    return value


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=UTC)
        return dt.astimezone(UTC).isoformat()
    dt = parse_ts(str(value))
    return dt.isoformat() if dt else str(value)


def parse_ts(value: str | None) -> datetime | None:
    """Parse an ISO-8601-ish timestamp (SQLite TEXT or PG) as an aware UTC time."""
    if not value:
        return None
    s = str(value).strip().replace("Z", "+00:00")
    try:
        dt = datetime.fromisoformat(s)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    return dt.astimezone(UTC)


async def list_runs(
    session: AsyncSession,
    *,
    deal_uuid: str,
    tenant_id: str,
    canonical_engines: Iterable[str],
    canonical_run_id: str | None,
) -> list[dict[str, Any]]:
    """One entry per ``run_id`` on the deal, newest first.

    ``deal_uuid`` is the persisted (coerced) deal id. Reads only the light
    columns — never the outputs blob.
    """
    rows = (
        await session.execute(
            text(
                # tenant-scope predicate required by tenant_middleware
                """
                SELECT run_id, engine_name, status, started_at, completed_at
                  FROM engine_outputs
                 WHERE deal_id = :deal AND tenant_id = :tenant
                   AND run_id IS NOT NULL
                """
            ),
            {"deal": deal_uuid, "tenant": tenant_id},
        )
    ).fetchall()

    scen_rows = (
        await session.execute(
            text(
                # tenant-scope predicate required by tenant_middleware
                """
                SELECT id, name, is_base, last_run_id
                  FROM scenarios
                 WHERE deal_id = :deal AND tenant_id = :tenant
                   AND last_run_id IS NOT NULL
                """
            ),
            {"deal": deal_uuid, "tenant": tenant_id},
        )
    ).fetchall()
    scenario_by_run: dict[str, dict[str, Any]] = {}
    for r in scen_rows:
        m = r._mapping
        scenario_by_run[str(m["last_run_id"])] = {
            "id": str(m["id"]),
            "name": m["name"],
            "is_base": bool(m["is_base"]),
        }

    canon = set(canonical_engines)
    runs: dict[str, dict[str, Any]] = {}
    for r in rows:
        m = r._mapping
        rid = str(m["run_id"])
        entry = runs.setdefault(
            rid,
            {"run_id": rid, "_started": [], "_completed": [], "_status": [], "engines": []},
        )
        entry["engines"].append(m["engine_name"])
        entry["_status"].append(m["status"])
        st = parse_ts(_iso(m["started_at"]))
        if st:
            entry["_started"].append(st)
        ct = parse_ts(_iso(m["completed_at"])) if m["completed_at"] is not None else None
        entry["_completed"].append(ct)

    out: list[dict[str, Any]] = []
    for rid, e in runs.items():
        statuses = e.pop("_status")
        started = e.pop("_started")
        completed = e.pop("_completed")
        if any(s == "running" for s in statuses) or any(c is None for c in completed):
            status = "running"
        elif any(s == "failed" for s in statuses):
            status = "failed" if all(s == "failed" for s in statuses) else "partial"
        else:
            status = "complete"
        done = [c for c in completed if c is not None]
        scen = scenario_by_run.get(rid)
        engines = sorted(set(e["engines"]))
        out.append(
            {
                "run_id": rid,
                "started_at": min(started).isoformat() if started else None,
                "completed_at": (
                    max(done).isoformat() if done and status != "running" else None
                ),
                "engines": engines,
                "status": status,
                "is_full_chain": canon.issubset(engines),
                "is_canonical": rid == canonical_run_id,
                "scenario": scen,
            }
        )
    out.sort(key=lambda r: r["started_at"] or "", reverse=True)
    return out


def pick_baseline(runs: list[dict[str, Any]], before: datetime) -> dict[str, Any] | None:
    """Newest complete, full-chain, non-scenario run completed before ``before``."""
    for r in runs:  # newest first
        if r["status"] != "complete" or not r["is_full_chain"]:
            continue
        scen = r.get("scenario")
        if scen and not scen.get("is_base"):
            continue
        done = parse_ts(r.get("completed_at"))
        if done and done < before:
            return r
    return None


async def load_run_engines(
    session: AsyncSession,
    *,
    deal_uuid: str,
    tenant_id: str,
    run_id: str,
) -> dict[str, dict[str, Any]]:
    """``{engine: {status, outputs, started_at, completed_at}}`` for one run."""
    rows = (
        await session.execute(
            text(
                # tenant-scope predicate required by tenant_middleware
                """
                SELECT engine_name, status, outputs, started_at, completed_at
                  FROM engine_outputs
                 WHERE deal_id = :deal AND tenant_id = :tenant AND run_id = :run
                 ORDER BY started_at ASC
                """
            ),
            {"deal": deal_uuid, "tenant": tenant_id, "run": run_id},
        )
    ).fetchall()
    out: dict[str, dict[str, Any]] = {}
    for r in rows:
        m = r._mapping
        out[m["engine_name"]] = {
            "status": m["status"],
            "outputs": _decode(m["outputs"]) if m["status"] == "complete" else None,
            "started_at": _iso(m["started_at"]),
            "completed_at": _iso(m["completed_at"]),
        }
    return out


def outputs_only(engines: Mapping[str, Mapping[str, Any]]) -> dict[str, Any]:
    return {name: (row.get("outputs") or None) for name, row in engines.items()}


# ─────────────────────────────── xlsx ───────────────────────────────────


def build_compare_workbook(result: Mapping[str, Any]) -> bytes:
    """Render a compare result (the API response dict) as an .xlsx workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from ..export.excel import (
        MULT_FMT,
        PCT_FMT,
        USD_FMT,
        _autosize,
        _freeze_top,
        _style_header_row,
        _style_subhead_row,
    )

    fmt_for = {"usd": USD_FMT, "pct": PCT_FMT, "x": MULT_FMT, "year": "0"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    base = result.get("base") or {}
    head = result.get("head") or {}
    summary = [
        ("Deal", result.get("deal_id")),
        ("Baseline run", base.get("run_id")),
        ("Baseline started", base.get("started_at")),
        ("Baseline completed", base.get("completed_at")),
        ("Head run", head.get("run_id")),
        ("Head started", head.get("started_at")),
        ("Head completed", head.get("completed_at")),
        ("Fields compared", result.get("fields_compared")),
        ("Fields changed", result.get("fields_changed")),
        ("Unchanged rows included", "yes" if result.get("include_unchanged") else "no"),
        ("Tolerance", f"rel {REL_TOL:g} / abs {ABS_TOL:g}"),
        ("Generated", datetime.now(UTC).isoformat()),
    ]
    ws.append(["Run comparison", ""])
    _style_header_row(ws, 1, 2)
    for k, v in summary:
        ws.append([k, "—" if v is None else v])
    _autosize(ws, [26, 60])

    ds = wb.create_sheet("Changed fields")
    headers = [
        "Engine", "Group", "Field", "Status", "Baseline", "Head",
        "Δ (abs)", "Δ (%)", "Note", "Baseline source", "Head source",
    ]
    ds.append(headers)
    _style_header_row(ds, 1, len(headers))
    current_group = None
    for row in result.get("fields") or []:
        if row["group"] != current_group:
            current_group = row["group"]
            ds.append([current_group])
            _style_subhead_row(ds, ds.max_row, len(headers))
        ds.append(
            [
                row["engine"],
                row["group"],
                row["label"],
                row["status"],
                "—" if row["base_value"] is None else row["base_value"],
                "—" if row["head_value"] is None else row["head_value"],
                "—" if row["delta_abs"] is None else row["delta_abs"],
                "—" if row["delta_pct"] is None else row["delta_pct"],
                row.get("delta_reason") or "",
                _source_text(row.get("base_provenance")),
                _source_text(row.get("head_provenance")),
            ]
        )
        r = ds.max_row
        nf = fmt_for.get(row["unit"])
        if nf:
            for col in (5, 6, 7):
                if isinstance(ds.cell(row=r, column=col).value, (int, float)):
                    ds.cell(row=r, column=col).number_format = nf
        if isinstance(ds.cell(row=r, column=8).value, (int, float)):
            ds.cell(row=r, column=8).number_format = PCT_FMT
    if not result.get("fields"):
        ds.append(["No headline field differs between the two runs."])
        ds.cell(row=ds.max_row, column=1).font = Font(italic=True)
    _freeze_top(ds)
    _autosize(ds, [12, 16, 44, 10, 16, 16, 16, 10, 28, 36, 36])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _source_text(prov: Mapping[str, Any] | None) -> str:
    if not prov:
        return ""
    parts = []
    if prov.get("source"):
        src = prov["source"]
        parts.append(json.dumps(src, default=str) if not isinstance(src, str) else src)
    if prov.get("state"):
        parts.append(str(prov["state"]))
    if prov.get("formula"):
        parts.append(str(prov["formula"]))
    return " · ".join(parts)
