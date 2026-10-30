"""FON-41 E-013 / E-017 — P&L Excel export + validated re-import.

Historical: the export carries stable ``document_id::field_name`` ids; the
import preview reports changed values, mapping errors and non-numeric cells;
Apply writes through ``review_extraction_field`` (reviewed="edited" + audit).

Projections: the Assumptions sheet carries every editable override key; the
import enforces the FON-74 note rule, reports engine-computed edits as
"computed", and Apply writes ``field_overrides`` through ``update_deal``.
"""

from __future__ import annotations

import io
import json
import os
from datetime import UTC, datetime
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import text

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite:///./fondok.db")

_TENANT = "23b0cff3-6f9b-57a9-8d2a-5511f3dd9f7e"
_OTHER = "8d0c0d4e-1b52-4d1e-9a31-77c2f0c3a001"


def _auth(tenant: str = _TENANT) -> object:
    from app.auth.context import AuthContext

    return AuthContext(
        tenant_id=UUID(tenant), user_id="user_analyst", role="member",
        source="jwt", org_id=None, email="analyst@fondok.test",
    )


FIELDS_2023 = [
    {"field_name": "rooms_revenue", "value": 9_810_000, "confidence": 0.55},
    {"field_name": "fb_revenue", "value": 2_100_000, "confidence": 0.95},
    {"field_name": "adr", "value": 247, "confidence": 0.96},
    {"field_name": "gop", "value": 4_000_000, "confidence": 0.9},
]
FIELDS_2024 = [
    {"field_name": "rooms_revenue", "value": "10,250,000", "confidence": 0.9},
    {"field_name": "adr", "value": 255.5, "confidence": 0.96},
]


async def _setup(overrides: dict | None = None) -> tuple[UUID, UUID, UUID]:
    from app.database import get_session_factory

    deal_id, d23, d24 = uuid4(), uuid4(), uuid4()
    async with get_session_factory()() as s:
        ts = datetime.now(UTC)
        await s.execute(
            text(
                "INSERT INTO deals (id, tenant_id, name, status, ai_confidence, field_overrides, "
                "created_at, updated_at) VALUES (:id,:t,'Round-trip','Draft',0.0,:fo,:ts,:ts)"
            ),
            {"id": str(deal_id), "t": _TENANT, "ts": ts, "fo": json.dumps(overrides or {})},
        )
        for doc_id, fy, fields, name in ((d23, 2023, FIELDS_2023, "pnl-2023.xlsx"),
                                         (d24, 2024, FIELDS_2024, "pnl-2024.xlsx")):
            await s.execute(
                text(
                    "INSERT INTO documents (id, deal_id, tenant_id, filename, doc_type, status, "
                    "fiscal_year, uploaded_at) VALUES (:id,:deal,:t,:fn,'PNL','EXTRACTED',:fy,:ts)"
                ),
                {"id": str(doc_id), "deal": str(deal_id), "t": _TENANT, "fn": name,
                 "fy": fy, "ts": ts},
            )
            await s.execute(
                text(
                    "INSERT INTO extraction_results (id, document_id, deal_id, tenant_id, fields, "
                    "confidence_report, agent_version, created_at) "
                    "VALUES (:id,:doc,:deal,:t,:f,'{}','v1',:ts)"
                ),
                {"id": str(uuid4()), "doc": str(doc_id), "deal": str(deal_id), "t": _TENANT,
                 "f": json.dumps(fields), "ts": ts},
            )
        await s.commit()
    return deal_id, d23, d24


async def _export_hist(deal_id: UUID, tenant: str = _TENANT) -> bytes:
    from app.api.pl_roundtrip import export_historicals
    from app.database import get_session_factory

    async with get_session_factory()() as s:
        resp = await export_historicals(deal_id=deal_id, session=s, auth=_auth(tenant))
    return resp.body


class _Upload:
    """Minimal UploadFile stand-in (the route only calls ``read``)."""

    def __init__(self, data: bytes) -> None:
        self._data = data

    async def read(self, n: int = -1) -> bytes:
        return self._data if n < 0 else self._data[:n]


def _sheet(data: bytes, name: str):
    return load_workbook(io.BytesIO(data))[name]


def _find_row(ws, line_id: str) -> int:
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == line_id:
            return r
    raise AssertionError(line_id)


def _save(ws) -> bytes:
    buf = io.BytesIO()
    ws.parent.save(buf)
    return buf.getvalue()


# ─────────────────────────── historical ───────────────────────────


@pytest.mark.asyncio
async def test_historical_export_shape_carries_stable_ids() -> None:
    deal_id, d23, d24 = await _setup()
    data = await _export_hist(deal_id)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Historical P&L", "README"]
    ws = wb["Historical P&L"]
    # Column pairs: value + hidden fondok_id; row 2 carries the document id.
    assert ws.cell(row=1, column=3).value == "FY 2023"
    assert ws.cell(row=1, column=4).value == "fondok_id"
    assert ws.cell(row=2, column=4).value == str(d23)
    assert ws.cell(row=1, column=5).value == "FY 2024"
    assert ws.cell(row=2, column=6).value == str(d24)
    assert ws.column_dimensions["A"].hidden and ws.column_dimensions["D"].hidden
    r = _find_row(ws, "rooms_revenue")
    assert ws.cell(row=r, column=2).value == "Rooms Revenue"
    assert ws.cell(row=r, column=3).value == 9_810_000
    assert ws.cell(row=r, column=4).value == f"{d23}::rooms_revenue"
    assert ws.cell(row=r, column=5).value == 10_250_000  # "10,250,000" → number
    # A line the 2024 statement never published: no value, no id.
    g = _find_row(ws, "gop")
    assert ws.cell(row=g, column=5).value is None and ws.cell(row=g, column=6).value is None
    readme = {row[0]: row[1] for row in wb["README"].iter_rows(values_only=True) if row[0]}
    assert readme["format"] == "fondok.historicals.v1"
    assert readme["deal_id"] == str(deal_id)


@pytest.mark.asyncio
async def test_historical_import_preview_diff_and_errors() -> None:
    from app.api.pl_roundtrip import preview_historicals_import
    from app.database import get_session_factory

    deal_id, d23, _d24 = await _setup()
    ws = _sheet(await _export_hist(deal_id), "Historical P&L")
    r_rooms, r_adr, r_gop, r_fb = (_find_row(ws, k) for k in ("rooms_revenue", "adr", "gop",
                                                                "fb_revenue"))
    ws.cell(row=r_rooms, column=3).value = 9_900_000          # changed
    ws.cell(row=r_adr, column=3).value = "two hundred"        # non-numeric
    ws.cell(row=r_gop, column=5).value = 123                  # no source line in 2024
    ws.cell(row=r_fb, column=4).value = f"{uuid4()}::fb_revenue"  # tampered id
    data = _save(ws)

    async with get_session_factory()() as s:
        preview = await preview_historicals_import(
            deal_id=deal_id, session=s, auth=_auth(), file=_Upload(data)
        )
    assert [(c["cell_id"], c["old_value"], c["new_value"]) for c in preview["changes"]] == [
        (f"{d23}::rooms_revenue", 9_810_000, 9_900_000)
    ]
    assert preview["changes"][0]["filename"] == "pnl-2023.xlsx"
    assert preview["changes"][0]["period_label"] == "FY 2023"
    reasons = sorted(e["reason"] for e in preview["mapping_errors"])
    assert reasons == ["no_source_line", "unknown_id"]
    assert [n["raw"] for n in preview["non_numeric"]] == ["two hundred"]

    # Nothing was written by the preview.
    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT fields FROM extraction_results WHERE document_id = :d AND tenant_id = :t"),
            {"d": str(d23), "t": _TENANT},
        )).first()
    assert json.loads(row[0])[0]["value"] == 9_810_000


@pytest.mark.asyncio
async def test_historical_import_refuses_another_deals_workbook() -> None:
    from app.api.pl_roundtrip import preview_historicals_import
    from app.database import get_session_factory

    deal_a, *_ = await _setup()
    deal_b, *_ = await _setup()
    data = await _export_hist(deal_a)
    async with get_session_factory()() as s:
        with pytest.raises(HTTPException) as exc:
            await preview_historicals_import(
                deal_id=deal_b, session=s, auth=_auth(), file=_Upload(data)
            )
    assert exc.value.status_code == 422
    assert exc.value.detail["code"] == "workbook_format"


@pytest.mark.asyncio
async def test_historical_apply_goes_through_the_review_path() -> None:
    from app.api.pl_roundtrip import HistApplyBody, apply_historicals_import
    from app.database import get_session_factory

    deal_id, d23, _ = await _setup()
    cid = f"{d23}::rooms_revenue"
    body = HistApplyBody(changes=[
        {"cell_id": cid, "new_value": 9_900_000, "old_value": 9_810_000,
         "note": "Workbook: rooms revenue restated per the 2023 audit"},  # FON-74 note
        {"cell_id": f"{d23}::adr", "new_value": 250, "old_value": 199, "note": "x"},  # stale preview
        {"cell_id": f"{uuid4()}::x", "new_value": 1, "note": "x"},
    ])
    async with get_session_factory()() as s:
        res = await apply_historicals_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert [a["cell_id"] for a in res["applied"]] == [cid]
    assert sorted(x["reason"] for x in res["skipped"]) == ["stale", "unknown_id"]

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT fields, confidence_report FROM extraction_results "
                 "WHERE document_id = :d AND tenant_id = :t"),
            {"d": str(d23), "t": _TENANT},
        )).first()
        fields = {f["field_name"]: f for f in json.loads(row[0])}
        audits = (await s.execute(
            text("SELECT action FROM audit_log WHERE tenant_id = :t AND resource_id = :r"),
            {"t": _TENANT, "r": str(d23)},
        )).fetchall()
    # Exactly the manual-correction semantics.
    assert fields["rooms_revenue"]["value"] == 9_900_000
    assert fields["rooms_revenue"]["reviewed"] == "edited"
    assert fields["rooms_revenue"]["confidence"] == 1.0
    assert "rooms_revenue" not in json.loads(row[1])["low_confidence_fields"]
    assert "reviewed" not in fields["adr"]
    assert ("extraction_field.edit",) in [tuple(a) for a in audits]


@pytest.mark.asyncio
async def test_historical_routes_are_tenant_scoped() -> None:
    deal_id, *_ = await _setup()
    with pytest.raises(HTTPException) as exc:
        await _export_hist(deal_id, tenant=_OTHER)
    assert exc.value.status_code == 404


# ─────────────────────────── projections ───────────────────────────


def _assumption_rows() -> list[dict]:
    from app.export.pl_workbooks import ASSUMPTIONS

    vals = {"revpar_growth": 0.045, "expense_growth": 0.035, "hold_years": 5,
            "stabilization_year": 3, "mgmt_fee_pct": 0.03}
    return [
        {"key": a.key, "label": a.label, "value": vals.get(a.key), "unit": a.unit,
         "source": "seed", "method": a.method,
         "note_required": a.key != "stabilization_year", "current_note": None}
        for a in ASSUMPTIONS
    ]


def _proj_workbook() -> tuple[bytes, dict]:
    from app.export.pl_workbooks import build_projections_workbook, projection_values

    outputs = {
        "revenue": {"years": [{"occupancy": 0.7, "adr": 300.0, "revpar": 210.0,
                               "rooms_revenue": 10_000_000.0}] * 2},
        "expense": {"years": [{"gop": 4_000_000.0, "dept_expenses": {"rooms": 2_000_000.0}}] * 2},
    }
    n, values = projection_values(outputs)
    assert n == 2
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
        values=values, assumptions=_assumption_rows(),
    )
    return data, values


def _parse(data: bytes, values: dict):
    from app.export.pl_workbooks import parse_projections_import

    cur = {r["key"]: r["value"] for r in _assumption_rows()}
    return parse_projections_import(
        data, deal_id="deal-1", current_values=values, current_assumptions=cur,
        needs_note=lambda k: k != "stabilization_year",
    )


def test_projection_export_lists_every_assumption_with_key_and_source() -> None:
    from app.export.pl_workbooks import ASSUMPTIONS

    data, _ = _proj_workbook()
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Future P&L", "Assumptions", "README"]
    wa = wb["Assumptions"]
    assert [c.value for c in wa[1]][:3] == ["Key", "Assumption", "Value"]
    keys = [wa.cell(row=r, column=1).value for r in range(2, wa.max_row + 1)]
    assert keys == [a.key for a in ASSUMPTIONS]
    assert wa.cell(row=2 + keys.index("revpar_growth"), column=5).value == "seed"
    # Unset assumptions stay blank — never a UI default.
    assert wa.cell(row=2 + keys.index("starting_adr"), column=3).value is None
    wp = wb["Future P&L"]
    # Row 2 is the hidden column-key row; engine lines follow the room counts.
    assert wp.cell(row=2, column=1).value == "fondok_column"
    assert wp.row_dimensions[2].hidden
    paths = [wp.cell(row=r, column=1).value for r in range(3, wp.max_row + 1)]
    occ = 3 + paths.index("revenue.occupancy")
    assert wp.cell(row=1, column=3).value == "Base year (Year 1) Amount"
    assert wp.cell(row=occ, column=3).value == 0.7


def test_projection_import_note_rule_and_computed_cells() -> None:
    data, values = _proj_workbook()
    wb = load_workbook(io.BytesIO(data))
    wa = wb["Assumptions"]
    keys = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    wa.cell(row=keys["revpar_growth"], column=3).value = 0.05
    wa.cell(row=keys["revpar_growth"], column=9).value = "Broker comp set supports 5%"
    wa.cell(row=keys["expense_growth"], column=3).value = 0.04            # no note → rejected
    wa.cell(row=keys["stabilization_year"], column=3).value = 4           # exempt → no note ok
    wa.cell(row=keys["hold_years"], column=3).value = 5.5                 # not a whole number
    wa.cell(row=keys["mgmt_fee_pct"], column=3).value = "3%"              # non-numeric
    r = wa.max_row + 1
    wa.cell(row=r, column=1).value = "made_up_key"
    wa.cell(row=r, column=3).value = 1
    wp = wb["Future P&L"]
    occ = next(r for r in range(3, wp.max_row + 1)
               if wp.cell(row=r, column=1).value == "revenue.occupancy")
    wp.cell(row=occ, column=3).value = 0.75                               # computed cell
    buf = io.BytesIO()
    wb.save(buf)

    p = _parse(buf.getvalue(), values)
    assert [(c["key"], c["old_value"], c["new_value"]) for c in p["changes"]] == [
        ("revpar_growth", 0.045, 0.05), ("stabilization_year", 3, 4),
    ]
    assert p["changes"][0]["note"] == "Broker comp set supports 5%"
    assert [(x["key"], x["reason"]) for x in p["rejected"]] == [("expense_growth", "note_required")]
    assert sorted(x["key"] for x in p["non_numeric"]) == ["hold_years", "mgmt_fee_pct"]
    assert [x["reason"] for x in p["mapping_errors"]] == ["unknown_key"]
    assert len(p["computed_edits"]) == 1
    ce = p["computed_edits"][0]
    assert ce["cell_key"] == "revenue.years[0].occupancy" and ce["reason"] == "computed"
    assert ce["old_value"] == 0.7 and ce["new_value"] == 0.75


@pytest.mark.asyncio
async def test_projection_apply_writes_field_overrides_through_update_deal() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory

    deal_id, *_ = await _setup(overrides={"exit_cap_rate": {"value": 0.08, "note": "kept"}})
    body = ProjApplyBody(changes=[
        {"key": "revpar_growth", "new_value": 0.0512, "note": "STR forecast"},
        {"key": "expense_growth", "new_value": 0.05, "note": ""},     # note required → skipped
        {"key": "stabilization_year", "new_value": 2, "note": ""},    # exempt
        {"key": "made_up", "new_value": 1, "note": "x"},
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert sorted(a["key"] for a in res["applied"]) == ["revpar_growth", "stabilization_year"]
    assert sorted(x["reason"] for x in res["skipped"]) == ["note_required", "unknown_key"]
    assert res["rerun_required"] is True

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
        audits = (await s.execute(
            text("SELECT action FROM audit_log WHERE tenant_id = :t AND resource_id = :r"),
            {"t": _TENANT, "r": str(deal_id)},
        )).fetchall()
    fo = json.loads(row[0])
    assert fo["revpar_growth"] == {"value": 0.0512, "note": "STR forecast"}
    assert fo["stabilization_year"] == 2
    assert fo["exit_cap_rate"] == {"value": 0.08, "note": "kept"}  # untouched
    assert "expense_growth" not in fo
    assert ("override.set",) in [tuple(a) for a in audits]


@pytest.mark.asyncio
async def test_projection_export_route_matches_assumption_sources() -> None:
    """The Assumptions sheet's values/sources are the loader's — the same map
    ``GET /deals/{id}/assumption_sources`` serves — and an override shows as one."""
    from app.api.deals import get_assumption_sources
    from app.api.pl_roundtrip import export_projections
    from app.database import get_session_factory

    deal_id, *_ = await _setup(overrides={"revpar_growth": {"value": 0.051, "note": "n"}})
    async with get_session_factory()() as s:
        resp = await export_projections(deal_id=deal_id, session=s, auth=_auth())
        srcs = await get_assumption_sources(deal_id=deal_id, session=s, tenant_id=UUID(_TENANT))
    wa = _sheet(resp.body, "Assumptions")
    rows = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    r = rows["revpar_growth"]
    assert wa.cell(row=r, column=3).value == 0.051
    assert wa.cell(row=r, column=5).value == "analyst_override"
    assert wa.cell(row=r, column=8).value == "n"
    r2 = rows["expense_growth"]
    assert wa.cell(row=r2, column=3).value == srcs.values.get("expense_growth")
    assert wa.cell(row=r2, column=5).value == srcs.sources.get("expense_growth")


# ─────────────────────── E-016 projection methods ───────────────────────


def test_projection_export_carries_per_line_method_rows() -> None:
    from app.export.pl_workbooks import PROJECTION_METHOD_LINE_LABELS

    data, _ = _proj_workbook()
    wa = _sheet(data, "Assumptions")
    keys = [wa.cell(row=r, column=1).value for r in range(2, wa.max_row + 1)]
    for line, _label in PROJECTION_METHOD_LINE_LABELS:
        assert f"projection_methods.{line}.method" in keys
        assert f"projection_methods.{line}.value" in keys


def test_projection_import_method_choice_rules() -> None:
    from app.export.pl_workbooks import parse_projections_import

    data, values = _proj_workbook()
    wb = load_workbook(io.BytesIO(data))
    wa = wb["Assumptions"]
    keys = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    m_rooms = "projection_methods.rooms_dept_expense.method"
    m_util = "projection_methods.utilities.method"
    m_ag = "projection_methods.administrative_general.method"
    wa.cell(row=keys[m_rooms], column=3).value = "POR"
    wa.cell(row=keys[m_rooms], column=9).value = "Brand standard staffing per occupied room"
    wa.cell(row=keys[m_util], column=3).value = "per key"          # not a method code
    wa.cell(row=keys[m_ag], column=3).value = None                  # cleared
    buf = io.BytesIO()
    wb.save(buf)

    cur = {r["key"]: r["value"] for r in _assumption_rows()}
    cur[m_rooms] = "growth"
    cur[m_util] = "growth"
    cur[m_ag] = "pct_revenue"
    p = parse_projections_import(
        buf.getvalue(), deal_id="deal-1", current_values=values,
        current_assumptions=cur, needs_note=lambda k: True,
    )
    assert [(c["key"], c["old_value"], c["new_value"]) for c in p["changes"]] == [
        (m_rooms, "growth", "por"),
    ]
    assert [x["key"] for x in p["non_numeric"]] == [m_util]
    assert (m_ag, "cleared") in [(x["key"], x["reason"]) for x in p["rejected"]]


@pytest.mark.asyncio
async def test_projection_apply_method_and_value_stay_coherent() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory

    m = "projection_methods.{}.method".format
    v = "projection_methods.{}.value".format
    deal_id, *_ = await _setup(overrides={
        m("utilities"): {"value": "growth", "note": "a"},
        v("utilities"): {"value": 0.04, "note": "a"},
    })
    body = ProjApplyBody(changes=[
        # method switched without a value → the old growth rate is dropped
        {"key": m("utilities"), "new_value": "PAR", "note": "per-key utility contracts"},
        # method + value together
        {"key": m("rooms_dept_expense"), "new_value": "por", "note": "staffing model"},
        {"key": v("rooms_dept_expense"), "new_value": 42.5, "note": "staffing model"},
        # value on a line with no method and no model run → refused
        {"key": v("sales_marketing"), "new_value": 0.06, "note": "x"},
        {"key": m("information_telecom"), "new_value": "per key", "note": "x"},
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert sorted(a["key"] for a in res["applied"]) == sorted(
        [m("utilities"), m("rooms_dept_expense"), v("rooms_dept_expense")]
    )
    assert sorted(x["reason"] for x in res["skipped"]) == ["invalid_choice", "method_required"]

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    fo = json.loads(row[0])
    assert fo[m("utilities")] == {"value": "par", "note": "per-key utility contracts"}
    assert v("utilities") not in fo
    assert fo[m("rooms_dept_expense")]["value"] == "por"
    assert fo[v("rooms_dept_expense")] == {"value": 42.5, "note": "staffing model"}
    assert v("sales_marketing") not in fo


# ─────────────────────── FON-97 Benchmark % column ───────────────────────


_BENCHMARK_FIELDS = [
    {"field_name": "pnl_benchmark.peer_set_size", "value": 11},
    {"field_name": "pnl_benchmark.peer.rooms_revenue.total_usd", "value": 10_000_000},
    {"field_name": "pnl_benchmark.peer.total_revenue.total_usd", "value": 15_000_000},
    {"field_name": "pnl_benchmark.peer.rooms_dept_expense.total_usd", "value": 2_500_000},
    {"field_name": "pnl_benchmark.peer.a_and_g.total_usd", "value": 1_200_000},
    # POR only — the category exists but publishes no ratio: the cell stays blank.
    {"field_name": "pnl_benchmark.peer.insurance.por_usd", "value": 6.25},
]


def test_benchmark_column_values_map_only_ratioed_lines() -> None:
    from app.api.documents import PnlBenchmarkCategory
    from app.export.pl_workbooks import benchmark_column_values
    from app.services.pnl_benchmark_map import map_benchmark_to_categories

    flat = {f["field_name"]: f["value"] for f in _BENCHMARK_FIELDS}
    cats = map_benchmark_to_categories(flat)
    assert {c["key"] for c in cats} == {"rooms", "administrative_general", "insurance"}
    bm = benchmark_column_values(cats)
    assert bm == {
        "expense.dept_expenses.rooms": pytest.approx(0.25),                   # dept basis
        "expense.undistributed.administrative_general": pytest.approx(0.08),  # total basis
    }
    # Pydantic models (the market-data block) read the same way as dicts.
    models = [PnlBenchmarkCategory(**c) for c in cats]
    assert benchmark_column_values(models) == bm
    assert benchmark_column_values(None) == {}


def test_projection_export_benchmark_column_is_read_only_and_after_the_years() -> None:
    from app.export.pl_workbooks import (
        BENCHMARK_HEADER, build_projections_workbook, projection_values,
    )

    outputs = {
        "revenue": {"years": [{"occupancy": 0.7, "rooms_revenue": 10_000_000.0}] * 2},
        "expense": {"years": [{"dept_expenses": {"rooms": 2_800_000.0}}] * 2},
    }
    n, values = projection_values(outputs)
    bm = {"expense.dept_expenses.rooms": 0.25}
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
        values=values, assumptions=_assumption_rows(), benchmark=bm,
    )
    wp = _sheet(data, "Future P&L")
    # The column sits AFTER every year group (Amount | % Rev | PAR | POR) and
    # the Exit Year column, so year indices are untouched.
    bm_col = 3 + 4 * n + 1
    assert wp.cell(row=1, column=bm_col).value == BENCHMARK_HEADER
    assert wp.cell(row=2, column=bm_col).value == "benchmark"
    assert wp.max_column == bm_col
    rows = {wp.cell(row=r, column=1).value: r for r in range(3, wp.max_row + 1)}
    rooms = rows["expense.dept_expenses.rooms"]
    assert wp.cell(row=rooms, column=bm_col).value == 0.25
    assert wp.cell(row=rooms, column=bm_col).number_format == "0.0%"
    # A line the benchmark did not publish is blank — never zero-filled.
    assert wp.cell(row=rows["revenue.occupancy"], column=bm_col).value is None
    # The year cells are exactly what they were without a benchmark.
    assert wp.cell(row=rooms, column=3).value == 2_800_000.0

    # No benchmark on the deal → no column at all (not a blank one).
    plain = build_projections_workbook(
        deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
        values=values, assumptions=_assumption_rows(),
    )
    assert _sheet(plain, "Future P&L").max_column == bm_col - 1


def test_projection_import_ignores_the_benchmark_column() -> None:
    from app.export.pl_workbooks import build_projections_workbook, projection_values

    outputs = {
        "revenue": {"years": [{"occupancy": 0.7}] * 2},
        "expense": {"years": [{"dept_expenses": {"rooms": 2_800_000.0}}] * 2},
    }
    n, values = projection_values(outputs)
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
        values=values, assumptions=_assumption_rows(),
        benchmark={"expense.dept_expenses.rooms": 0.25},
    )
    wb = load_workbook(io.BytesIO(data))
    wp = wb["Future P&L"]
    rows = {wp.cell(row=r, column=1).value: r for r in range(2, wp.max_row + 1)}
    header = [c.value for c in wp[1]]
    bm_col = header.index("Benchmark %") + 1
    assert bm_col == 3 + 4 * n + 1
    # An analyst types over the benchmark cell (and even a blank one) …
    wp.cell(row=rows["expense.dept_expenses.rooms"], column=bm_col).value = 0.99
    wp.cell(row=rows["revenue.occupancy"], column=bm_col).value = "n/a"
    buf = io.BytesIO()
    wb.save(buf)

    p = _parse(buf.getvalue(), values)
    # … and nothing is diffed, rejected or reported: the column is not a year.
    assert p["computed_edits"] == []
    assert p["mapping_errors"] == []
    assert p["non_numeric"] == []
    assert p["changes"] == []


@pytest.mark.asyncio
async def test_projection_export_route_feeds_benchmark_from_market_data() -> None:
    """The export's Benchmark % column comes from the same ``pnl_benchmark``
    block ``GET /deals/{id}/market-data`` serves — and is absent without one."""
    from app.api.documents import get_market_data
    from app.api.pl_roundtrip import export_projections
    from app.database import get_session_factory
    from app.export.pl_workbooks import BENCHMARK_HEADER

    deal_id, *_ = await _setup()
    async with get_session_factory()() as s:
        before = await export_projections(deal_id=deal_id, session=s, auth=_auth())
    wp0 = _sheet(before.body, "Future P&L")
    assert BENCHMARK_HEADER not in [c.value for c in wp0[1]]

    bench_doc = uuid4()
    async with get_session_factory()() as s:
        ts = datetime.now(UTC)
        await s.execute(
            text(
                "INSERT INTO documents (id, deal_id, tenant_id, filename, doc_type, status, "
                "uploaded_at) VALUES (:id,:deal,:t,:fn,'PNL_BENCHMARK','EXTRACTED',:ts)"
            ),
            {"id": str(bench_doc), "deal": str(deal_id), "t": _TENANT,
             "fn": "CBRE-Trends-2025-upper-upscale.pdf", "ts": ts},
        )
        await s.execute(
            text(
                "INSERT INTO extraction_results (id, document_id, deal_id, tenant_id, fields, "
                "confidence_report, agent_version, created_at) "
                "VALUES (:id,:doc,:deal,:t,:f,'{}','v1',:ts)"
            ),
            {"id": str(uuid4()), "doc": str(bench_doc), "deal": str(deal_id), "t": _TENANT,
             "f": json.dumps(_BENCHMARK_FIELDS), "ts": ts},
        )
        await s.commit()

    async with get_session_factory()() as s:
        resp = await export_projections(deal_id=deal_id, session=s, auth=_auth())
        market = await get_market_data(deal_id=deal_id, session=s, tenant_id=UUID(_TENANT))
    assert market.pnl_benchmark is not None
    assert market.pnl_benchmark.source_filename == "CBRE-Trends-2025-upper-upscale.pdf"
    by_key = {c.key: c for c in market.pnl_benchmark.categories}

    wp = _sheet(resp.body, "Future P&L")
    header = [c.value for c in wp[1]]
    assert header[-1] == BENCHMARK_HEADER
    bm_col = len(header)
    rows = {wp.cell(row=r, column=1).value: r for r in range(2, wp.max_row + 1)}
    assert wp.cell(row=rows["expense.dept_expenses.rooms"], column=bm_col).value == pytest.approx(
        by_key["rooms"].ratio
    )
    assert wp.cell(
        row=rows["expense.undistributed.administrative_general"], column=bm_col
    ).value == pytest.approx(by_key["administrative_general"].ratio)
    # Insurance was priced POR-only: present in the block with no ratio → blank cell.
    assert by_key["insurance"].ratio is None
    assert wp.cell(row=rows["expense.fixed_charges.insurance"], column=bm_col).value is None


# ─────────────────── FON-95 part 2 — penetration targets by year ───────────────────


@pytest.mark.asyncio
async def test_projection_export_carries_by_year_target_rows_for_the_hold() -> None:
    """One row per metric per model year; the by-year entry is the value (the
    scalar shorthand is its own key), its note rides along, blank = not set."""
    from app.api.pl_roundtrip import export_projections
    from app.database import get_session_factory

    deal_id, *_ = await _setup(overrides={
        "hold_years": 3,
        "index_mpi_target_by_year": {"value": [None, 110.0], "note": "Share gain in Year 2"},
        "index_ari_target": {"value": 100.0, "note": "parity"},
    })
    async with get_session_factory()() as s:
        resp = await export_projections(deal_id=deal_id, session=s, auth=_auth())
    wa = _sheet(resp.body, "Assumptions")
    rows = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    assert [k for k in rows if k.startswith("index_mpi_target_by_year.")] == [
        "index_mpi_target_by_year.1", "index_mpi_target_by_year.2", "index_mpi_target_by_year.3",
    ]
    assert "index_ari_target_by_year.3" in rows and "index_ari_target_by_year.4" not in rows
    r2 = rows["index_mpi_target_by_year.2"]
    assert wa.cell(row=r2, column=3).value == 110.0
    assert wa.cell(row=r2, column=5).value == "analyst_override"
    assert wa.cell(row=r2, column=7).value == "Yes"
    assert wa.cell(row=r2, column=8).value == "Share gain in Year 2"
    # The scalar ARI never leaks into the by-year ARI rows.
    assert wa.cell(row=rows["index_ari_target_by_year.1"], column=3).value is None
    assert wa.cell(row=rows["index_ari_target_by_year.1"], column=5).value == "not set"


@pytest.mark.asyncio
async def test_projection_apply_folds_by_year_rows_into_the_list_override() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory

    deal_id, *_ = await _setup(overrides={
        "index_mpi_target_by_year": {"value": [None, 110.0], "note": "old"},
    })
    body = ProjApplyBody(changes=[
        {"key": "index_mpi_target_by_year.3", "new_value": 115, "note": "Post-PIP share gain"},
        {"key": "index_mpi_target_by_year.2", "new_value": 110, "note": "x"},   # unchanged
        {"key": "index_ari_target_by_year.1", "new_value": 98, "note": ""},     # note required → skipped
        {"key": "index_ari_target_by_year.2", "new_value": 0, "note": "zero"},  # not an index → skipped
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert [a["key"] for a in res["applied"]] == ["index_mpi_target_by_year.3"]
    assert sorted(x["reason"] for x in res["skipped"]) == ["non_numeric", "note_required", "unchanged"]

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    fo = json.loads(row[0])
    # Untouched years keep their entry; the list grows to the edited year; the
    # list's legacy note survives in front of the edited year's own note.
    assert fo["index_mpi_target_by_year"] == {
        "value": [None, 110.0, 115.0], "note": "old · Y3: Post-PIP share gain",
    }
    assert "index_ari_target_by_year" not in fo
    assert "index_mpi_target_by_year.3" not in fo  # never a dotted key on the deal


def test_by_year_note_merge_keeps_every_untouched_years_note() -> None:
    from app.api.pl_roundtrip import merge_by_year_notes, split_by_year_notes

    assert merge_by_year_notes(None, {2: "a"}) == "Y2: a"
    assert merge_by_year_notes("", {3: "c", 2: "b"}) == "Y2: b · Y3: c"
    # An edited year replaces only its own segment; the rest stay.
    assert merge_by_year_notes("Y2: keep · Y4: keep4", {3: "new"}) == "Y2: keep · Y3: new · Y4: keep4"
    assert merge_by_year_notes("Y2: old2 · Y3: old3", {2: "new2"}) == "Y2: new2 · Y3: old3"
    # A legacy whole-list note is never lost.
    assert merge_by_year_notes("Share gain", {2: "b"}) == "Share gain · Y2: b"
    assert split_by_year_notes("legacy · Y2: a · Y10: j") == ("legacy", {2: "a", 10: "j"})
    assert split_by_year_notes(None) == (None, {})


@pytest.mark.asyncio
async def test_projection_apply_keeps_each_years_note_and_refuses_out_of_range() -> None:
    """Several by-year rows in one apply: each year keeps ITS note ("Y2: … ·
    Y3: …"), untouched years keep theirs, and a value that is neither index
    points (25–250) nor a ratio (0.25–2.5) is refused with the reason — "5"
    is never read as 5 index points and driven into the model."""
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory
    from app.services.index_methodology import coerce_target_list

    deal_id, *_ = await _setup(overrides={
        "index_mpi_target_by_year": {
            "value": [None, 110.0, None, 120.0], "note": "Y2: share gain · Y4: stabilised",
        },
    })
    body = ProjApplyBody(changes=[
        {"key": "index_mpi_target_by_year.3", "new_value": 115, "note": "Post-PIP"},
        {"key": "index_mpi_target_by_year.4", "new_value": 1.25, "note": "ratio form"},  # 1.25 → 125 pts
        {"key": "index_mpi_target_by_year.1", "new_value": 5, "note": "typo"},           # refused
        {"key": "index_ari_target_by_year.1", "new_value": 300, "note": "typo"},         # refused
        {"key": "index_ari_target_by_year.2", "new_value": 12, "note": "ambiguous"},     # refused
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert [a["key"] for a in res["applied"]] == [
        "index_mpi_target_by_year.3", "index_mpi_target_by_year.4",
    ]
    skipped = {x["key"]: x for x in res["skipped"]}
    assert set(skipped) == {
        "index_mpi_target_by_year.1", "index_ari_target_by_year.1", "index_ari_target_by_year.2",
    }
    for sk in skipped.values():
        assert sk["reason"] == "out_of_range"
        assert "25–250" in sk["detail"] and "0.25–2.5" in sk["detail"]

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    fo = json.loads(row[0])
    mpi = fo["index_mpi_target_by_year"]
    assert mpi["value"] == [None, 110.0, 115.0, 1.25]
    assert mpi["note"] == "Y2: share gain · Y3: Post-PIP · Y4: ratio form"
    # The loader reads the ratio as 125 index points, never as 1.25 points.
    assert coerce_target_list(mpi["value"]) == [None, 110.0, 115.0, 125.0]
    # Nothing was stored for the refused ARI rows.
    assert "index_ari_target_by_year" not in fo


@pytest.mark.asyncio
async def test_projection_export_survives_a_market_data_failure(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture,
) -> None:
    """The Benchmark column is a convenience: when ``load_market_data`` raises
    (an aggregation bug on a malformed benchmark extraction), the export still
    serves the analyst's projections — without the column — and logs."""
    import logging

    import app.api.pl_roundtrip as mod
    from app.database import get_session_factory
    from app.export.pl_workbooks import BENCHMARK_HEADER

    async def _boom(*_a, **_k):
        raise RuntimeError("benchmark aggregation blew up")

    monkeypatch.setattr(mod, "load_market_data", _boom)
    deal_id, *_ = await _setup()
    with caplog.at_level(logging.ERROR, logger="app.api.pl_roundtrip"):
        async with get_session_factory()() as s:
            resp = await mod.export_projections(deal_id=deal_id, session=s, auth=_auth())
    assert resp.status_code == 200
    wp = _sheet(resp.body, "Future P&L")
    header = [c.value for c in wp[1]]
    assert BENCHMARK_HEADER not in header
    assert wp.max_row > 1  # the projections themselves are there
    assert any("market data unavailable" in r.getMessage() for r in caplog.records)


@pytest.mark.asyncio
async def test_projection_by_year_routes_are_tenant_scoped() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import, export_projections
    from app.database import get_session_factory

    deal_id, *_ = await _setup(overrides={"index_mpi_target_by_year": {"value": [120.0], "note": "n"}})
    body = ProjApplyBody(changes=[{"key": "index_mpi_target_by_year.1", "new_value": 125, "note": "n"}])
    async with get_session_factory()() as s:
        with pytest.raises(HTTPException) as exc:
            await export_projections(deal_id=deal_id, session=s, auth=_auth(_OTHER))
        assert exc.value.status_code == 404
        with pytest.raises(HTTPException) as exc2:
            await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth(_OTHER))
        assert exc2.value.status_code == 404
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    assert json.loads(row[0])["index_mpi_target_by_year"] == {"value": [120.0], "note": "n"}



@pytest.mark.asyncio
async def test_historical_apply_skips_a_change_without_a_justification() -> None:
    """FON-74 — the workbook round-trip is the same edit as "Correct this
    value": no justification, no write; the cell is reported as skipped."""
    from app.api.pl_roundtrip import HistApplyBody, apply_historicals_import
    from app.database import get_session_factory

    deal_id, d23, _ = await _setup()
    cid = f"{d23}::rooms_revenue"
    body = HistApplyBody(changes=[{"cell_id": cid, "new_value": 9_950_000, "old_value": 9_810_000}])
    async with get_session_factory()() as s:
        res = await apply_historicals_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert res["applied"] == []
    assert [x["reason"] for x in res["skipped"]] == ["note_required"]
    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT fields FROM extraction_results WHERE document_id = :d AND tenant_id = :t"),
            {"d": str(d23), "t": _TENANT},
        )).first()
    fields = {f["field_name"]: f for f in json.loads(row[0])}
    assert "reviewed" not in fields["rooms_revenue"]
