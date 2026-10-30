"""FON-118 — "Returns by Exit Year" in the exports (live payload, memo PDF, Excel).

Every export reads the same rows the engine persisted. A row the engine did
not price is a typed refusal in each artifact (glyph + reason label in the
memo, ``not available: <reason>`` in the workbook, ``status`` / ``reason`` on
the live payload and an entry in ``model["refusals"]``) — never a 0, 0% or
$0. A payload from a run that predates the schedule says so.
"""

from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite:///./fondok-tests-exports-exit.db")

from fondok_schemas.reasons import REASON_META, REFUSAL_GLYPH, ReasonCode

from app.export.labels import (
    EXIT_SCHEDULE_NO_RETURNS_NOTE,
    EXIT_SCHEDULE_STALE_NOTE,
    EXIT_SCHEDULE_TITLE,
)
from app.export.refusals import NOT_AVAILABLE_PREFIX, collect_refusals
from tests.fixtures.kimpton_export_payload import kimpton_memo, kimpton_model

# Raw engine rows exactly as ``ReturnsEngineOutputExt.exit_schedule`` dumps.
_ENGINE_ROWS = [
    {
        "exit_year": 1, "exit_month": 12, "is_selected": False, "is_peak_levered": False,
        "status": "not_available", "reason": "capital_deploying",
        "reason_detail": "capital still deploying through month 15; a sale before completion is not modelled",
        "exit_noi_basis": "forward_12m", "exit_noi_period_label": "Forward 12-month NOI (Year 2, projected)",
        "terminal_noi": None, "gross_sale_price": None, "selling_costs": None, "transfer_tax": None,
        "loan_payoff": None, "refi_included": False, "net_proceeds": None,
        "levered_irr": None, "unlevered_irr": None, "equity_multiple": None,
    },
    {
        "exit_year": 2, "exit_month": 24, "is_selected": False, "is_peak_levered": True,
        "status": "computed", "reason": None, "reason_detail": None,
        "exit_noi_basis": "forward_12m", "exit_noi_period_label": "Forward 12-month NOI (Year 3, projected)",
        "terminal_noi": 4_344_427.35, "gross_sale_price": 62_063_247.83, "selling_costs": 1_241_264.96,
        "transfer_tax": 0.0, "loan_payoff": 23_142_967.79, "refi_included": False,
        "net_proceeds": 37_679_015.08, "levered_irr": 0.5153, "unlevered_irr": 0.2806,
        "equity_multiple": 2.185,
    },
    {
        "exit_year": 5, "exit_month": 60, "is_selected": True, "is_peak_levered": False,
        "status": "computed", "reason": None, "reason_detail": None,
        "exit_noi_basis": "forward_12m", "exit_noi_period_label": "Forward 12-month NOI (Year 6)",
        "terminal_noi": 5_133_453.81, "gross_sale_price": 73_335_054.46, "selling_costs": 1_466_701.09,
        "transfer_tax": 0.0, "loan_payoff": 22_223_267.50, "refi_included": False,
        "net_proceeds": 49_645_085.87, "levered_irr": 0.3051, "unlevered_irr": 0.1967,
        "equity_multiple": 3.173,
    },
]


def _engine_returns() -> dict:
    return {
        "hold_years": 5, "exit_cap_rate": 0.07, "terminal_noi": 5_133_453.81,
        "gross_sale_price": 73_335_054.46, "selling_costs": 1_466_701.09,
        "net_proceeds": 49_645_085.87, "levered_irr": 0.3051, "unlevered_irr": 0.1967,
        "equity_multiple": 3.173, "year_one_coc": 0.08, "avg_coc": 0.1,
        "exit_schedule": _ENGINE_ROWS, "selected_exit_month": 60,
        "selected_exit_year": 5, "peak_levered_exit_year": 2,
    }


def _live_rows() -> list[dict]:
    from app.export.live_payload import _build_returns

    return _build_returns(_engine_returns())["exit_schedule"]


def _model_without_schedule() -> dict:
    """The fixture payload as a run made BEFORE the schedule existed wrote it."""
    model = kimpton_model()
    ret = dict(model["returns_engine"])
    for key in ("exit_schedule", "selected_exit_month", "selected_exit_year", "peak_levered_exit_year"):
        ret.pop(key, None)
    model["returns_engine"] = ret
    return model


def _model_with_schedule() -> dict:
    model = kimpton_model()
    ret = dict(model["returns_engine"])
    ret["exit_schedule"] = _live_rows()
    ret["selected_exit_month"] = 60
    ret["selected_exit_year"] = 5
    ret["peak_levered_exit_year"] = 2
    model["returns_engine"] = ret
    return model


# ═════════════════ live payload ═════════════════


def test_live_payload_renames_the_rows_and_records_the_row_refusal() -> None:
    from app.export.live_payload import _build_returns

    with collect_refusals() as log:
        out = _build_returns(_engine_returns())

    rows = out["exit_schedule"]
    assert len(rows) == 3
    assert out["selected_exit_month"] == 60 and out["selected_exit_year"] == 5
    assert out["peak_levered_exit_year"] == 2
    dashed, peak, selected = rows
    # The dashed row keeps its status/reason and carries NO numbers at all.
    assert dashed["status"] == "not_available" and dashed["reason"] == "capital_deploying"
    assert dashed["exit_year"] == 1 and dashed["exit_month"] == 12
    for key in ("terminal_noi_usd", "gross_sale_price_usd", "loan_payoff_usd",
                "net_sale_proceeds_usd", "levered_irr", "unlevered_irr", "equity_multiple"):
        assert key not in dashed, key
    # Live renames on the computed rows.
    assert peak["is_peak_levered"] is True and peak["is_selected"] is False
    assert peak["terminal_noi_usd"] == 4_344_427.35
    assert peak["net_sale_proceeds_usd"] == 37_679_015.08
    assert peak["loan_payoff_usd"] == 23_142_967.79
    assert peak["levered_irr"] == 0.5153
    assert selected["is_selected"] is True
    assert selected["equity_multiple"] == 3.173
    assert selected["exit_noi_period_label"] == "Forward 12-month NOI (Year 6)"
    # The row refusal is recorded for model["refusals"].
    codes = [(r.code, r.concept) for r in log]
    assert (ReasonCode.CAPITAL_DEPLOYING, "returns.exit_schedule[0]") in codes


def test_live_payload_without_a_schedule_adds_no_table() -> None:
    from app.export.live_payload import _build_returns

    ret = _engine_returns()
    for key in ("exit_schedule", "selected_exit_month", "selected_exit_year", "peak_levered_exit_year"):
        ret.pop(key)
    out = _build_returns(ret)
    assert "exit_schedule" not in out
    assert "selected_exit_month" not in out


# ═════════════════ memo PDF ═════════════════


def test_memo_renders_the_table_with_glyph_and_reason_on_the_dashed_row() -> None:
    from app.export.memo_pdf import _render_html_body

    memo, model = kimpton_memo(), _model_with_schedule()
    with collect_refusals() as log:
        html_out = _render_html_body(memo, model)

    assert EXIT_SCHEDULE_TITLE in html_out
    assert "Year 1</td>" in html_out
    assert "Year 2 · peak" in html_out
    assert "Year 5 · selected" in html_out
    # The dashed row: the glyph + the reason's label, no number.
    label = REASON_META[ReasonCode.CAPITAL_DEPLOYING]["label"]
    assert f"{REFUSAL_GLYPH} {label}" in html_out
    # The computed rows carry real figures (memo precision: 2 dp %, $ whole).
    assert "51.53%" in html_out and "30.51%" in html_out
    assert "$37,679,015" in html_out and "$49,645,086" in html_out
    assert "2.19x" in html_out and "3.17x" in html_out
    assert "Forward 12-month NOI (Year 3, projected)" in html_out
    codes = {(r.code, r.concept) for r in log}
    assert (ReasonCode.CAPITAL_DEPLOYING, "returns.exit_schedule[0]") in codes
    # Nothing in the table reads as a zero.
    assert "0.00%" not in html_out.split(EXIT_SCHEDULE_TITLE, 1)[1].split("</table>", 1)[0]


def test_memo_without_the_schedule_says_the_run_predates_it() -> None:
    from app.export.memo_pdf import _render_html_body

    memo, model = kimpton_memo(), _model_without_schedule()
    assert "exit_schedule" not in model["returns_engine"]
    with collect_refusals() as log:
        html_out = _render_html_body(memo, model)
    assert EXIT_SCHEDULE_TITLE in html_out
    assert f"{NOT_AVAILABLE_PREFIX}: {EXIT_SCHEDULE_STALE_NOTE}" in html_out
    assert any(r.code is ReasonCode.STALE_RUN and r.concept == "returns.exit_schedule" for r in log)


def test_memo_without_returns_says_the_engine_did_not_run() -> None:
    from app.export.memo_pdf import _render_html_body

    memo, model = kimpton_memo(), kimpton_model()
    model.pop("returns_engine")
    with collect_refusals() as log:
        html_out = _render_html_body(memo, model)
    assert f"{NOT_AVAILABLE_PREFIX}: {EXIT_SCHEDULE_NO_RETURNS_NOTE}" in html_out
    assert any(r.code is ReasonCode.ENGINE_SKIPPED and r.concept == "returns.exit_schedule" for r in log)


# ═════════════════ Excel ═════════════════


def _sheet_rows(path: Path, sheet: str) -> list[list[object]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return [[c.value for c in row] for row in wb[sheet].iter_rows()]
    finally:
        wb.close()


def test_excel_returns_sheet_carries_the_block_with_typed_refusals(tmp_path: Path) -> None:
    from app.export import build_excel

    out = build_excel("x", _model_with_schedule(), tmp_path / "exit.xlsx")
    rows = _sheet_rows(out, "Returns")
    flat = [str(v) for r in rows for v in r if v is not None]
    assert EXIT_SCHEDULE_TITLE.upper() in flat
    assert "Levered IRR" in flat and "Loan payoff" in flat
    by_label = {str(r[0]): r for r in rows if r and r[0] is not None}
    dashed = by_label["Year 1"]
    assert str(dashed[1]).startswith(f"{NOT_AVAILABLE_PREFIX}: capital still deploying")
    assert all(v is None for v in dashed[2:8])  # no 0 anywhere on the dashed row
    peak = by_label["Year 2 · peak"]
    assert peak[1] == 4_344_427.35 and peak[4] == 37_679_015.08
    assert peak[6] == 0.5153 and peak[7] == 2.185
    assert peak[8] == "Forward 12-month NOI (Year 3, projected)"
    selected = by_label["Year 5 · selected"]
    assert selected[6] == 0.3051 and selected[3] == 22_223_267.50
    # Nothing on the sheet reads a dashed figure as a zero.
    assert 0 not in [v for r in rows for v in r[1:8] if isinstance(v, (int, float))]


def test_excel_returns_sheet_states_why_when_the_run_predates_the_schedule(tmp_path: Path) -> None:
    from app.export import build_excel

    out = build_excel("x", _model_without_schedule(), tmp_path / "old.xlsx")
    flat = [str(v) for r in _sheet_rows(out, "Returns") for v in r if v is not None]
    assert EXIT_SCHEDULE_TITLE.upper() in flat
    assert f"{NOT_AVAILABLE_PREFIX}: {EXIT_SCHEDULE_STALE_NOTE}" in flat
