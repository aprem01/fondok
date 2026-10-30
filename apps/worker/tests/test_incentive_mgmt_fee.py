"""Sam #9 — base vs incentive management fee.

The expense engine has always carried ONE management fee: ``mgmt_fee_pct`` ×
total revenue (the base fee). Hotel management agreements usually add an
INCENTIVE fee on top, earned on profit. This pins the opt-in
``incentive_mgmt_fee_pct_of_gop`` assumption:

* unset (the default) → every expense output, NOI and provenance trace is
  identical to before, and ``incentive_mgmt_fee`` is ``None`` — the Kimpton
  goldens cannot move;
* set → each year's incentive fee = max(0, GOP − base fee) × rate, deducted
  between the base fee and NOI (both NOI bases), with a provenance trace;
* a year whose GOP does not cover the base fee (or is negative) earns 0;
* the override routes through ``engine_runner`` into the engine and on to
  returns, and requires an analyst justification like every engine input;
* the Excel round-trip carries the Assumptions row and, only when modelled,
  the Future P&L row.
"""

from __future__ import annotations

import io
import json
import os
from datetime import UTC, datetime
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite:///./fondok.db")
os.environ.setdefault("EVALS_MOCK", "true")

from app.engines.expense import ExpenseEngine, ExpenseEngineInput  # noqa: E402
from app.engines.fb_revenue import FBRevenueOutput, FBRevenueYear  # noqa: E402

KEY = "incentive_mgmt_fee_pct_of_gop"


def _fb_payload(totals: list[float]) -> FBRevenueOutput:
    years = [
        FBRevenueYear(
            year=i + 1,
            rooms_revenue=t * 0.74,
            fb_revenue=t * 0.22,
            other_revenue=t - t * 0.74 - t * 0.22,
            total_revenue=t,
        )
        for i, t in enumerate(totals)
    ]
    return FBRevenueOutput(
        deal_id=uuid4(), years=years, fb_ratio_used=0.22, other_ratio_used=0.04
    )


_TOTALS = [13_600_000.0, 14_200_000.0, 14_900_000.0, 15_400_000.0, 15_900_000.0]


def _input(**overrides) -> ExpenseEngineInput:
    kwargs = dict(
        deal_id=uuid4(),
        revenue=_fb_payload(_TOTALS),
        hotel_type="full",
        mgmt_fee_pct=0.03,
        ffe_reserve_pct=0.04,
        expense_growth=0.035,
        grow_opex_independently=True,
    )
    kwargs.update(overrides)
    return ExpenseEngineInput(**kwargs)


# ── unset → identical ────────────────────────────────────────────────────


def test_unset_incentive_fee_is_byte_identical() -> None:
    payload = _input()
    baseline = ExpenseEngine().run(payload).model_dump(mode="json")
    explicit = ExpenseEngine().run(
        payload.model_copy(update={KEY: None})
    ).model_dump(mode="json")
    assert explicit == baseline
    for y in baseline["years"]:
        assert y["incentive_mgmt_fee"] is None
        # NOI is still exactly gop − base fee − fixed (− reserve), no extra term.
        assert y["noi_institutional"] == (
            y["gop"] - y["mgmt_fee"] - y["fixed_charges"]["total"]
        )
        assert y["noi"] == (
            y["gop"] - y["mgmt_fee"] - y["ffe_reserve"] - y["fixed_charges"]["total"]
        )
    assert not any("incentive" in k for k in baseline["provenance"])
    assert "incentive" not in baseline["provenance"]["years[0].noi"]["formula"]


# ── set → arithmetic per year ────────────────────────────────────────────


def test_set_incentive_fee_arithmetic_per_year() -> None:
    rate = 0.10
    base = ExpenseEngine().run(_input())
    out = ExpenseEngine().run(_input(**{KEY: rate}))
    assert len(out.years) == len(_TOTALS)
    for b, y in zip(base.years, out.years, strict=True):
        # Everything above the fee is unchanged.
        assert y.gop == b.gop
        assert y.mgmt_fee == b.mgmt_fee
        assert y.ffe_reserve == b.ffe_reserve
        assert y.fixed_charges == b.fixed_charges
        expected = max(0.0, y.gop - y.mgmt_fee) * rate
        assert y.gop - y.mgmt_fee > 0  # a healthy year — the fee is positive
        assert y.incentive_mgmt_fee == pytest.approx(expected)
        assert y.noi_institutional == pytest.approx(b.noi_institutional - expected)
        assert y.noi == pytest.approx(b.noi - expected)

    # Provenance: the fee carries its own trace, and NOI names it.
    trace = out.provenance["years[1].incentive_mgmt_fee"]
    assert trace.value == pytest.approx(out.years[1].incentive_mgmt_fee)
    assert trace.assumption_key == KEY
    assert trace.state == "calculated"
    assert [i.name for i in trace.inputs] == [
        "gop", "management_fee", "incentive_mgmt_fee_pct_of_gop",
    ]
    noi_trace = out.provenance["years[1].noi"]
    assert "incentive_management_fee" in noi_trace.formula
    assert any(
        i.traces_to == "years[1].incentive_mgmt_fee" for i in noi_trace.inputs
    )


def test_zero_rate_emits_zero_fee_and_unchanged_noi() -> None:
    base = ExpenseEngine().run(_input())
    out = ExpenseEngine().run(_input(**{KEY: 0.0}))
    for b, y in zip(base.years, out.years, strict=True):
        assert y.incentive_mgmt_fee == 0.0
        assert y.noi == pytest.approx(b.noi)


# ── GOP below the base fee / negative GOP clamps to 0 ─────────────────────


def test_gop_below_base_fee_earns_no_incentive() -> None:
    # A Year-1 T-12 management fee larger than GOP: GOP − base fee < 0.
    out = ExpenseEngine().run(
        _input(**{KEY: 0.20}, t12_actuals={"mgmt_fee": 50_000_000.0})
    )
    y1 = out.years[0]
    assert y1.gop - y1.mgmt_fee < 0
    assert y1.incentive_mgmt_fee == 0.0
    assert y1.noi_institutional == pytest.approx(
        y1.gop - y1.mgmt_fee - y1.fixed_charges.total
    )
    # Later years (base fee back on the revenue ratio) still earn the fee.
    assert out.years[1].incentive_mgmt_fee > 0


def test_negative_gop_clamps_incentive_to_zero() -> None:
    # Departmental + undistributed costs above revenue → GOP < 0 every year.
    out = ExpenseEngine().run(
        _input(
            **{KEY: 0.15},
            overrides={
                "rooms_dept_pct": 0.95,
                "fb_dept_pct": 0.95,
                "other_dept_pct": 0.95,
                "undistributed_pct_revenue": 0.40,
            },
        )
    )
    for y in out.years:
        assert y.gop < 0
        assert y.incentive_mgmt_fee == 0.0
        assert y.noi_institutional == pytest.approx(
            y.gop - y.mgmt_fee - y.fixed_charges.total
        )


# ── engine_runner routing + note rule ────────────────────────────────────


def test_runner_coerces_the_override_value() -> None:
    from app.services.engine_runner import (
        INCENTIVE_MGMT_FEE_KEY,
        _coerce_incentive_mgmt_fee_pct,
    )

    assert INCENTIVE_MGMT_FEE_KEY == KEY
    assert _coerce_incentive_mgmt_fee_pct(None) is None
    assert _coerce_incentive_mgmt_fee_pct("") is None
    assert _coerce_incentive_mgmt_fee_pct("abc") is None
    assert _coerce_incentive_mgmt_fee_pct(True) is None
    assert _coerce_incentive_mgmt_fee_pct(float("nan")) is None
    assert _coerce_incentive_mgmt_fee_pct(0.1) == pytest.approx(0.1)
    assert _coerce_incentive_mgmt_fee_pct("0.12") == pytest.approx(0.12)
    assert _coerce_incentive_mgmt_fee_pct(-0.2) == 0.0
    assert _coerce_incentive_mgmt_fee_pct(3.0) == 0.5


def test_incentive_override_requires_a_justification() -> None:
    from app.api.deals import _NOTE_EXEMPT_KEYS, _override_needs_note
    from app.services.engine_runner import _OVERRIDE_NON_ENGINE_KEYS

    assert KEY not in _NOTE_EXEMPT_KEYS
    assert KEY not in _OVERRIDE_NON_ENGINE_KEYS
    assert _override_needs_note(KEY, _OVERRIDE_NON_ENGINE_KEYS) is True


async def _run_deal(overrides: dict) -> dict:
    from app.database import get_session_factory
    from app.migrations import run_startup_migrations
    from app.services.engine_runner import run_all_engines

    await run_startup_migrations()
    factory = get_session_factory()
    async with factory() as session:
        deal_id, tenant_id = str(uuid4()), str(uuid4())
        await session.execute(
            text(
                """
                INSERT INTO deals (id, tenant_id, name, status, keys, purchase_price, field_overrides,
                                   created_at, updated_at)
                VALUES (:id, :tenant, 'Incentive Fee Hotel', 'Draft', 132, 36400000, :ov, :now, :now)
                """
            ),
            {
                "id": deal_id,
                "tenant": tenant_id,
                "ov": json.dumps(overrides),
                "now": datetime.now(UTC),
            },
        )
        await session.commit()
        return await run_all_engines(
            session, deal_id=deal_id, tenant_id=tenant_id, run_id=str(uuid4())
        )


@pytest.mark.asyncio
async def test_persisted_override_flows_to_noi_and_returns() -> None:
    baseline = await _run_deal({})
    flexed = await _run_deal({KEY: {"value": 0.10, "note": "HMA: 10% of GOP after base"}})

    base_years = baseline["expense"]["outputs"]["years"]
    flex_years = flexed["expense"]["outputs"]["years"]
    assert all(y["incentive_mgmt_fee"] is None for y in base_years)
    for b, f in zip(base_years, flex_years, strict=True):
        fee = max(0.0, f["gop"] - f["mgmt_fee"]) * 0.10
        assert f["incentive_mgmt_fee"] == pytest.approx(fee)
        assert f["noi_institutional"] == pytest.approx(b["noi_institutional"] - fee)
        assert f["noi"] == pytest.approx(b["noi"] - fee)
    # Revenue never moves; returns follow the lower NOI.
    assert (
        flexed["revenue"]["outputs"]["years"] == baseline["revenue"]["outputs"]["years"]
    )
    b_ret, f_ret = baseline["returns"]["outputs"], flexed["returns"]["outputs"]
    assert f_ret["terminal_noi"] < b_ret["terminal_noi"]
    assert f_ret["levered_irr"] < b_ret["levered_irr"]


# ── Excel round-trip ─────────────────────────────────────────────────────


def test_roundtrip_assumptions_row_is_present_and_note_gated() -> None:
    from app.export.pl_workbooks import (
        ASSUMPTION_BY_KEY,
        ASSUMPTIONS,
        build_projections_workbook,
        parse_projections_import,
        projection_values,
    )

    spec = ASSUMPTION_BY_KEY[KEY]
    assert spec.label == "Incentive management fee"
    keys = [a.key for a in ASSUMPTIONS]
    # Sits right after the base management fee.
    assert keys.index(KEY) == keys.index("mgmt_fee_pct") + 1

    rows = [
        {"key": a.key, "label": a.label, "value": None, "unit": a.unit,
         "source": "", "method": a.method, "note_required": True,
         "current_note": None}
        for a in ASSUMPTIONS
    ]
    n, values = projection_values({"expense": {"years": [{"gop": 4_000_000.0}] * 2}})
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
        values=values, assumptions=rows,
    )
    wb = load_workbook(io.BytesIO(data))
    wa = wb["Assumptions"]
    row_of = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    assert wa.cell(row=row_of[KEY], column=3).value is None  # unset → blank

    def _parse(note: str | None):
        wa.cell(row=row_of[KEY], column=3).value = 0.1
        wa.cell(row=row_of[KEY], column=9).value = note
        buf = io.BytesIO()
        wb.save(buf)
        return parse_projections_import(
            buf.getvalue(), deal_id="deal-1", current_values=values,
            current_assumptions={KEY: None}, needs_note=lambda k: True,
        )

    refused = _parse(None)
    assert [r["key"] for r in refused["rejected"]] == [KEY]
    assert refused["rejected"][0]["reason"] == "note_required"
    accepted = _parse("HMA incentive: 10% of GOP after base")
    assert [(c["key"], c["new_value"]) for c in accepted["changes"]] == [(KEY, 0.1)]


def test_future_pl_sheet_row_only_when_modelled() -> None:
    from app.export.pl_workbooks import build_projections_workbook, projection_values

    def _paths(expense_years: list[dict]) -> list[str]:
        _n, values = projection_values({"expense": {"years": expense_years}})
        data = build_projections_workbook(
            deal_id="deal-1", deal_name="X",
            year_headers=["Base year (Year 1)", "Year 2"], values=values,
            assumptions=[],
        )
        ws = load_workbook(io.BytesIO(data))["Future P&L"]
        return [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]

    unset = _paths([{"gop": 4e6, "mgmt_fee": 4e5, "incentive_mgmt_fee": None}] * 2)
    assert "expense.incentive_mgmt_fee" not in unset
    assert "expense.mgmt_fee" in unset
    modelled = _paths([{"gop": 4e6, "mgmt_fee": 4e5, "incentive_mgmt_fee": 3.6e5}] * 2)
    i = modelled.index("expense.incentive_mgmt_fee")
    assert modelled[i - 1] == "expense.mgmt_fee"
