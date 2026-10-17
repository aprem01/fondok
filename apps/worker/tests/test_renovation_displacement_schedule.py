"""Renovation displacement schedule — Sam, 2026-10-15.

"Renovation impacts should be driven by explicit assumptions around timing,
room displacement, and rate adjustments." The schedule is that statement, as
periods in months from the acquisition close × rooms offline × rate discount.

Pinned here:

* OFF by default — no schedule → the revenue engine's output is identical
  (model_dump equality), and the loader's ``displacement_status`` is the
  pre-existing verdict.
* One-period arithmetic per model year (``renovation_displacement_factors``):
  equal 1/12 month weights, occupancy factor = mean open share, ADR factor =
  open-room-weighted (1 − discount); a period in months 13–18 displaces Year 2
  only; overlapping periods add and cap at 100%; a year closed entirely has
  occupancy factor 0 and ADR factor 1.
* Precedence — with both the scalar ``y1_*`` and the schedule set, the
  schedule wins (the scalar is NOT applied on top) and the status says so.
* Provenance — ``renovation_displacement_by_year`` on the output and the
  occupancy / ADR lineage traces name the schedule.
* Loader — the ``renovation_displacement_schedule`` list override is routed
  to the revenue input; the status gains ``on_schedule`` and names periods.
* API — the FON-74 note gate applies to the key (it moves a number) and a
  malformed period is refused with ``renovation_schedule_invalid``.
* Excel round-trip — schedule rows appear on the Assumptions sheet and the
  Apply folds them back into the list with per-period notes.
"""

from __future__ import annotations

import json
import os
from datetime import UTC, datetime
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from sqlalchemy import text

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite:///./fondok.db")

from app.engines.revenue import RevenueEngine, renovation_displacement_factors  # noqa: E402
from fondok_schemas.underwriting import (  # noqa: E402
    RenovationDisplacementPeriod,
    RevenueEngineInput,
)

_TENANT = "00000000-0000-0000-0000-00000000d15c"
KEYS = 200


_DEAL = UUID("00000000-0000-0000-0000-00000000c0de")


def _input(**overrides) -> RevenueEngineInput:
    kwargs = dict(
        deal_id=_DEAL,
        keys=KEYS,
        starting_occupancy=0.80,
        starting_adr=250.0,
        occupancy_growth=0.0,
        adr_growth=0.0,
        fb_revenue_per_occupied_room=0.0,
        other_revenue_pct_of_rooms=0.0,
        hold_years=3,
    )
    kwargs.update(overrides)
    return RevenueEngineInput(**kwargs)


def _period(start: int, end: int, offline: float = 0.0, discount: float = 0.0, note: str | None = None):
    return RenovationDisplacementPeriod(
        start_month=start, end_month=end, rooms_offline_pct=offline,
        rate_discount_pct=discount, note=note,
    )


# ───────────────────────── 1. OFF by default ─────────────────────────


def test_no_schedule_is_byte_identical() -> None:
    a = RevenueEngine().run(_input())
    b = RevenueEngine().run(_input(renovation_displacement_schedule=[]))
    assert a.model_dump() == b.model_dump()
    assert a.renovation_displacement_by_year == []
    # The new output field is an EMPTY list — no new leaf for the golden differ.
    assert "renovation_displacement_by_year" in a.model_dump()
    assert a.model_dump()["renovation_displacement_by_year"] == []


def test_schedule_with_a_scalar_set_but_schedule_empty_keeps_the_scalar_path() -> None:
    """An empty schedule must not switch the scalar lever off."""
    out = RevenueEngine().run(
        _input(y1_occupancy_displacement_pct=0.10, renovation_displacement_schedule=[])
    )
    assert out.years[0].occupancy == pytest.approx(0.80 * 0.90)


# ───────────────────── 2. Arithmetic per model year ─────────────────────


def test_one_period_six_months_quarter_offline_ten_percent_discount() -> None:
    sched = [_period(1, 6, 0.25, 0.10)]
    by_year = renovation_displacement_factors(sched, 3)
    assert sorted(by_year) == [1]
    y1 = by_year[1]
    # open share: 6 × 0.75 + 6 × 1.0 = 10.5 → ÷ 12
    assert y1.occupancy_factor == pytest.approx(10.5 / 12)
    assert y1.rooms_offline_share == pytest.approx(1 - 10.5 / 12)
    # ADR: (6 × 0.75 × 0.9 + 6 × 1.0) ÷ 10.5
    assert y1.adr_factor == pytest.approx((6 * 0.75 * 0.9 + 6) / 10.5)
    assert y1.months_in_period == 6 and y1.entries == [0]

    out = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    y = out.years[0]
    assert y.occupancy == pytest.approx(0.80 * 10.5 / 12)
    assert y.adr == pytest.approx(250.0 * (6 * 0.75 * 0.9 + 6) / 10.5)
    # Rooms revenue = Σ twelve monthly (open share × discounted rate) terms.
    monthly = KEYS * 365 / 12 * 0.80 * 250.0
    expected = monthly * (6 * 0.75 * 0.9 + 6 * 1.0)
    assert y.rooms_revenue == pytest.approx(expected)
    # RevPAR still reconciles on the full inventory.
    assert y.revpar == pytest.approx(y.rooms_revenue / (KEYS * 365))
    # Year 2 compounds from the UN-displaced baseline — no cascade.
    assert out.years[1].occupancy == pytest.approx(0.80)
    assert out.years[1].adr == pytest.approx(250.0)
    assert [d.year for d in out.renovation_displacement_by_year] == [1]


def test_period_in_months_13_to_18_displaces_year_two_only() -> None:
    sched = [_period(13, 18, 0.50, 0.0)]
    out = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    assert out.years[0].occupancy == pytest.approx(0.80)
    assert out.years[1].occupancy == pytest.approx(0.80 * (6 * 0.5 + 6) / 12)
    assert out.years[1].adr == pytest.approx(250.0)
    assert out.years[2].occupancy == pytest.approx(0.80)
    assert [d.year for d in out.renovation_displacement_by_year] == [2]
    tr = out.provenance["years[1].occupancy"]
    assert "schedule_occupancy_factor" in tr.formula
    assert "Year 2" in (tr.note or "") and "months 13–18" in (tr.note or "")


def test_period_spanning_a_year_boundary_is_split_by_months() -> None:
    sched = [_period(10, 15, 1.0, 0.0)]  # 3 months of Y1, 3 months of Y2, fully closed
    by_year = renovation_displacement_factors(sched, 3)
    assert by_year[1].occupancy_factor == pytest.approx(9 / 12)
    assert by_year[2].occupancy_factor == pytest.approx(9 / 12)
    assert by_year[1].months_in_period == 3 and by_year[2].months_in_period == 3
    # Closed rooms carry no rate — the ADR factor stays 1.0 when the discount is 0.
    assert by_year[1].adr_factor == pytest.approx(1.0)


def test_overlapping_periods_add_and_cap_at_full_closure() -> None:
    sched = [_period(1, 12, 0.60, 0.05), _period(1, 12, 0.60, 0.10)]
    by_year = renovation_displacement_factors(sched, 1)
    y1 = by_year[1]
    assert y1.occupancy_factor == pytest.approx(0.0)
    assert y1.rooms_offline_share == pytest.approx(1.0)
    assert y1.adr_factor == pytest.approx(1.0)  # nothing sold → no rate to discount
    assert y1.entries == [0, 1]
    out = RevenueEngine().run(_input(renovation_displacement_schedule=sched, hold_years=2))
    assert out.years[0].rooms_revenue == pytest.approx(0.0)
    assert out.years[1].occupancy == pytest.approx(0.80)


def test_rate_discount_only_period_keeps_occupancy_whole() -> None:
    sched = [_period(1, 12, 0.0, 0.20)]
    out = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    assert out.years[0].occupancy == pytest.approx(0.80)
    assert out.years[0].adr == pytest.approx(250.0 * 0.80)


def test_periods_past_the_hold_do_nothing() -> None:
    sched = [_period(100, 110, 0.5, 0.5)]
    out = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    base = RevenueEngine().run(_input())
    assert [y.rooms_revenue for y in out.years] == [y.rooms_revenue for y in base.years]
    assert out.renovation_displacement_by_year == []


def test_schedule_applies_inside_segments() -> None:
    from fondok_schemas.underwriting import RevenueSegment

    segs = [
        RevenueSegment(name="transient_bar", adr=300.0, mix_pct=0.5, channel_cost_pct=0.0),
        RevenueSegment(name="group", adr=200.0, mix_pct=0.5, channel_cost_pct=0.0),
    ]
    sched = [_period(1, 12, 0.0, 0.10)]
    out = RevenueEngine().run(_input(segments=segs, renovation_displacement_schedule=sched))
    y1 = out.years[0]
    assert [s.adr for s in y1.segment_breakdown] == pytest.approx([270.0, 180.0])


# ───────────────────── 3. Precedence — schedule wins ─────────────────────


def test_schedule_wins_over_the_scalar_y1_percentages() -> None:
    sched = [_period(1, 6, 0.25, 0.10)]
    with_scalar = RevenueEngine().run(
        _input(
            y1_occupancy_displacement_pct=0.30,
            y1_adr_displacement_pct=0.30,
            renovation_displacement_schedule=sched,
        )
    )
    without_scalar = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    assert with_scalar.model_dump() == without_scalar.model_dump()
    tr = with_scalar.provenance["years[0].occupancy"]
    assert tr.formula == "occupancy = starting_occupancy × schedule_occupancy_factor"
    assert "y1_occupancy_displacement_pct" not in tr.formula
    assert "not applied on top" in (tr.note or "")
    by = {i.name: i for i in tr.inputs}
    assert by["schedule_occupancy_factor"].assumption_key == "renovation_displacement_schedule"


def test_schedule_wins_over_the_structured_pip_closure() -> None:
    from fondok_schemas.underwriting import PIPDisplacement

    pip = PIPDisplacement(closure_strategy="rolling", pct_rooms_offline_by_month=[0.5] * 12)
    sched = [_period(1, 6, 0.25, 0.10)]
    a = RevenueEngine().run(_input(pip_displacement=pip, renovation_displacement_schedule=sched))
    b = RevenueEngine().run(_input(renovation_displacement_schedule=sched))
    assert a.model_dump() == b.model_dump()


def test_validation_rejects_bad_periods() -> None:
    from pydantic import ValidationError

    with pytest.raises(ValidationError):
        _period(6, 1)
    with pytest.raises(ValidationError):
        _period(1, 6, offline=1.5)
    with pytest.raises(ValidationError):
        _period(0, 6)


# ───────────────────────── 4. Loader + status ─────────────────────────


@pytest.fixture(autouse=True)
async def _reset_db() -> None:
    from app.database import get_session_factory
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    async with get_session_factory()() as session:
        for tbl in ("extraction_results", "documents", "deals"):
            try:
                await session.execute(text(f"DELETE FROM {tbl} WHERE tenant_id = :t"), {"t": _TENANT})
            except Exception:  # noqa: BLE001
                pass
        await session.commit()
    yield


async def _insert_deal(deal_id: UUID, *, field_overrides: dict | None = None) -> None:
    from app.database import get_session_factory

    async with get_session_factory()() as session:
        await session.execute(
            text(
                "INSERT INTO deals (id, tenant_id, name, status, ai_confidence, keys, "
                "purchase_price, field_overrides, created_at, updated_at) VALUES "
                "(:id, :tenant, 'Schedule deal', 'Underwriting', 0.0, 132, 36400000, :fo, :ts, :ts)"
            ),
            {"id": str(deal_id), "tenant": _TENANT, "fo": json.dumps(field_overrides or {}),
             "ts": datetime.now(UTC)},
        )
        await session.commit()


async def _base_for(deal_id: UUID) -> dict:
    from app.database import get_session_factory
    from app.services.engine_runner import _load_engine_inputs

    async with get_session_factory()() as session:
        return await _load_engine_inputs(session, str(deal_id), tenant_id=_TENANT)


_SCHED = [
    {"start_month": 1, "end_month": 6, "rooms_offline_pct": 0.25, "rate_discount_pct": 0.10,
     "note": "Tower A guestrooms"},
]


@pytest.mark.asyncio
async def test_loader_without_a_schedule_keeps_the_existing_status() -> None:
    deal_id = uuid4()
    await _insert_deal(deal_id)
    base = await _base_for(deal_id)
    assert "renovation_displacement_schedule" not in base
    st = base["__displacement_status__"]
    assert st["state"] == "off_no_assumption"
    assert st["schedule"] == [] and st["schedule_by_year"] == []


@pytest.mark.asyncio
async def test_loader_routes_the_schedule_and_reports_on_schedule() -> None:
    from app.services.engine_runner import (
        DISPLACEMENT_ON_SCHEDULE,
        SOURCE_ANALYST_OVERRIDE,
        _build_input_for,
    )

    deal_id = uuid4()
    await _insert_deal(
        deal_id,
        field_overrides={
            "renovation_displacement_schedule": {"value": _SCHED, "note": "GC phasing plan"},
        },
    )
    base = await _base_for(deal_id)
    assert base["renovation_displacement_schedule"] == _SCHED
    assert base["__sources__"]["renovation_displacement_schedule"] == SOURCE_ANALYST_OVERRIDE
    st = base["__displacement_status__"]
    assert st["active"] is True
    assert st["state"] == DISPLACEMENT_ON_SCHEDULE == "on_schedule"
    assert "months 1–6: 25% rooms offline, 10% rate discount" in st["detail"]
    assert "Year 1: occupancy −12.5%, ADR −4.3%" in st["detail"]
    assert st["schedule"][0]["start_month"] == 1
    assert st["schedule_by_year"][0]["occupancy_factor"] == pytest.approx(0.875)
    # The scalar levers were never set, so the status does not claim to override them.
    assert "takes precedence" not in st["detail"]
    # The revenue input carries the typed periods.
    payload = _build_input_for("revenue", str(deal_id), base, {})
    assert [p.start_month for p in payload.renovation_displacement_schedule] == [1]


@pytest.mark.asyncio
async def test_loader_status_says_the_schedule_wins_over_the_scalar() -> None:
    deal_id = uuid4()
    await _insert_deal(
        deal_id,
        field_overrides={
            "renovation_displacement_schedule": {"value": _SCHED, "note": "GC phasing plan"},
            "y1_occupancy_displacement_pct": {"value": 0.15, "note": "old flat haircut"},
            "y1_adr_displacement_pct": {"value": 0.08, "note": "old flat haircut"},
        },
    )
    base = await _base_for(deal_id)
    st = base["__displacement_status__"]
    assert st["state"] == "on_schedule"
    assert "The renovation schedule takes precedence" in st["detail"]
    assert "occupancy −15%, ADR −8%" in st["detail"]
    assert "not applied in addition" in st["detail"]
    # Client prose: no shorthand (Y1 / PIP) and no raw ×0.xxx factors.
    assert "Y1" not in st["detail"] and "PIP" not in st["detail"]
    assert "×" not in st["detail"]
    # The scalars are still reported (they are what the analyst typed) ...
    assert st["y1_occupancy_displacement_pct"] == pytest.approx(0.15)
    # ... but the engine runs the schedule, not the scalar, for Year 1.
    from app.services.engine_runner import _build_input_for

    out = RevenueEngine().run(_build_input_for("revenue", str(deal_id), base, {}))
    assert out.years[0].occupancy == pytest.approx(base["starting_occupancy"] * 0.875)


@pytest.mark.asyncio
async def test_loader_drops_a_malformed_period_and_keeps_the_rest() -> None:
    deal_id = uuid4()
    await _insert_deal(
        deal_id,
        field_overrides={
            "renovation_displacement_schedule": {
                "value": [{"start_month": 9, "end_month": 2}, _SCHED[0]],
                "note": "x",
            },
        },
    )
    base = await _base_for(deal_id)
    st = base["__displacement_status__"]
    assert st["state"] == "on_schedule"
    assert len(st["schedule"]) == 1 and st["schedule"][0]["start_month"] == 1


# ───────────────────────── 5. API note + shape gate ─────────────────────────


def _auth() -> object:
    from app.auth.context import AuthContext

    return AuthContext(
        tenant_id=UUID(_TENANT), user_id="user_analyst", role="member",
        source="jwt", org_id=None, email="analyst@fondok.test",
    )


async def _patch(deal_id: UUID, overrides: dict) -> None:
    from app.api.deals import UpdateDealBody, update_deal
    from app.database import get_session_factory

    async with get_session_factory()() as s:
        await update_deal(
            deal_id=deal_id, body=UpdateDealBody(field_overrides=overrides),
            session=s, tenant_id=UUID(_TENANT),
        )


@pytest.mark.asyncio
async def test_api_requires_a_note_and_refuses_a_malformed_period() -> None:
    from app.api.deals import _NOTE_EXEMPT_KEYS, _override_needs_note
    from app.services.engine_runner import _OVERRIDE_NON_ENGINE_KEYS

    assert "renovation_displacement_schedule" not in _NOTE_EXEMPT_KEYS
    assert _override_needs_note("renovation_displacement_schedule", _OVERRIDE_NON_ENGINE_KEYS)

    deal_id = uuid4()
    await _insert_deal(deal_id)
    with pytest.raises(HTTPException) as exc:
        await _patch(deal_id, {"renovation_displacement_schedule": _SCHED})
    assert exc.value.status_code == 422
    assert exc.value.detail["code"] == "override_note_required"

    with pytest.raises(HTTPException) as exc:
        await _patch(deal_id, {"renovation_displacement_schedule": {
            "value": [{"start_month": 6, "end_month": 1}], "note": "oops"}})
    assert exc.value.status_code == 422
    assert exc.value.detail["code"] == "renovation_schedule_invalid"
    assert "end_month" in exc.value.detail["problems"]["renovation_displacement_schedule"]

    # Well-formed + note → stored as entered.
    await _patch(deal_id, {"renovation_displacement_schedule": {"value": _SCHED, "note": "GC plan"}})
    from app.database import get_session_factory

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    assert json.loads(row[0])["renovation_displacement_schedule"] == {"value": _SCHED, "note": "GC plan"}

    # Clearing (null) is a revert — no note, no gate.
    await _patch(deal_id, {})
    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    assert "renovation_displacement_schedule" not in json.loads(row[0])


# ───────────────────────── 6. Excel round-trip ─────────────────────────


@pytest.mark.asyncio
async def test_excel_export_carries_schedule_rows_plus_one_blank_period() -> None:
    import io

    from openpyxl import load_workbook

    from app.api.pl_roundtrip import export_projections
    from app.database import get_session_factory

    deal_id = uuid4()
    await _insert_deal(
        deal_id,
        field_overrides={"renovation_displacement_schedule": {"value": _SCHED, "note": "GC plan"}},
    )
    async with get_session_factory()() as s:
        resp = await export_projections(deal_id=deal_id, session=s, auth=_auth())
    wa = load_workbook(io.BytesIO(resp.body))["Assumptions"]
    rows = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    sched_keys = [k for k in rows if k.startswith("renovation_displacement_schedule.")]
    assert sched_keys == [
        "renovation_displacement_schedule.1.start_month",
        "renovation_displacement_schedule.1.end_month",
        "renovation_displacement_schedule.1.rooms_offline_pct",
        "renovation_displacement_schedule.1.rate_discount_pct",
        "renovation_displacement_schedule.2.start_month",
        "renovation_displacement_schedule.2.end_month",
        "renovation_displacement_schedule.2.rooms_offline_pct",
        "renovation_displacement_schedule.2.rate_discount_pct",
    ]
    r = rows["renovation_displacement_schedule.1.rooms_offline_pct"]
    assert wa.cell(row=r, column=3).value == 0.25
    assert wa.cell(row=r, column=5).value == "analyst_override"
    assert wa.cell(row=r, column=7).value == "Yes"          # note required
    assert wa.cell(row=r, column=8).value == "Tower A guestrooms"  # the period's own note
    blank = rows["renovation_displacement_schedule.2.start_month"]
    assert wa.cell(row=blank, column=3).value is None
    assert wa.cell(row=blank, column=5).value == "not set"


@pytest.mark.asyncio
async def test_excel_apply_folds_period_rows_into_the_list_with_notes() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory

    deal_id = uuid4()
    await _insert_deal(
        deal_id,
        field_overrides={"renovation_displacement_schedule": {"value": _SCHED, "note": "GC plan"}},
    )
    body = ProjApplyBody(changes=[
        # edit period 1's discount
        {"key": "renovation_displacement_schedule.1.rate_discount_pct", "new_value": 0.15,
         "old_value": 0.10, "note": "Deeper disruption pricing"},
        # add period 2 (both months present)
        {"key": "renovation_displacement_schedule.2.start_month", "new_value": 13, "note": "Tower B"},
        {"key": "renovation_displacement_schedule.2.end_month", "new_value": 15, "note": "Tower B"},
        {"key": "renovation_displacement_schedule.2.rooms_offline_pct", "new_value": 0.5, "note": "Tower B"},
        # period 3 without both months → incomplete
        {"key": "renovation_displacement_schedule.3.start_month", "new_value": 20, "note": "x"},
        # note required → skipped
        {"key": "renovation_displacement_schedule.1.rooms_offline_pct", "new_value": 0.30, "note": ""},
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert sorted(a["key"] for a in res["applied"]) == [
        "renovation_displacement_schedule.1.rate_discount_pct",
        "renovation_displacement_schedule.2.end_month",
        "renovation_displacement_schedule.2.rooms_offline_pct",
        "renovation_displacement_schedule.2.start_month",
    ]
    assert sorted(x["reason"] for x in res["skipped"]) == ["note_required", "period_incomplete"]

    async with get_session_factory()() as s:
        row = (await s.execute(
            text("SELECT field_overrides FROM deals WHERE id = :id AND tenant_id = :t"),
            {"id": str(deal_id), "t": _TENANT},
        )).first()
    fo = json.loads(row[0])
    saved = fo["renovation_displacement_schedule"]
    assert saved["value"] == [
        {"start_month": 1, "end_month": 6, "rooms_offline_pct": 0.25, "rate_discount_pct": 0.15,
         "note": "Deeper disruption pricing"},
        {"start_month": 13, "end_month": 15, "rooms_offline_pct": 0.5, "rate_discount_pct": 0.0,
         "note": "Tower B"},
    ]
    assert saved["note"] == "GC plan · P1: Deeper disruption pricing · P2: Tower B"
    assert not any(k.startswith("renovation_displacement_schedule.") for k in fo)

    # The loader now runs on both periods: Year 1 and Year 2 displaced.
    base = await _base_for(deal_id)
    st = base["__displacement_status__"]
    assert [y["year"] for y in st["schedule_by_year"]] == [1, 2]


@pytest.mark.asyncio
async def test_excel_apply_refuses_a_period_the_engine_would_reject() -> None:
    from app.api.pl_roundtrip import ProjApplyBody, apply_projections_import
    from app.database import get_session_factory

    deal_id = uuid4()
    await _insert_deal(deal_id)
    body = ProjApplyBody(changes=[
        {"key": "renovation_displacement_schedule.1.start_month", "new_value": 9, "note": "x"},
        {"key": "renovation_displacement_schedule.1.end_month", "new_value": 2, "note": "x"},
    ])
    async with get_session_factory()() as s:
        res = await apply_projections_import(deal_id=deal_id, body=body, session=s, auth=_auth())
    assert res["applied"] == []
    assert {x["reason"] for x in res["skipped"]} == {"out_of_range"}
    assert "end_month" in res["skipped"][0]["detail"]
