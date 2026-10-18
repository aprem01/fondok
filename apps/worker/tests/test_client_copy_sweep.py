"""Client-readiness sweep (2026-10-17) — what an institutional client sees.

Locks the worker side of the sweep whose brief was "a product a committee
member can open without seeing anything unfinished": no developer jargon,
raw ids or field paths in copy, no placeholder / stub data, no hotel name
from the demo fixture under a real deal id, no infrastructure on a public
probe, and the IC Memo's two slow reads served from a self-invalidating
cache.
"""

from __future__ import annotations

import asyncio
import json
import os
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import pytest
from sqlalchemy import text

_TMP_DB = Path(tempfile.gettempdir()) / "fondok-tests-client-copy-sweep.db"
if _TMP_DB.exists():
    _TMP_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_TMP_DB}"
os.environ.setdefault("ALLOW_TENANT_HEADER_WITHOUT_JWT", "1")
os.environ.setdefault("EVALS_MOCK", "true")

_TENANT = "aaaaaaaa-0000-0000-0000-00000000c0ff"


@pytest.fixture(autouse=True)
async def _migrated() -> None:
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    yield


def _client():
    from httpx import ASGITransport, AsyncClient

    from app.main import app

    return AsyncClient(transport=ASGITransport(app=app), base_url="http://t")


def _h(tenant: str = _TENANT) -> dict[str, str]:
    return {"X-Tenant-Id": tenant}


async def _seed_deal(*, tenant: str = _TENANT, **cols) -> str:
    from app.database import get_session_factory

    deal_id = str(uuid4())
    row = {
        "id": deal_id, "tenant": tenant, "name": "Sweep Hotel", "keys": 150,
        "purchase_price": 40_000_000, "fo": json.dumps(cols.pop("field_overrides", {})),
        "ts": datetime.now(UTC),
    }
    async with get_session_factory()() as session:
        await session.execute(
            text(
                "INSERT INTO deals (id, tenant_id, name, status, ai_confidence, keys, "
                "purchase_price, field_overrides, created_at, updated_at) VALUES "
                "(:id, :tenant, :name, 'Underwriting', 0.0, :keys, :purchase_price, "
                ":fo, :ts, :ts)"
            ),
            row,
        )
        await session.commit()
    return deal_id


# ─────────────────────── stub routes are gone (findings 4, 5) ───────────────────────


@pytest.mark.asyncio
async def test_legacy_stub_routes_are_not_mounted() -> None:
    """``/data-library/*`` (200 placeholder for any id) and ``/model/*``
    (empty ``outputs`` / ``status: queued`` for nothing) are deleted, not
    hidden; nothing in apps/web called them."""
    from fastapi.routing import APIRoute

    from app.main import app

    paths = {r.path for r in app.routes if isinstance(r, APIRoute)}
    assert not any(p.startswith("/data-library") for p in paths), paths
    assert not any(p.startswith("/model/") for p in paths), paths

    async with _client() as c:
        assert (await c.get("/data-library/entries", headers=_h())).status_code == 404
        assert (await c.get("/data-library/entries/anything", headers=_h())).status_code == 404
        r = await c.post(f"/model/{uuid4()}/engines/run", json={"engine": "revenue"}, headers=_h())
        assert r.status_code == 404
        r = await c.post(f"/model/{uuid4()}/scenarios", json={"name": "x"}, headers=_h())
        assert r.status_code == 404


# ─────────────────────── demo fixture never leaks (findings 6, 7, 8) ───────────────────────


@pytest.mark.asyncio
async def test_export_of_a_non_uuid_deal_id_is_404_unless_demo_slug_enabled(monkeypatch) -> None:
    from app.api.export import DEAL_NOT_FOUND_DETAIL
    from app.config import get_settings

    monkeypatch.setenv("DEMO_EXPORT_SLUGS", "")
    get_settings.cache_clear()
    try:
        async with _client() as c:
            for path in ("export/excel", "export/memo.pdf", "export/presentation.pptx"):
                r = await c.get(f"/deals/kimpton-angler-2026/{path}", headers=_h())
                assert r.status_code == 404, (path, r.status_code, r.text[:200])
                assert r.json()["detail"] == DEAL_NOT_FOUND_DETAIL
                assert "kimpton" not in r.text.lower()
                r = await c.get(f"/deals/typo-d-url/{path}", headers=_h())
                assert r.status_code == 404, (path, r.status_code)
            # A UUID this tenant does not own answers the same sentence.
            r = await c.get(f"/deals/{uuid4()}/export/excel", headers=_h())
            assert r.status_code == 404 and r.json()["detail"] == DEAL_NOT_FOUND_DETAIL

        # Explicitly enabled (a demo environment) → the fixture workbook serves.
        monkeypatch.setenv("DEMO_EXPORT_SLUGS", "kimpton-angler-2026, other-demo")
        get_settings.cache_clear()
        assert get_settings().demo_export_slugs() == {"kimpton-angler-2026", "other-demo"}
        async with _client() as c:
            r = await c.get("/deals/kimpton-angler-2026/export/excel", headers=_h())
            assert r.status_code == 200, r.text[:200]
            r = await c.get("/deals/still-not-listed/export/excel", headers=_h())
            assert r.status_code == 404
    finally:
        get_settings.cache_clear()


@pytest.mark.asyncio
async def test_memo_payload_loader_never_builds_the_fixture_for_a_db_backed_call() -> None:
    """``_load_deal_payload`` with a session: a deal that is not on the tenant
    raises (code ``deal_not_found``) instead of falling through to the Kimpton
    fixture; a real deal gets only its own documents (an empty list, never the
    fixture's fictional ``STR_MarketReport_Q1.pdf`` excerpts)."""
    from app.api.deals import MemoInputMissing, _load_deal_payload
    from app.database import get_session_factory

    async with get_session_factory()() as session:
        with pytest.raises(MemoInputMissing) as exc:
            await _load_deal_payload(str(uuid4()), session=session, tenant_id=_TENANT)
        assert exc.value.code == "deal_not_found"
        assert exc.value.message == "Deal not found."
        # A stranger's deal is the same refusal (tenant predicate on the lookup).
        foreign = await _seed_deal(tenant=str(uuid4()))
        with pytest.raises(MemoInputMissing) as exc:
            await _load_deal_payload(foreign, session=session, tenant_id=_TENANT)
        assert exc.value.code == "deal_not_found"

    # The fixture path still exists for the session-less demo / smoke flow.
    payload = await _load_deal_payload("kimpton-angler-2026")
    assert all(d.document_id.startswith("doc-") for d in payload.source_documents)


def test_dossier_field_source_vocabulary_names_no_hotel() -> None:
    from typing import get_args

    from app.dossier.schema import DossierField

    allowed = set(get_args(DossierField.model_fields["source"].annotation))
    assert "kimpton_seed" not in allowed
    assert "seed" in allowed
    schema = json.dumps(DossierField.model_json_schema()).lower()
    assert "kimpton" not in schema


# ─────────────────────── export labels (findings 10, 18, 19) ───────────────────────


def test_line_and_source_labels_are_words_not_slugs() -> None:
    from app.export.labels import (
        COMP_CAP_METHOD_DISPLAY,
        line_label,
        source_display,
    )

    assert line_label("noi") == "NOI"
    assert line_label("gop") == "GOP"
    assert line_label("fnb_revenue") == "F&B Revenue"
    assert line_label("ebitda") == "EBITDA"
    assert line_label("total_revenue") == "Total Revenue"
    assert line_label("some_new_line") == "Some New Line"  # graceful, never a slug
    for code, word in (
        ("deal_row", "Deal record"), ("om_broker", "OM"), ("extraction", "Extracted"),
        ("engine_output", "Model"), ("computed", "Derived"), ("seed", "Seed"),
        ("t12_actual", "T-12"),
    ):
        assert source_display(code) == word, code
    # Every source constant the engine runner emits has a word.
    from app.services import engine_runner as er

    codes = {
        v for k, v in vars(er).items()
        if k.startswith("SOURCE_") and isinstance(v, str)
    }
    unmapped = {c for c in codes if source_display(c) == "Other"}
    assert not unmapped, unmapped
    assert source_display("something_unexpected") == "Other"
    assert source_display(None) == "Other"
    assert COMP_CAP_METHOD_DISPLAY["none"].startswith("Not derived")


def test_memo_pdf_historical_walk_prints_noi_and_gop_in_caps() -> None:
    from app.export.fixtures import kimpton_model
    from app.export.memo_pdf import _render_historical_walk

    html_str = _render_historical_walk(kimpton_model()["historical_baseline"])
    assert "<td>NOI</td>" in html_str and "<td>GOP</td>" in html_str
    assert "Noi" not in html_str and "Gop" not in html_str
    # YoY chips carry the same labels.
    assert "NOI FY2023" in html_str


def test_memo_pdf_op_ratio_source_chip_never_title_cases_a_code() -> None:
    from app.export.memo_pdf import _render_op_ratio_provenance

    html_str = _render_op_ratio_provenance(
        {"lines": [
            {"field": "rooms_dept_expense", "value": 0.27, "source": "deal_row"},
            {"field": "utilities", "value": 0.04, "source": "om_broker"},
            {"field": "insurance", "value": 0.01, "source": "brand_new_code"},
        ]}
    )
    assert ">Deal record<" in html_str and ">OM<" in html_str and ">Other<" in html_str
    assert "Deal_Row" not in html_str and "Om_Broker" not in html_str
    assert "Brand_New_Code" not in html_str


def test_excel_historical_walk_and_provenance_and_comps_use_display_labels(tmp_path: Path) -> None:
    import copy

    from openpyxl import load_workbook

    from app.export.excel import build_excel
    from app.export.fixtures import kimpton_model

    model = copy.deepcopy(kimpton_model())
    # Fewer than 3 comps → method "none": the label is a sentence and the cap
    # cells stay blank (a 0% here would be a placeholder number).
    model["comp_sales"]["derived_cap_rate_method"] = "none"
    model["comp_sales"]["coverage_quality"] = "low"
    model["comp_sales"]["derived_cap_rate_median"] = None
    model["comp_sales"]["derived_cap_rate_weighted"] = None
    model["op_ratio_provenance"] = {
        "lines": [{"field": "utilities", "value": 0.04, "source": "om_broker",
                   "document_id": "doc-1"}]
    }
    out = tmp_path / "sweep.xlsx"
    build_excel("sweep-deal", model, out)
    wb = load_workbook(out, read_only=True)

    ws = wb["Historical Baseline"]
    col_a = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "NOI" in col_a and "Noi" not in col_a and "Gop" not in col_a

    ws = wb["Comparable Sales"]
    rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    assert ws.cell(row=rows["Method"], column=2).value == "Not derived — fewer than 3 comparable sales"
    assert ws.cell(row=rows["Method"], column=4).value == "Low"
    assert ws.cell(row=rows["Median Cap"], column=2).value is None
    assert ws.cell(row=rows["Median Cap"], column=4).value is None

    if "Op-Ratio Provenance" in wb.sheetnames:
        ws = wb["Op-Ratio Provenance"]
        vals = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert "OM" in vals and "Om_Broker" not in vals
    wb.close()


# ─────────────────────── error copy (findings 11 to 17, 22, 23) ───────────────────────


def test_grid_422_copy_is_fixed_sentences() -> None:
    from app.api.analysis import _max_price_grid_422_detail
    from app.engines.max_price_grid import TOO_MANY_CELLS_MESSAGE
    from app.engines.price_solver import NO_TARGET_MESSAGE

    assert _max_price_grid_422_detail(ValueError(NO_TARGET_MESSAGE)) == NO_TARGET_MESSAGE
    assert "25 cells" in _max_price_grid_422_detail(ValueError(TOO_MANY_CELLS_MESSAGE))
    assert _max_price_grid_422_detail(ValueError("Grid axes must be non-empty")).startswith("Each axis")
    generic = _max_price_grid_422_detail(ValueError("cap_axis[3] exceeds _CAP_MAX=0.30"))
    assert generic == "The pricing grid could not be computed for these inputs."


def test_doc_type_labels_cover_every_enum_value() -> None:
    from fondok_schemas import DocType

    from app.api.documents import _DOC_TYPE_LABELS, doc_type_label

    for dt in DocType:
        assert dt.value in _DOC_TYPE_LABELS, dt.value
        assert "_" not in doc_type_label(dt.value)
    assert doc_type_label("OM") == "Offering Memorandum"
    assert doc_type_label(None) == "Unclassified"


@pytest.mark.asyncio
async def test_classification_patch_with_unknown_type_lists_document_types() -> None:
    deal_id = await _seed_deal()
    from app.database import get_session_factory

    doc_id = str(uuid4())
    async with get_session_factory()() as session:
        await session.execute(
            text(
                "INSERT INTO documents (id, deal_id, tenant_id, filename, status, "
                "content_hash, storage_key) VALUES (:id, :deal, :tenant, 'x.pdf', "
                "'EXTRACTED', 'cafe', 'k/x.pdf')"
            ),
            {"id": doc_id, "deal": deal_id, "tenant": _TENANT},
        )
        await session.commit()
    async with _client() as c:
        r = await c.patch(
            f"/deals/{deal_id}/documents/{doc_id}/classification",
            json={"doc_type": "xyz"},
            headers=_h(),
        )
    assert r.status_code == 422, r.text
    detail = r.json()["detail"]
    assert detail.startswith("Unrecognised document type. Choose one of: Offering Memorandum, T-12")
    assert "doc_type" not in detail and "xyz" not in detail


@pytest.mark.asyncio
async def test_document_404_and_backfill_copy_name_no_internals() -> None:
    from app.api.documents import DOCUMENT_NOT_FOUND_DETAIL, SEMANTIC_SEARCH_DISABLED_REASON

    deal_id = await _seed_deal()
    async with _client() as c:
        r = await c.get(f"/deals/{deal_id}/documents/{uuid4()}/download", headers=_h())
        assert r.status_code == 404
        assert r.json()["detail"] == DOCUMENT_NOT_FOUND_DETAIL
        r = await c.post(f"/deals/{deal_id}/backfill_embeddings", headers=_h())
        assert r.status_code == 200, r.text
        reason = r.json()["skipped_reason"]
        assert reason == SEMANTIC_SEARCH_DISABLED_REASON
        for word in ("document_chunks", "SQLite", "pgvector", "VOYAGE"):
            assert word not in reason


def test_client_facing_details_carry_no_jargon() -> None:
    """The constants every rewritten detail resolves to."""
    from app.api.documents import (
        DOCUMENT_NOT_FOUND_DETAIL,
        EXTRACTION_NOT_FOUND_DETAIL,
        SEMANTIC_SEARCH_DISABLED_REASON,
    )
    from app.api.export import DEAL_NOT_FOUND_DETAIL
    from app.api.model import NO_RUN_TO_COMPARE_DETAIL, RUN_NOT_ON_DEAL_DETAIL

    for sentence in (
        DOCUMENT_NOT_FOUND_DETAIL, EXTRACTION_NOT_FOUND_DETAIL,
        SEMANTIC_SEARCH_DISABLED_REASON, DEAL_NOT_FOUND_DETAIL,
        NO_RUN_TO_COMPARE_DETAIL, RUN_NOT_ON_DEAL_DETAIL,
    ):
        assert sentence[0].isupper() and sentence.endswith("."), sentence
        assert "_" not in sentence and "{" not in sentence, sentence


# ─────────────────────── lineage seed label (finding 21) ───────────────────────


def test_lineage_seed_label_uses_display_names() -> None:
    from app.services.lineage import _seed_label

    assert _seed_label("seed", "exit_cap_rate", "exit_cap_rate") == "Seed — Exit cap rate"
    assert _seed_label("deal_row", "purchase_price", "purchase_price") == (
        "Deal record — Purchase price (asking price)"
    )
    assert _seed_label(None, "hold_years", None) == "Seed — Hold years"
    # A registry source without a bespoke word still reads as its label.
    assert _seed_label("cbre_horizons", "revpar_growth", "revpar") == "CBRE benchmark — Revenue per Available Room"


# ─────────────────────── public vs operator health (finding 9) ───────────────────────


@pytest.mark.asyncio
async def test_public_health_is_a_probe_and_admin_health_is_gated() -> None:
    from app.tenants import upsert_tenant

    async with _client() as c:
        body = (await c.get("/health")).json()
        assert {"status", "version", "db", "degraded_reasons"} <= set(body)
        for leaked in ("auth", "raw_store", "semantic_search", "extraction_cache", "read_cache"):
            assert leaked not in body, leaked
        serialized = json.dumps(body).lower()
        assert "bucket" not in serialized and "jwt" not in serialized

        # An unregistered caller and an external tenant are both refused.
        assert (await c.get("/admin/health", headers=_h(str(uuid4())))).status_code == 403
        external = uuid4()
        await upsert_tenant(tenant_id=external, clerk_org_id="org_sweep_external",
                            name="Harbor Point", kind="external", status="active")
        r = await c.get("/admin/health", headers=_h(str(external)))
        assert r.status_code == 403 and "bucket" not in r.text

        internal = uuid4()
        await upsert_tenant(tenant_id=internal, clerk_org_id="org_sweep_internal",
                            name="Fondok Internal", kind="internal", status="active")
        r = await c.get("/admin/health", headers=_h(str(internal)))
        assert r.status_code == 200, r.text
        ops = r.json()
        assert set(ops["auth"]) == {
            "header_tenant_path_enabled", "jwt_required",
            "registered_tenant_required", "rls_tenant_guc",
        }
        assert ops["auth"]["jwt_required"] is False  # header path open in the suite
        assert {"kind", "bucket", "region"} <= set(ops["raw_store"])
        assert "mode" in ops["semantic_search"]
        assert "per_tenant" in ops["extraction_cache"]
        assert {"hits", "misses"} <= set(ops["read_cache"])


# ─────────────────────── scenario compare is a read (finding 2) ───────────────────────


@pytest.mark.asyncio
async def test_scenario_compare_is_available_as_a_get() -> None:
    async with _client() as c:
        r = await c.post("/deals", json={"name": "Compare Hotel", "city": "Denver", "keys": 120},
                         headers=_h())
        assert r.status_code == 201, r.text
        deal_id = r.json()["id"]
        base = (await c.get(f"/deals/{deal_id}/scenarios", headers=_h())).json()
        base_id = base[0]["id"]
        r = await c.post(
            f"/deals/{deal_id}/scenarios",
            json={"name": "downside", "overrides": [{"field_path": "exit_cap_rate", "value": 0.085}]},
            headers=_h(),
        )
        assert r.status_code == 201, r.text
        down_id = r.json()["id"]

        got = await c.get(
            f"/deals/{deal_id}/scenarios/compare",
            params=[("scenario_ids", base_id), ("scenario_ids", down_id)],
            headers=_h(),
        )
        assert got.status_code == 200, got.text
        posted = await c.post(
            f"/deals/{deal_id}/scenarios/compare",
            json={"scenario_ids": [base_id, down_id]},
            headers=_h(),
        )
        assert posted.status_code == 200, posted.text
        g, p = got.json(), posted.json()
        assert [s["scenario_id"] for s in g["scenarios"]] == [s["scenario_id"] for s in p["scenarios"]]
        # Once materialised, the read is idempotent: same run ids both ways.
        assert [s["last_run_id"] for s in g["scenarios"]] == [s["last_run_id"] for s in p["scenarios"]]
        # Cross-tenant ids are still 404 on the GET form.
        r = await c.get(
            f"/deals/{deal_id}/scenarios/compare",
            params=[("scenario_ids", base_id)],
            headers=_h(str(uuid4())),
        )
        assert r.status_code == 404
        # The query form is validated like the body form (1 to 4 ids).
        r = await c.get(f"/deals/{deal_id}/scenarios/compare", headers=_h())
        assert r.status_code == 422


# ─────────────────────── deal read cache (finding 1) ───────────────────────


@pytest.mark.asyncio
async def test_read_cache_serves_repeat_views_and_invalidates_on_every_input_change() -> None:
    from app.database import get_session_factory
    from app.services import read_cache as rc

    deal_id = await _seed_deal()
    calls = {"n": 0}

    async def compute():
        calls["n"] += 1
        return {"n": calls["n"]}

    async with get_session_factory()() as session:
        kw = {"name": "probe", "deal_id": deal_id, "tenant_id": _TENANT, "compute": compute}
        assert await rc.cached_deal_read(session, **kw) == {"n": 1}
        assert await rc.cached_deal_read(session, **kw) == {"n": 1}  # hit
        assert rc.read_cache_stats()["hits"] == 1

        # Another tenant never sees this tenant's entry (and 404s upstream anyway).
        other = dict(kw, tenant_id=str(uuid4()))
        assert await rc.cached_deal_read(session, **other) == {"n": 2}

        # Override write → deal row changed → recomputed.
        await session.execute(
            text("UPDATE deals SET field_overrides = :fo WHERE id = :id AND tenant_id = :tenant"),
            {"fo": json.dumps({"exit_cap_rate": {"value": 0.07, "note": "IC"}}),
             "id": deal_id, "tenant": _TENANT},
        )
        await session.commit()
        assert await rc.cached_deal_read(session, **kw) == {"n": 3}
        assert await rc.cached_deal_read(session, **kw) == {"n": 3}

        # New document / extraction → recomputed.
        doc_id = str(uuid4())
        await session.execute(
            text("INSERT INTO documents (id, deal_id, tenant_id, filename, status, content_hash, "
                 "storage_key) VALUES (:id, :deal, :tenant, 't12.pdf', 'EXTRACTED', 'h', 'k')"),
            {"id": doc_id, "deal": deal_id, "tenant": _TENANT},
        )
        await session.commit()
        assert await rc.cached_deal_read(session, **kw) == {"n": 4}
        await session.execute(
            text("INSERT INTO extraction_results (id, document_id, deal_id, tenant_id, fields, "
                 "agent_version) VALUES (:id, :doc, :deal, :tenant, '[]', 'v')"),
            {"id": str(uuid4()), "doc": doc_id, "deal": deal_id, "tenant": _TENANT},
        )
        await session.commit()
        assert await rc.cached_deal_read(session, **kw) == {"n": 5}

        # Reclassification → recomputed.
        await session.execute(
            text("UPDATE documents SET doc_type = 'OM' WHERE id = :id AND tenant_id = :tenant"),
            {"id": doc_id, "tenant": _TENANT},
        )
        await session.commit()
        assert await rc.cached_deal_read(session, **kw) == {"n": 6}

        # Engine run → recomputed.
        await session.execute(
            text("INSERT INTO engine_outputs (id, deal_id, tenant_id, run_id, engine_name, status, "
                 "inputs, outputs, error, started_at, completed_at, runtime_ms) VALUES "
                 "(:id, :deal, :tenant, :run, 'returns', 'complete', NULL, '{}', NULL, "
                 "'2026-10-17T00:00:00+00:00', '2026-10-17T00:00:05+00:00', 1)"),
            {"id": str(uuid4()), "deal": deal_id, "tenant": _TENANT, "run": str(uuid4())},
        )
        await session.commit()
        assert await rc.cached_deal_read(session, **kw) == {"n": 7}

        # In-place extraction rewrite (field review) → the generation bump.
        assert await rc.cached_deal_read(session, **kw) == {"n": 7}
        rc.bump_generation(tenant_id=_TENANT, deal_id=deal_id)
        assert await rc.cached_deal_read(session, **kw) == {"n": 8}

        # TTL lapse → recomputed; TTL 0 → bypassed entirely.
        from app.config import get_settings

        os.environ["READ_CACHE_TTL_S"] = "0"
        get_settings.cache_clear()
        try:
            assert await rc.cached_deal_read(session, **kw) == {"n": 9}
            assert await rc.cached_deal_read(session, **kw) == {"n": 10}
        finally:
            os.environ.pop("READ_CACHE_TTL_S", None)
            get_settings.cache_clear()


@pytest.mark.asyncio
async def test_read_cache_coalesces_concurrent_identical_reads() -> None:
    from app.database import get_session_factory
    from app.services import read_cache as rc

    deal_id = await _seed_deal()
    calls = {"n": 0}

    async def slow_compute():
        calls["n"] += 1
        await asyncio.sleep(0.05)
        return {"n": calls["n"]}

    async def one():
        async with get_session_factory()() as session:
            return await rc.cached_deal_read(
                session, name="probe", deal_id=deal_id, tenant_id=_TENANT, compute=slow_compute
            )

    results = await asyncio.gather(one(), one(), one())
    assert results == [{"n": 1}] * 3
    assert calls["n"] == 1
    assert rc.read_cache_stats()["coalesced"] == 2


@pytest.mark.asyncio
async def test_read_cache_returns_copies_and_propagates_errors() -> None:
    from pydantic import BaseModel

    from app.database import get_session_factory
    from app.services import read_cache as rc

    class Out(BaseModel):
        items: list[int]

    deal_id = await _seed_deal()

    async def compute():
        return Out(items=[1])

    async def boom():
        raise RuntimeError("engine fell over")

    async with get_session_factory()() as session:
        first = await rc.cached_deal_read(session, name="m", deal_id=deal_id, tenant_id=_TENANT, compute=compute)
        first.items.append(99)  # a caller mutating its copy cannot poison the cache
        second = await rc.cached_deal_read(session, name="m", deal_id=deal_id, tenant_id=_TENANT, compute=compute)
        assert second.items == [1]
        with pytest.raises(RuntimeError):
            await rc.cached_deal_read(session, name="err", deal_id=deal_id, tenant_id=_TENANT, compute=boom)
        # A failure caches nothing: the next call computes again.
        with pytest.raises(RuntimeError):
            await rc.cached_deal_read(session, name="err", deal_id=deal_id, tenant_id=_TENANT, compute=boom)


@pytest.mark.asyncio
async def test_assumption_sources_and_variance_routes_are_served_from_the_cache() -> None:
    """End to end: the two IC-memo reads hit the cache on a repeat view and
    pick up an override written through the API."""
    from app.services import read_cache as rc

    deal_id = await _seed_deal()
    async with _client() as c:
        r1 = await c.get(f"/deals/{deal_id}/assumption_sources", headers=_h())
        assert r1.status_code == 200, r1.text
        r2 = await c.get(f"/deals/{deal_id}/assumption_sources", headers=_h())
        assert r2.json() == r1.json()
        v1 = await c.get(f"/analysis/{deal_id}/variance", headers=_h())
        assert v1.status_code == 200, v1.text
        v2 = await c.get(f"/analysis/{deal_id}/variance", headers=_h())
        assert v2.json() == v1.json()
        stats = rc.read_cache_stats()
        assert stats["hits"] >= 2 and stats["misses"] >= 2, stats

        # A stranger is refused before the cache is consulted.
        assert (await c.get(f"/deals/{deal_id}/assumption_sources", headers=_h(str(uuid4())))).status_code == 404
        assert (await c.get(f"/analysis/{deal_id}/variance", headers=_h(str(uuid4())))).status_code == 404

        # An override through the API is visible on the very next read.
        r = await c.patch(
            f"/deals/{deal_id}/overrides",
            json={"overrides": {"exit_cap_rate": {"value": 0.0925, "note": "committee haircut"}}},
            headers=_h(),
        )
        if r.status_code in (200, 201):
            r3 = await c.get(f"/deals/{deal_id}/assumption_sources", headers=_h())
            assert r3.json()["values"]["exit_cap_rate"] == pytest.approx(0.0925)
            assert r3.json()["sources"]["exit_cap_rate"] == "analyst_override"
