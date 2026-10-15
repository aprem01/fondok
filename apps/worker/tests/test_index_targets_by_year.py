"""FON-95 part 2 — penetration targets BY YEAR drive the projection.

Sam (2026-10-14): "adjust market growth and target penetration assumptions by
year … flow consistently into the property's financial projections."

Pinned here:

* ``resolve_index_by_year`` — the benchmark series per method (STR / CoStar
  comp set grown at the market growth assumptions; market benchmark = CBRE
  Horizons by calendar year), the effective target per year (by-year beats
  the scalar all-years shorthand), the DERIVED RGI, and ``benchmark_unavailable``
  for a year with a target but no benchmark — never an invented figure;
* the revenue engine's per-year drivers: a driven year replaces the compounded
  figure and re-anchors the years after it; with no drivers the output is
  byte-identical (every golden deal);
* the runner seam: nothing moves with the toggle off; with it on the lists
  land on the revenue input, Year 1 is re-pointed (by-year beats scalar), the
  explicit analyst ``starting_*`` is never overwritten;
* the endpoint carries the by-year grid, says why it is disabled, and is
  tenant-scoped;
* the Future P&L Assumptions sheet round-trips ``index_*_target_by_year.<n>``.
"""

from __future__ import annotations

import io
import json
import os
from datetime import UTC, datetime
from typing import Any
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import text

os.environ.setdefault("DATABASE_URL", "sqlite+aiosqlite:///./fondok.db")

from app.services.index_methodology import (  # noqa: E402
    BASE_ADR_BY_YEAR,
    BASE_DRIVERS_BY_YEAR,
    BASE_OCC_BY_YEAR,
    KEY_ADR_GROWTH,
    KEY_ARI_TARGET,
    KEY_ARI_TARGET_BY_YEAR,
    KEY_MPI_TARGET,
    KEY_MPI_TARGET_BY_YEAR,
    KEY_OCC_GROWTH,
    METHOD_COSTAR_COMP_SET,
    METHOD_MARKET_BENCHMARK,
    SOURCE_INDEX_ASSUMPTION,
    build_index_drivers,
    coerce_target_list,
    read_targets_by_year,
    resolve_index_by_year,
    wants_benchmark,
)
from app.services.market_fields import FieldRow  # noqa: E402

STR_DOC = "Subject STR Trend May 2025.xlsx"
MARKET_DOC = "Miami Beach Submarket 2025.pdf"
CBRE_DOC = "CBRE Hotel Horizons Miami Q3 2025.pdf"
STR_BASIS = frozenset({"str_forecast", "str_subject_ttm", "str_comp_set"})


def _row(name: str, value: Any, *, doc: str, doc_type: str, page: int | None = 2, period: str | None = None) -> FieldRow:
    return FieldRow(
        field_name=name, value=value, unit=None, page=page, doc_name=doc,
        doc_id=f"doc-{doc_type.lower()}", extraction_id=f"ext-{doc_type.lower()}",
        doc_type=doc_type, period=period,
    )


STR_ROWS = [
    _row("ttm_performance.subject.occupancy_pct", 84.6, doc=STR_DOC, doc_type="STR_TREND", page=3),
    _row("ttm_performance.subject.adr_usd", 255.0, doc=STR_DOC, doc_type="STR_TREND", page=3),
    _row("ttm_performance.indices.mpi_occupancy_index", 110.0, doc=STR_DOC, doc_type="STR_TREND", page=4),
    _row("ttm_performance.indices.ari_adr_index", 102.0, doc=STR_DOC, doc_type="STR_TREND", page=4),
]
COMP_OCC = 0.846 / 1.10
COMP_ADR = 255.0 / 1.02
MARKET_ROWS = [
    _row("pnl_benchmark.market.occupancy_pct", 0.712, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=5),
    _row("pnl_benchmark.market.adr_usd", 231.0, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=5),
    _row("pnl_benchmark.market.adr_change_2026_forecast", 0.031, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=6),
    _row("pnl_benchmark.market.occupancy_change_2026_forecast", 0.005, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=6),
]
CBRE_ROWS = [
    _row("cbre_horizons.segment_all.2026.occupancy_pct", 73.0, doc=CBRE_DOC, doc_type="CBRE_HORIZONS", page=4, period="forecast"),
    _row("cbre_horizons.segment_all.2026.adr_usd", 240.0, doc=CBRE_DOC, doc_type="CBRE_HORIZONS", page=4, period="forecast"),
    _row("cbre_horizons.segment_all.2026.period", "forecast", doc=CBRE_DOC, doc_type="CBRE_HORIZONS", page=4),
    _row("cbre_horizons.segment_upper_priced.2026.occupancy_pct", 76.0, doc=CBRE_DOC, doc_type="CBRE_HORIZONS", page=5, period="forecast"),
    _row("cbre_horizons.segment_upper_priced.2026.adr_usd", 310.0, doc=CBRE_DOC, doc_type="CBRE_HORIZONS", page=5, period="forecast"),
]


# ─────────────────────────── targets ───────────────────────────


def test_target_lists_coerce_points_ratios_json_and_blanks() -> None:
    assert coerce_target_list([None, 115, "1.05", 0, -3, "x"]) == [None, 115.0, 105.0, None, None, None]
    assert coerce_target_list("[110, null]") == [110.0, None]
    assert coerce_target_list("not json") == []
    assert coerce_target_list(110) == []
    # Out-of-range entries read as UNSET, never as points: "5" used to become
    # MPI 5 → occupancy 3.8% marked ``driven``.
    assert coerce_target_list([5, 12, 300, 0.2, 2.6, 24.9]) == [None] * 6
    assert coerce_target_list([25, 250, 0.25, 2.5]) == [25.0, 250.0, 25.0, 250.0]


def test_index_target_validator_bands_and_reasons() -> None:
    from app.services.index_methodology import (
        INDEX_POINTS_RANGE,
        INDEX_RATIO_RANGE,
        target_list_problems,
        validate_index_target,
    )

    assert INDEX_POINTS_RANGE == (25.0, 250.0) and INDEX_RATIO_RANGE == (0.25, 2.5)
    assert validate_index_target(110) == (110.0, None)
    assert validate_index_target("1.10") == (pytest.approx(110.0), None)
    assert validate_index_target("103.2") == (103.2, None)
    for bad in (5, 12, 300, 0.2, 2.6, 24.99, 250.01, 0, -3):
        pts, why = validate_index_target(bad)
        assert pts is None and why and "25–250" in why and "0.25–2.5" in why, bad
    assert validate_index_target("x") == (None, "not a number")
    assert validate_index_target(None) == (None, "not a number")

    assert target_list_problems([None, 110, "", 1.05]) == []
    assert target_list_problems(None) == []
    problems = target_list_problems([None, 5, 110, "abc"])
    assert len(problems) == 2
    assert problems[0].startswith("Year 2: 5 is outside")
    assert problems[1] == "Year 4: not a number"
    assert target_list_problems(110) == [
        "must be a list with one entry per model year (null = unset)"
    ]


@pytest.mark.asyncio
async def test_patch_refuses_out_of_range_penetration_targets_and_keeps_valid_ones() -> None:
    """``update_deal`` is the PATCH path every override takes: a scalar or
    by-year penetration target outside both bands is a 422
    ``override_out_of_range`` and NOTHING is written; in-range values (points
    or ratios) and cleared entries pass."""
    from fastapi import HTTPException

    from app.api.deals import UpdateDealBody, update_deal
    from app.database import get_session_factory

    async with get_session_factory()() as s:
        deal_id, tenant_id = await _str_deal(s, overrides={})

    async def patch(fo: dict):
        async with get_session_factory()() as s:
            return await update_deal(
                deal_id=UUID(deal_id), body=UpdateDealBody(field_overrides=fo),
                session=s, tenant_id=UUID(tenant_id),
            )

    async def stored() -> dict:
        async with get_session_factory()() as s:
            row = (await s.execute(
                text("SELECT field_overrides FROM deals WHERE id = :id"), {"id": deal_id},
            )).first()
        return json.loads(row[0]) if isinstance(row[0], str) else dict(row[0] or {})

    # Scalar "5" — neither 5 points nor a 5× ratio.
    with pytest.raises(HTTPException) as exc:
        await patch({KEY_MPI_TARGET: {"value": 5, "note": "typo"}})
    assert exc.value.status_code == 422
    assert exc.value.detail["code"] == "override_out_of_range"
    assert exc.value.detail["keys"] == [KEY_MPI_TARGET]
    assert "25–250" in exc.value.detail["problems"][KEY_MPI_TARGET]
    assert await stored() == {}

    # By-year list with one bad year — named by year; a good key beside it is
    # refused too (the PATCH is all-or-nothing).
    with pytest.raises(HTTPException) as exc2:
        await patch({
            KEY_MPI_TARGET_BY_YEAR: {"value": [None, 5, 110], "note": "n"},
            KEY_ARI_TARGET: {"value": 1.05, "note": "fine"},
        })
    assert exc2.value.detail["keys"] == [KEY_MPI_TARGET_BY_YEAR]
    assert exc2.value.detail["problems"][KEY_MPI_TARGET_BY_YEAR].startswith("Year 2: 5 is outside")
    assert await stored() == {}

    # In range: points, a ratio, blanks and nulls in the list.
    ok = {
        KEY_MPI_TARGET: {"value": 110, "note": "share gain"},
        KEY_ARI_TARGET: {"value": 1.05, "note": "rate premium"},
        KEY_MPI_TARGET_BY_YEAR: {"value": [None, 115, "", 1.2], "note": "Y2: a · Y4: b"},
    }
    await patch(ok)
    assert await stored() == ok

    # Clearing a key (it leaves the blob) is a revert, never a range problem.
    await patch({KEY_ARI_TARGET: {"value": 1.05, "note": "rate premium"}})
    assert set(await stored()) == {KEY_ARI_TARGET}


def test_by_year_beats_scalar_and_only_analyst_set_values_count() -> None:
    mpi, ari = read_targets_by_year(
        {KEY_MPI_TARGET_BY_YEAR: {"value": [None, 115], "note": "Year 2 share gain"}, KEY_MPI_TARGET: {"value": 110, "note": "all"}},
        3,
    )
    assert [(t.value, t.source) for t in mpi] == [(110.0, "scalar"), (115.0, "by_year"), (110.0, "scalar")]
    assert mpi[1].note == "Year 2 share gain" and mpi[0].note == "all"
    assert [(t.value, t.source) for t in ari] == [(None, None)] * 3


# ─────────────────────────── benchmark series ───────────────────────────


def test_str_comp_set_series_is_the_ttm_level_grown_at_the_market_growth_assumption() -> None:
    """The benchmark SERIES is shown grown at the market growth assumption
    (document default or override); whether an out-year is DRIVEN is a
    separate question — see the scalar gate test below."""
    r = resolve_index_by_year(
        [*STR_ROWS, *MARKET_ROWS],
        {KEY_MPI_TARGET: {"value": 115, "note": "n"}, KEY_ARI_TARGET: {"value": 100, "note": "n"}},
        hold_years=3, calendar_years=[2025, 2026, 2027],
    )
    y1, y2, y3 = r.years
    assert y1.benchmark_occupancy.value == pytest.approx(COMP_OCC) and y1.benchmark_occupancy.source == "computed"
    # Market occupancy growth default = the market report's 2026 forecast (+0.5%).
    assert y2.benchmark_occupancy.value == pytest.approx(COMP_OCC * 1.005)
    assert y3.benchmark_occupancy.value == pytest.approx(COMP_OCC * 1.005**2)
    assert "grown 2 years at +0.50%" in (y3.benchmark_occupancy.detail or "")
    assert y3.benchmark_adr.value == pytest.approx(COMP_ADR * 1.031**2)
    # The growth row rides along as provenance.
    assert any(ref.field_name == "pnl_benchmark.market.occupancy_change_2026_forecast" for ref in y3.benchmark_occupancy.inputs)
    # Year 1: occupancy = benchmark × MPI ÷ 100; ADR at parity; RGI derived.
    assert y1.occupancy == pytest.approx(COMP_OCC * 1.15) and y1.adr == pytest.approx(COMP_ADR)
    assert y1.rgi_derived == pytest.approx(115.0) and (y1.occupancy_status, y1.adr_status) == ("driven", "driven")
    # Year 2 still shows the scalar target and the derived RGI, but a scalar
    # target with only the DOCUMENT-DEFAULT growth drives Year 1 only.
    assert (y2.mpi_target.value, y2.mpi_target.source) == (115.0, "scalar")
    assert y2.rgi_derived == pytest.approx(115.0)
    assert y2.occupancy is None and y2.adr is None
    assert (y2.occupancy_status, y2.adr_status) == ("scalar_year1_only", "scalar_year1_only")
    assert y2.calendar_year == 2026


def test_scalar_target_drives_year_one_only_unless_the_analyst_asks_for_the_out_years() -> None:
    """Review 2026-10-14: a live deal with only ``index_mpi_target`` and the STR
    toggle on must NOT be re-based in Years 2+ on benchmark × document-default
    growth × target with no analyst action. Out-years are driven only by (a) a
    by-year target for that year or (b) an analyst market-growth OVERRIDE for
    that metric; otherwise the previous behaviour (Year 1 only) holds."""
    scalar = {KEY_MPI_TARGET: {"value": 115, "note": "n"}, KEY_ARI_TARGET: {"value": 100, "note": "n"}}
    r = resolve_index_by_year([*STR_ROWS, *MARKET_ROWS], scalar, hold_years=3)
    for y in r.years[1:]:
        assert y.occupancy is None and y.occupancy_status == "scalar_year1_only"
        assert y.adr is None and y.adr_status == "scalar_year1_only"
        assert "drives Year 1 only" in (y.reason or "") and f"Year {y.year}" in (y.reason or "")
        assert "growth override" in (y.reason or "")  # the comp set is a grown series
    # (b) an occupancy growth OVERRIDE drives the out-year occupancy on the
    # overridden growth — ADR (no override) stays Year-1-only.
    r_ov = resolve_index_by_year([*STR_ROWS, *MARKET_ROWS], {**scalar, KEY_OCC_GROWTH: {"value": 0.02, "note": "g"}}, hold_years=3)
    y2 = r_ov.years[1]
    assert y2.occupancy == pytest.approx(COMP_OCC * 1.02 * 1.15) and y2.occupancy_status == "driven"
    assert y2.adr is None and y2.adr_status == "scalar_year1_only"
    # (a) a by-year target for Year 2 drives it on the document-default growth,
    # and Year 3 (back to the scalar) is Year-1-only again.
    r_by = resolve_index_by_year(
        [*STR_ROWS, *MARKET_ROWS],
        {**scalar, KEY_ARI_TARGET_BY_YEAR: {"value": [None, 102], "note": "y2"}},
        hold_years=3,
    )
    assert r_by.years[1].adr == pytest.approx(COMP_ADR * 1.031 * 1.02) and r_by.years[1].adr_status == "driven"
    assert r_by.years[1].ari_target.source == "by_year"
    assert r_by.years[2].adr_status == "scalar_year1_only"
    # The seam: a scalar-only deal parks Year 1 and leaves Years 2+ to the
    # model's own growth — exactly the pre-by-year numbers.
    base = _base(**{KEY_MPI_TARGET: 115.0, KEY_ARI_TARGET: 100.0})
    sources = {"starting_occupancy": "str_subject_ttm", "starting_adr": "str_subject_ttm"}
    changed = build_index_drivers(
        base, sources, {KEY_MPI_TARGET, KEY_ARI_TARGET},
        rows=[*STR_ROWS, *MARKET_ROWS], str_basis_sources=STR_BASIS, hold_years=3, calendar_years=[],
    )
    assert set(changed) == {"starting_occupancy", "starting_adr"}
    assert base[BASE_OCC_BY_YEAR] == [pytest.approx(COMP_OCC * 1.15), None, None]
    assert base[BASE_ADR_BY_YEAR] == [pytest.approx(COMP_ADR), None, None]
    assert [d["occupancy_status"] for d in base[BASE_DRIVERS_BY_YEAR]] == ["driven", "scalar_year1_only", "scalar_year1_only"]


def test_no_growth_assumption_flags_the_out_years_unavailable_never_invents() -> None:
    r = resolve_index_by_year(STR_ROWS, {KEY_MPI_TARGET_BY_YEAR: {"value": [None, 115, None], "note": "n"}}, hold_years=3)
    y1, y2, y3 = r.years
    assert y1.occupancy is None and y1.occupancy_status == "no_target"
    assert y2.occupancy is None and y2.occupancy_status == "benchmark_unavailable"
    assert "No market occupancy growth assumption" in (y2.reason or "")
    assert y3.occupancy_status == "no_target"
    # An analyst growth override restores the series.
    r2 = resolve_index_by_year(
        STR_ROWS,
        {KEY_MPI_TARGET_BY_YEAR: {"value": [None, 115, None], "note": "n"}, KEY_OCC_GROWTH: {"value": 0.02, "note": "g"}},
        hold_years=3,
    )
    assert r2.years[1].occupancy == pytest.approx(COMP_OCC * 1.02 * 1.15)
    assert "your override" in (r2.years[1].benchmark_occupancy.detail or "")


def test_market_benchmark_series_is_cbre_horizons_by_calendar_year() -> None:
    r = resolve_index_by_year(
        [*STR_ROWS, *MARKET_ROWS, *CBRE_ROWS],
        {"index_methodology": METHOD_MARKET_BENCHMARK, KEY_MPI_TARGET_BY_YEAR: {"value": [110, 110, 110], "note": "n"}},
        hold_years=3, calendar_years=[2025, 2026, 2027],
    )
    y1, y2, y3 = r.years
    # Year 1 anchors on the reported level (as the scalar path does).
    assert y1.benchmark_occupancy.value == pytest.approx(0.712) and y1.occupancy == pytest.approx(0.712 * 1.10)
    # Year 2 = CBRE All Hotels 2026 forecast, document + page.
    assert y2.benchmark_occupancy.value == pytest.approx(0.73)
    assert y2.benchmark_occupancy.inputs[0].doc_name == CBRE_DOC and y2.benchmark_occupancy.inputs[0].page == 4
    assert "CBRE Horizons All Hotels 2026 forecast" in (y2.benchmark_occupancy.detail or "")
    assert y2.occupancy == pytest.approx(0.73 * 1.10)
    # Year 3: the report carries no 2027 → unavailable, said so.
    assert y3.occupancy is None and y3.occupancy_status == "benchmark_unavailable"
    assert "2027" in (y3.reason or "")
    # The scalar gate applies here too: the all-years target alone drives Year 1
    # only, and the reason does not suggest a growth override (CBRE is not grown).
    r_scalar = resolve_index_by_year(
        [*STR_ROWS, *MARKET_ROWS, *CBRE_ROWS],
        {"index_methodology": METHOD_MARKET_BENCHMARK, KEY_MPI_TARGET: {"value": 110, "note": "n"}},
        hold_years=3, calendar_years=[2025, 2026, 2027],
    )
    assert r_scalar.years[0].occupancy == pytest.approx(0.712 * 1.10)
    assert r_scalar.years[1].occupancy_status == "scalar_year1_only" and r_scalar.years[1].benchmark_occupancy.value == pytest.approx(0.73)
    assert "growth override" not in (r_scalar.years[1].reason or "")


def test_market_benchmark_uses_the_chosen_segment_table_when_cbre_carries_it() -> None:
    rows = [
        *STR_ROWS,
        _row("ttm_performance.segment.upper_priced.occupancy_pct", 74.0, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=7),
        _row("ttm_performance.segment.upper_priced.adr_usd", 290.0, doc=MARKET_DOC, doc_type="MARKET_STUDY", page=7),
        *CBRE_ROWS,
    ]
    r = resolve_index_by_year(
        rows,
        {"index_methodology": METHOD_MARKET_BENCHMARK, "index_market_segment": "upper_priced", KEY_ARI_TARGET: {"value": 100, "note": "n"}},
        hold_years=2, calendar_years=[2025, 2026],
    )
    assert r.years[1].benchmark_adr.value == pytest.approx(310.0)
    assert "upper priced" in (r.years[1].benchmark_adr.detail or "")


def test_market_benchmark_without_a_close_date_cannot_place_the_forecast() -> None:
    r = resolve_index_by_year(
        [*STR_ROWS, *MARKET_ROWS, *CBRE_ROWS],
        {"index_methodology": METHOD_MARKET_BENCHMARK, KEY_MPI_TARGET_BY_YEAR: {"value": [110, 110], "note": "n"}},
        hold_years=2, calendar_years=[],
    )
    assert r.calendar_known is False
    assert r.years[0].occupancy == pytest.approx(0.712 * 1.10)
    assert r.years[1].occupancy_status == "benchmark_unavailable"
    assert "close date" in (r.years[1].reason or "")


def test_disabled_method_leaves_every_year_unavailable_with_the_reason() -> None:
    r = resolve_index_by_year(
        STR_ROWS, {"index_methodology": METHOD_COSTAR_COMP_SET, KEY_MPI_TARGET_BY_YEAR: {"value": [110, 110], "note": "n"}}, hold_years=2
    )
    assert all(y.occupancy_status == "benchmark_unavailable" for y in r.years)
    assert "CoStar" in (r.years[0].reason or "")
    # With only the scalar target, Year 1 says CoStar is unavailable and Year 2
    # says the scalar drives Year 1 only (the gate is the earlier question).
    r2 = resolve_index_by_year(
        STR_ROWS, {"index_methodology": METHOD_COSTAR_COMP_SET, KEY_MPI_TARGET: {"value": 110, "note": "n"}}, hold_years=2
    )
    assert r2.years[0].occupancy_status == "benchmark_unavailable" and "CoStar" in (r2.years[0].reason or "")
    assert r2.years[1].occupancy_status == "scalar_year1_only"


def test_an_impossible_occupancy_is_refused_per_year() -> None:
    r = resolve_index_by_year(STR_ROWS, {KEY_MPI_TARGET_BY_YEAR: {"value": [200], "note": "n"}}, hold_years=1)
    assert r.years[0].occupancy is None and r.years[0].occupancy_status == "benchmark_unavailable"


# ─────────────────────────── runner seam ───────────────────────────


def _base(**kw: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "starting_occupancy": 0.846, "starting_adr": 255.0,
        "occupancy_growth": 0.008, "adr_growth": 0.04, "hold_years": 3,
    }
    base.update(kw)
    return base


def test_wants_benchmark_covers_the_by_year_lists() -> None:
    assert wants_benchmark({KEY_MPI_TARGET_BY_YEAR: [None, 115]}, {KEY_MPI_TARGET_BY_YEAR}) is True
    assert wants_benchmark({KEY_MPI_TARGET_BY_YEAR: [None, None]}, {KEY_MPI_TARGET_BY_YEAR}) is False
    assert wants_benchmark({KEY_MPI_TARGET_BY_YEAR: [None, 115]}, set()) is False


def test_seam_does_nothing_without_a_target() -> None:
    base = _base()
    sources = {"starting_occupancy": "str_subject_ttm", "starting_adr": "str_subject_ttm"}
    before = dict(base)
    assert build_index_drivers(base, sources, set(), rows=STR_ROWS, str_basis_sources=STR_BASIS, hold_years=3, calendar_years=[]) == []
    assert base == before and BASE_OCC_BY_YEAR not in base


def test_seam_parks_the_lists_and_repoints_year_one_by_year_beating_scalar() -> None:
    base = _base(**{
        KEY_MPI_TARGET_BY_YEAR: [120.0, 115.0, None], KEY_MPI_TARGET: 110.0,
        KEY_ARI_TARGET: 100.0, KEY_OCC_GROWTH: 0.012,
    })
    sources = {"starting_occupancy": "str_subject_ttm", "starting_adr": "str_subject_ttm"}
    paths = {KEY_MPI_TARGET_BY_YEAR, KEY_MPI_TARGET, KEY_ARI_TARGET, KEY_OCC_GROWTH}
    changed = build_index_drivers(base, sources, paths, rows=STR_ROWS, str_basis_sources=STR_BASIS, hold_years=3, calendar_years=[])
    assert set(changed) == {"starting_occupancy", "starting_adr"}
    # Year 1: by-year 120 beats scalar 110.
    assert base["starting_occupancy"] == pytest.approx(COMP_OCC * 1.20)
    assert sources["starting_occupancy"] == SOURCE_INDEX_ASSUMPTION
    assert base["starting_adr"] == pytest.approx(COMP_ADR)
    occ = base[BASE_OCC_BY_YEAR]
    assert occ[0] == pytest.approx(COMP_OCC * 1.20)
    assert occ[1] == pytest.approx(COMP_OCC * 1.012 * 1.15)
    # Year 3 falls back to the scalar 110 on the grown benchmark — driven because
    # the analyst OVERRODE the market occupancy growth (the out-year gate).
    assert occ[2] == pytest.approx(COMP_OCC * 1.012**2 * 1.10)
    # ADR: scalar ARI with no ADR growth override → Year 1 only, flagged after.
    assert base[BASE_ADR_BY_YEAR] == [pytest.approx(COMP_ADR), None, None]
    drivers = base[BASE_DRIVERS_BY_YEAR]
    assert [d["year"] for d in drivers] == [1, 2, 3]
    assert drivers[1]["adr_status"] == "scalar_year1_only" and drivers[1]["occupancy_status"] == "driven"
    assert "ADR: the all-years target drives Year 1 only" in drivers[1]["reason"]
    assert drivers[1]["mpi_target_source"] == "by_year" and drivers[2]["mpi_target_source"] == "scalar"
    assert drivers[1]["rgi_derived"] == pytest.approx(115.0)
    assert drivers[0]["benchmark_doc"] == STR_DOC


def test_seam_never_overwrites_an_explicit_analyst_rate() -> None:
    base = _base(**{KEY_MPI_TARGET_BY_YEAR: [120.0, 115.0], KEY_ARI_TARGET: 100.0})
    sources = {"starting_occupancy": "analyst_override", "starting_adr": "str_subject_ttm"}
    changed = build_index_drivers(
        base, sources, {KEY_MPI_TARGET_BY_YEAR, KEY_ARI_TARGET, "starting_occupancy"},
        rows=STR_ROWS, str_basis_sources=STR_BASIS, hold_years=2, calendar_years=[],
    )
    assert changed == ["starting_adr"]
    assert base["starting_occupancy"] == pytest.approx(0.846)
    assert base[BASE_OCC_BY_YEAR] == [None, None]
    assert base[BASE_DRIVERS_BY_YEAR][0]["occupancy_status"] == "no_target"
    assert "explicit analyst value" in base[BASE_DRIVERS_BY_YEAR][0]["reason"]


def test_seam_explicit_rate_reason_joins_the_other_metrics_reason() -> None:
    # Occupancy is an explicit analyst value (not driven); ADR is on the STR
    # basis with a scalar-only ARI target, so Year 2 carries the out-year gate
    # reason for ADR. The explicit-value note must JOIN it, not replace it.
    base = _base(**{KEY_MPI_TARGET_BY_YEAR: [120.0, 115.0], KEY_ARI_TARGET: 100.0, KEY_OCC_GROWTH: 0.012})
    sources = {"starting_occupancy": "analyst_override", "starting_adr": "str_subject_ttm"}
    build_index_drivers(
        base, sources, {KEY_MPI_TARGET_BY_YEAR, KEY_ARI_TARGET, KEY_OCC_GROWTH, "starting_occupancy"},
        rows=STR_ROWS, str_basis_sources=STR_BASIS, hold_years=2, calendar_years=[],
    )
    y2 = base[BASE_DRIVERS_BY_YEAR][1]
    assert y2["occupancy_status"] == "no_target" and y2["adr_status"] == "scalar_year1_only"
    assert "explicit analyst value" in y2["reason"]
    assert "ADR: the all-years target drives Year 1 only" in y2["reason"]


def test_seam_needs_the_str_basis_to_have_landed() -> None:
    base = _base(**{KEY_MPI_TARGET_BY_YEAR: [120.0]})
    sources = {"starting_occupancy": "t12_actual", "starting_adr": "t12_actual"}
    assert build_index_drivers(base, sources, {KEY_MPI_TARGET_BY_YEAR}, rows=STR_ROWS, str_basis_sources=STR_BASIS, hold_years=1, calendar_years=[]) == []
    assert BASE_OCC_BY_YEAR not in base


# ─────────────────────────── revenue engine ───────────────────────────


def _engine_input(**kw: Any):
    from fondok_schemas.underwriting import RevenueEngineInput

    base = dict(deal_id=UUID("11111111-1111-1111-1111-111111111111"), keys=100, starting_occupancy=0.70,
                starting_adr=200.0, occupancy_growth=0.01, adr_growth=0.03, hold_years=5)
    base.update(kw)
    return RevenueEngineInput(**base)


def test_engine_output_is_byte_identical_without_drivers() -> None:
    from app.engines.revenue import RevenueEngine

    a = RevenueEngine().run(_engine_input())
    b = RevenueEngine().run(_engine_input(occupancy_by_year_override=None, adr_by_year_override=[], index_drivers_by_year=[]))
    assert a.model_dump_json() == b.model_dump_json()
    assert a.index_drivers_by_year == []


def test_engine_drivers_replace_the_year_and_reanchor_the_years_after() -> None:
    from app.engines.revenue import RevenueEngine
    from fondok_schemas.underwriting import IndexDriverYear

    out = RevenueEngine().run(_engine_input(
        occupancy_by_year_override=[None, None, 0.80, None, None],
        adr_by_year_override=[210.0, None, None, 260.0, None],
        index_drivers_by_year=[
            IndexDriverYear(year=1, benchmark_adr=210.0, ari_target=100.0, ari_target_source="scalar", adr_status="driven"),
            IndexDriverYear(year=3, benchmark_occupancy=0.72, mpi_target=111.1, mpi_target_source="by_year", occupancy_status="driven"),
            IndexDriverYear(year=4, benchmark_adr=250.0, ari_target=104.0, ari_target_source="by_year", adr_status="driven"),
        ],
    ))
    occ = [y.occupancy for y in out.years]
    adr = [y.adr for y in out.years]
    # A Year-1 ADR driver replaces the baseline; undriven years compound from it.
    assert adr[0] == pytest.approx(210.0) and adr[1] == pytest.approx(210.0 * 1.03) and adr[2] == pytest.approx(210.0 * 1.03**2)
    # Year 4 driven, Year 5 compounds from Year 4 — not from the baseline.
    assert adr[3] == pytest.approx(260.0) and adr[4] == pytest.approx(260.0 * 1.03)
    # Occupancy: legacy compounding until Year 3, then the driver, then re-anchored growth.
    assert occ[1] == pytest.approx(0.70 * 1.01)
    assert occ[2] == pytest.approx(0.80) and occ[3] == pytest.approx(0.80 * 1.01)
    # The drivers are echoed with the figures that ran, and the lineage names the target.
    echoed = {d.year: d for d in out.index_drivers_by_year}
    assert echoed[3].occupancy == pytest.approx(0.80) and echoed[3].adr == pytest.approx(210.0 * 1.03**2)
    trace = out.provenance["years[2].rooms_revenue"]
    assert "Year 3 is indexed" in (trace.note or "")
    assert {i.name: i.assumption_key for i in trace.inputs}["mpi_target"] == "index_mpi_target_by_year"
    assert {i.name: i.assumption_key for i in trace.inputs}["ari_target"] == "index_ari_target" if "ari_target" in {i.name for i in trace.inputs} else True
    y1 = out.provenance["years[0].rooms_revenue"]
    assert {i.name: i.assumption_key for i in y1.inputs}["ari_target"] == "index_ari_target"
    # An undriven year's trace is the legacy text — no driver note.
    assert "indexed" not in (out.provenance["years[1].rooms_revenue"].note or "")


def test_engine_caps_a_driven_occupancy_at_the_model_ceiling() -> None:
    from app.engines.revenue import RevenueEngine

    out = RevenueEngine().run(_engine_input(occupancy_by_year_override=[None, 0.99]))
    assert out.years[1].occupancy == pytest.approx(0.95)


# ─────────────────────────── runner integration (DB) ───────────────────────────


_STR_FIELDS = [
    {"field_name": "ttm_performance.subject.occupancy_pct", "value": 78.5},
    {"field_name": "ttm_performance.subject.adr_usd", "value": 412.0},
    {"field_name": "ttm_performance.indices.mpi_occupancy_index", "value": 103.2},
    {"field_name": "ttm_performance.indices.ari_adr_index", "value": 94.2},
]
_COMP_OCC_DB = 0.785 / 1.032
_COMP_ADR_DB = 412.0 / 0.942


async def _str_deal(session, *, overrides: dict, tenant_id: str | None = None) -> tuple[str, str]:
    deal_id, tenant_id, doc_id = str(uuid4()), tenant_id or str(uuid4()), str(uuid4())
    now = datetime.now(UTC)
    await session.execute(
        text(
            "INSERT INTO deals (id, tenant_id, name, status, keys, field_overrides, created_at, updated_at) "
            "VALUES (:id, :tenant, 'By-Year Hotel', 'Draft', 132, :ov, :now, :now)"
        ),
        {"id": deal_id, "tenant": tenant_id, "ov": json.dumps(overrides), "now": now},
    )
    await session.execute(
        text(
            "INSERT INTO documents (id, deal_id, tenant_id, filename, doc_type, status, uploaded_at) "
            "VALUES (:id, :deal, :tenant, 'str_trend.pdf', 'STR_TREND', 'EXTRACTED', :now)"
        ),
        {"id": doc_id, "deal": deal_id, "tenant": tenant_id, "now": now},
    )
    await session.execute(
        text(
            "INSERT INTO extraction_results (id, document_id, deal_id, tenant_id, fields, confidence_report, "
            "agent_version, created_at) VALUES (:id, :doc, :deal, :tenant, :f, '{}', 'test', :now)"
        ),
        {"id": str(uuid4()), "doc": doc_id, "deal": deal_id, "tenant": tenant_id, "f": json.dumps(_STR_FIELDS), "now": now},
    )
    await session.commit()
    return deal_id, tenant_id


_BY_YEAR_OVERRIDES = {
    KEY_MPI_TARGET_BY_YEAR: {"value": [None, 115.0, None, None, None], "note": "Year-2 share gain after the PIP"},
    KEY_ARI_TARGET: {"value": 100.0, "note": "Parity on rate"},
    KEY_OCC_GROWTH: {"value": 0.012, "note": "Submarket demand recovery"},
}


@pytest.mark.asyncio
async def test_by_year_targets_move_nothing_with_the_toggle_off() -> None:
    from app.database import get_session_factory
    from app.services.engine_runner import _build_input_for, _load_engine_inputs

    async with get_session_factory()() as session:
        plain_id, plain_t = await _str_deal(session, overrides={})
        plain = await _load_engine_inputs(session, plain_id, tenant_id=plain_t)
        deal_id, tenant_id = await _str_deal(session, overrides=dict(_BY_YEAR_OVERRIDES))
        base = await _load_engine_inputs(session, deal_id, tenant_id=tenant_id)
    for key in ("starting_occupancy", "starting_adr", "occupancy_growth", "adr_growth"):
        assert base[key] == plain[key], key
    assert BASE_OCC_BY_YEAR not in base and BASE_DRIVERS_BY_YEAR not in base
    rev_in = _build_input_for("revenue", deal_id, base, {})
    assert rev_in.occupancy_by_year_override is None and rev_in.index_drivers_by_year == []


@pytest.mark.asyncio
async def test_by_year_targets_drive_the_revenue_projection_with_the_toggle_on() -> None:
    from app.database import get_session_factory
    from app.engines.revenue import RevenueEngine
    from app.services.engine_runner import _build_input_for, _load_engine_inputs

    async with get_session_factory()() as session:
        deal_id, tenant_id = await _str_deal(
            session, overrides={**_BY_YEAR_OVERRIDES, "revenue_seed_from_str_forecast": True}
        )
        base = await _load_engine_inputs(session, deal_id, tenant_id=tenant_id)

    sources = base["__sources__"]
    assert base["occupancy_growth"] == pytest.approx(0.012)
    # Year 1: the scalar ARI (parity) re-points ADR; no MPI for Year 1 → the subject TTM occupancy stays.
    assert base["starting_adr"] == pytest.approx(_COMP_ADR_DB) and sources["starting_adr"] == SOURCE_INDEX_ASSUMPTION
    assert base["starting_occupancy"] == pytest.approx(0.785) and sources["starting_occupancy"] == "str_subject_ttm"
    occ = base[BASE_OCC_BY_YEAR]
    assert occ[0] is None and occ[1] == pytest.approx(_COMP_OCC_DB * 1.012 * 1.15) and occ[2:] == [None, None, None]
    # Scalar ARI with no ADR growth override → ADR driven in Year 1 only, later years flagged.
    assert base[BASE_ADR_BY_YEAR] == [pytest.approx(_COMP_ADR_DB), None, None, None, None]
    drivers = base[BASE_DRIVERS_BY_YEAR]
    assert drivers[1]["occupancy_status"] == "driven" and drivers[1]["adr_status"] == "scalar_year1_only"
    assert "drives Year 1 only" in drivers[1]["reason"]

    rev_in = _build_input_for("revenue", deal_id, base, {})
    out = RevenueEngine().run(rev_in)
    years = out.years
    # Year 2 occupancy = benchmark × MPI; Year 3 keeps the model's growth from Year 2 (the last driven year).
    assert years[1].occupancy == pytest.approx(_COMP_OCC_DB * 1.012 * 1.15)
    assert years[2].occupancy == pytest.approx(years[1].occupancy * 1.012)
    # ADR: Year 1 at parity with the comp set, then the model's own ADR growth — nothing invented.
    assert years[0].adr == pytest.approx(_COMP_ADR_DB)
    assert years[1].adr == pytest.approx(_COMP_ADR_DB * (1 + base["adr_growth"]))
    echoed = {d.year: d for d in out.index_drivers_by_year}
    assert echoed[2].occupancy == pytest.approx(years[1].occupancy)
    assert echoed[2].calendar_year is None  # no close date on this deal — never guessed


@pytest.mark.asyncio
async def test_by_year_beats_scalar_in_the_loader() -> None:
    from app.database import get_session_factory
    from app.services.engine_runner import _load_engine_inputs

    async with get_session_factory()() as session:
        deal_id, tenant_id = await _str_deal(
            session,
            overrides={
                KEY_MPI_TARGET_BY_YEAR: {"value": [120.0], "note": "Year 1 share"},
                KEY_MPI_TARGET: {"value": 110.0, "note": "all years"},
                "revenue_seed_from_str_forecast": True,
            },
        )
        base = await _load_engine_inputs(session, deal_id, tenant_id=tenant_id)
    assert base["starting_occupancy"] == pytest.approx(_COMP_OCC_DB * 1.20)
    assert base["__sources__"]["starting_occupancy"] == SOURCE_INDEX_ASSUMPTION
    assert base[BASE_OCC_BY_YEAR][0] == pytest.approx(_COMP_OCC_DB * 1.20)


# ─────────────────────────── endpoint ───────────────────────────


@pytest.mark.asyncio
async def test_endpoint_carries_the_by_year_grid_on_the_deal_calendar_and_is_tenant_scoped() -> None:
    from app.api.market import index_methodology
    from app.database import get_session_factory

    async with get_session_factory()() as session:
        deal_id, tenant_id = await _str_deal(
            session,
            overrides={
                **_BY_YEAR_OVERRIDES,
                "hold_years": 3,
                "acquisition_close_date": "2025-09-30",
            },
        )
        out = await index_methodology(deal_id=UUID(deal_id), session=session, tenant_id=UUID(tenant_id))
        with pytest.raises(HTTPException) as exc:
            await index_methodology(deal_id=UUID(deal_id), session=session, tenant_id=uuid4())
    assert exc.value.status_code == 404
    payload = out.model_dump(mode="json")
    assert payload["hold_years"] == 3 and payload["calendar_known"] is True
    assert [r["calendar_year"] for r in payload["by_year"]] == [2025, 2026, 2027]
    assert payload["by_year_applies"] is False and "Use STR rates" in payload["by_year_disabled_reason"]
    y2 = payload["by_year"][1]
    assert y2["mpi_target"] == {"value": 115.0, "source": "by_year", "note": "Year-2 share gain after the PIP"}
    assert y2["ari_target"]["source"] == "scalar" and y2["rgi_derived"] == pytest.approx(115.0)
    assert y2["occupancy"] == pytest.approx(_COMP_OCC_DB * 1.012 * 1.15) and y2["occupancy_status"] == "driven"
    assert y2["adr_status"] == "scalar_year1_only" and "drives Year 1 only" in y2["reason"]
    assert y2["benchmark_occupancy"]["source"] == "computed"


def test_by_year_target_keys_require_a_justification() -> None:
    from app.api.deals import _NOTE_EXEMPT_KEYS, _override_needs_note

    for key in (KEY_MPI_TARGET_BY_YEAR, KEY_ARI_TARGET_BY_YEAR, "index_mpi_target_by_year.3"):
        assert key not in _NOTE_EXEMPT_KEYS
        assert _override_needs_note(key, frozenset()) is True


# ─────────────────────────── Excel round-trip (pure) ───────────────────────────


def test_assumptions_sheet_rows_and_import_rules_for_by_year_targets() -> None:
    from app.export.pl_workbooks import (
        ASSUMPTIONS,
        assumption_spec,
        build_projections_workbook,
        by_year_target_assumptions,
        parse_projections_import,
        projection_values,
        split_by_year_target_key,
    )

    specs = ASSUMPTIONS + by_year_target_assumptions(2)
    keys = [s.key for s in specs]
    assert keys[-4:] == [
        "index_mpi_target_by_year.1", "index_mpi_target_by_year.2",
        "index_ari_target_by_year.1", "index_ari_target_by_year.2",
    ]
    assert split_by_year_target_key("index_ari_target_by_year.2") == ("index_ari_target_by_year", 2)
    assert split_by_year_target_key("index_ari_target_by_year.25") is None
    assert assumption_spec("index_mpi_target_by_year.2").label == "Occupancy penetration target (MPI) — Year 2"
    assert "benchmark" in assumption_spec("index_mpi_target_by_year.1").unit

    cur = {s.key: None for s in specs}
    cur["index_mpi_target_by_year.2"] = 110.0
    rows = [
        {"key": s.key, "label": s.label, "value": cur[s.key], "unit": s.unit, "source": "", "method": s.method,
         "note_required": True, "current_note": None}
        for s in specs
    ]
    _n, values = projection_values({"revenue": {"years": [{"occupancy": 0.7}] * 2}})
    data = build_projections_workbook(deal_id="deal-1", deal_name="X", year_headers=["Base year (Year 1)", "Year 2"],
                                      values=values, assumptions=rows)
    wb = load_workbook(io.BytesIO(data))
    wa = wb["Assumptions"]
    at = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    assert wa.cell(row=at["index_mpi_target_by_year.2"], column=3).value == 110.0
    wa.cell(row=at["index_mpi_target_by_year.2"], column=3).value = 115
    wa.cell(row=at["index_mpi_target_by_year.2"], column=9).value = "Post-PIP share gain"
    wa.cell(row=at["index_ari_target_by_year.1"], column=3).value = 98      # no note → rejected
    r = wa.max_row + 1
    wa.cell(row=r, column=1).value = "index_mpi_target_by_year.25"          # past the model cap
    wa.cell(row=r, column=3).value = 1
    buf = io.BytesIO()
    wb.save(buf)
    p = parse_projections_import(buf.getvalue(), deal_id="deal-1", current_values=values, current_assumptions=cur,
                                 needs_note=lambda _k: True)
    assert [(c["key"], c["old_value"], c["new_value"], c["note"]) for c in p["changes"]] == [
        ("index_mpi_target_by_year.2", 110.0, 115, "Post-PIP share gain"),
    ]
    assert [(x["key"], x["reason"]) for x in p["rejected"]] == [("index_ari_target_by_year.1", "note_required")]
    assert [(x["key"], x["reason"]) for x in p["mapping_errors"]] == [("index_mpi_target_by_year.25", "unknown_key")]
