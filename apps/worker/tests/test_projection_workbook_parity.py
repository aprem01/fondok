"""Task 15 — Future P&L workbook parity with the retired browser builder.

The worker's ``GET /deals/{id}/exports/projections.xlsx`` now carries what the
deleted web builder (``ProjectionsSection.onExport``) carried:

* the Days / Number of Rooms / Available Rooms / Occupied Rooms rows, read
  from the revenue engine's published occupied-rooms and RevPAR traces;
* per-year ``% Rev`` / ``PAR`` / ``POR`` columns, computed from the row's own
  amount and that year's revenue / available rooms / occupied rooms;
* the Net Cash Flow row, from the cash-flow engine's unlevered statement;
* the Exit Year column with the returns engine's terminal NOI and its basis.

Every one of them is read-only: the importer ignores them. An input the run
did not publish is "not available: <reason>", never 0.
"""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook

from tests.fixtures.kimpton_seed import KIMPTON_FIXTURE_DEAL_ID, kimpton_seed

_TMP_DB = Path(tempfile.gettempdir()) / "fondok-tests-projection-parity.db"
if _TMP_DB.exists():
    _TMP_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_TMP_DB}"

_RUN: dict[str, dict] = {}


async def _kimpton_outputs() -> dict[str, dict]:
    """One real full-chain run on the Kimpton fixture (cached per module)."""
    if _RUN:
        return _RUN
    from app.database import get_session_factory
    from app.migrations import run_startup_migrations
    from app.services.engine_runner import run_all_engines

    await run_startup_migrations()
    async with get_session_factory()() as s:
        results = await run_all_engines(
            s, deal_id=KIMPTON_FIXTURE_DEAL_ID, tenant_id=str(uuid4()),
            run_id=str(uuid4()), assumptions=kimpton_seed(),
        )
    for name, env in results.items():
        assert env["status"] == "complete", (name, env)
        _RUN[name] = env["outputs"]
    return _RUN


def _headers(n: int) -> list[str]:
    return ["Base year (Year 1)" if i == 0 else f"Year {i + 1}" for i in range(n)]


def _assumptions() -> list[dict]:
    from app.export.pl_workbooks import ASSUMPTIONS

    return [
        {"key": a.key, "label": a.label, "value": 0.045 if a.key == "revpar_growth" else None,
         "unit": a.unit, "source": "seed", "method": a.method,
         "note_required": True, "current_note": None}
        for a in ASSUMPTIONS
    ]


def _build(outputs: dict, n: int | None = None) -> tuple[bytes, dict, int]:
    from app.export.pl_workbooks import (
        build_projections_workbook,
        projection_extras,
        projection_values,
    )

    count, values = projection_values(outputs)
    if n is not None:
        count = n
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="Kimpton", year_headers=_headers(count),
        values=values, assumptions=_assumptions(),
        extras=projection_extras(outputs, count),
    )
    return data, values, count


def _grid(ws) -> tuple[dict[str, int], dict[str, int]]:
    """``({row label: row}, {column header: column})``."""
    rows = {ws.cell(row=r, column=2).value: r for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=2).value is not None}
    cols = {ws.cell(row=1, column=c).value: c for c in range(1, ws.max_column + 1)}
    return rows, cols


def _is_refusal(v: object) -> bool:
    return isinstance(v, str) and v.startswith("not available: ") and len(v) > len("not available: ")


@pytest.mark.asyncio
async def test_real_run_workbook_carries_every_parity_row_and_column() -> None:
    from app.export.pl_workbooks import FORWARD_NOI_LABEL

    out = await _kimpton_outputs()
    data, _values, n = _build(out)
    assert n == 5
    ws = load_workbook(io.BytesIO(data))["Future P&L"]
    rows, cols = _grid(ws)

    # Column groups in the browser builder's order: Amount | % Rev | PAR | POR
    # per model year, then the display-only Exit Year column.
    header = [ws.cell(row=1, column=c).value for c in range(3, ws.max_column + 1)]
    expected: list[str] = []
    for h in _headers(n):
        expected += [f"{h} Amount", f"{h} % Rev", f"{h} PAR", f"{h} POR"]
    expected.append("Exit Year")
    assert header == expected

    rev = out["revenue"]
    exp_years = out["expense"]["years"]
    for i, h in enumerate(_headers(n)):
        trace = rev["provenance"][f"years[{i}].occupied_rooms"]
        ins = {x["name"]: x["value"] for x in trace["inputs"]}
        avail = {x["name"]: x["value"] for x in rev["provenance"][f"years[{i}].revpar"]["inputs"]}
        amt = cols[f"{h} Amount"]
        assert ws.cell(row=rows["Days"], column=amt).value == ins["days"] == 365
        assert ws.cell(row=rows["Number of Rooms"], column=amt).value == ins["keys"] == 132
        assert ws.cell(row=rows["Available Rooms"], column=amt).value == avail["rooms_available"]
        assert ws.cell(row=rows["Occupied Rooms"], column=amt).value == pytest.approx(trace["value"])

        rooms_rev = rev["years"][i]["rooms_revenue"]
        total = exp_years[i]["total_revenue"]
        r = rows["Rooms Revenue"]
        assert ws.cell(row=r, column=cols[f"{h} % Rev"]).value == pytest.approx(rooms_rev / total)
        # Rooms revenue per available room IS RevPAR; per occupied room IS ADR.
        assert ws.cell(row=r, column=cols[f"{h} PAR"]).value == pytest.approx(rev["years"][i]["revpar"])
        assert ws.cell(row=r, column=cols[f"{h} POR"]).value == pytest.approx(rev["years"][i]["adr"])
        # Departmental expense % Rev is over its own department's revenue (USALI,
        # the same basis as the Benchmark % column).
        dr = rows["Rooms Department Expense"]
        assert ws.cell(row=dr, column=cols[f"{h} % Rev"]).value == pytest.approx(
            exp_years[i]["dept_expenses"]["rooms"] / rooms_rev
        )
        # A rate row carries no ratio — a dash, not a number.
        assert ws.cell(row=rows["Occupancy"], column=cols[f"{h} % Rev"]).value == "—"
        assert ws.cell(row=rows["Days"], column=cols[f"{h} PAR"]).value == "—"

        # Net Cash Flow = the cash-flow engine's NOI (before reserve) + FF&E row.
        unlev = {line["label"]: line["values"] for line in out["cash_flow"]["unlevered"]}
        ncf = unlev["NOI (before FF&E reserve)"][i + 1] + unlev["FF&E Reserve"][i + 1]
        nr = rows["Net Cash Flow"]
        assert ws.cell(row=nr, column=amt).value == pytest.approx(ncf)
        assert ws.cell(row=nr, column=amt).value == pytest.approx(out["returns"]["noi_by_year"][i])
        assert ws.cell(row=nr, column=cols[f"{h} % Rev"]).value == pytest.approx(ncf / total)
        # Nothing but the exit NOI sits in the Exit Year column.
        assert ws.cell(row=nr, column=cols["Exit Year"]).value == "—"

    ret = out["returns"]
    fwd_label = f"{FORWARD_NOI_LABEL} — {ret['exit_noi_period_label']}"
    assert ret["exit_noi_basis"] == "forward_12m"
    fr = rows[fwd_label]
    assert ws.cell(row=fr, column=cols["Exit Year"]).value == pytest.approx(ret["terminal_noi"])
    assert ws.cell(row=fr, column=cols["Base year (Year 1) Amount"]).value == "—"

    # Every new cell is read-only styled, like the computed lines around it.
    for label in ("Days", "Net Cash Flow", fwd_label):
        assert ws.cell(row=rows[label], column=3).fill.fgColor.rgb.endswith("F3F2EE")
    # Analyst-recognisable order: operating stats first, NCF then exit NOI last.
    order = sorted(rows, key=rows.get)
    assert order.index("Days") < order.index("Number of Rooms") < order.index(
        "Available Rooms") < order.index("Occupied Rooms") < order.index("Occupancy")
    assert order.index("FF&E Reserve") < order.index("Net Cash Flow") < order.index(fwd_label)
    assert order[-1] == fwd_label


@pytest.mark.asyncio
async def test_absent_inputs_render_the_refusal_sentence_never_zero() -> None:
    """A run that published no occupied-rooms trace, no cash-flow statement and
    no terminal NOI: every dependent cell says why it is empty."""
    out = await _kimpton_outputs()
    stripped = {
        "revenue": {k: v for k, v in out["revenue"].items() if k != "provenance"},
        "fb": out["fb"],
        "expense": {
            **out["expense"],
            "years": [{**y, "total_revenue": None} if i == 1 else y
                      for i, y in enumerate(out["expense"]["years"])],
        },
    }
    data, _values, n = _build(stripped)
    ws = load_workbook(io.BytesIO(data))["Future P&L"]
    rows, cols = _grid(ws)
    hs = _headers(n)

    for label in ("Days", "Number of Rooms", "Available Rooms", "Occupied Rooms",
                  "Net Cash Flow"):
        for h in hs:
            v = ws.cell(row=rows[label], column=cols[f"{h} Amount"]).value
            assert _is_refusal(v), (label, h, v)
    assert "revenue" in ws.cell(row=rows["Days"], column=cols[f"{hs[0]} Amount"]).value
    assert "cash-flow" in ws.cell(row=rows["Net Cash Flow"], column=cols[f"{hs[0]} Amount"]).value
    exit_rows = [r for lbl, r in rows.items() if isinstance(lbl, str) and "NOI" in lbl
                 and ws.cell(row=r, column=1).value == "fondok_readonly:exit_noi"]
    assert len(exit_rows) == 1
    assert _is_refusal(ws.cell(row=exit_rows[0], column=cols["Exit Year"]).value)
    assert "returns" in ws.cell(row=exit_rows[0], column=cols["Exit Year"]).value

    # PAR / POR have no denominator without the rooms traces …
    rr = rows["Rooms Revenue"]
    assert _is_refusal(ws.cell(row=rr, column=cols[f"{hs[0]} PAR"]).value)
    assert _is_refusal(ws.cell(row=rr, column=cols[f"{hs[0]} POR"]).value)
    # … % Rev has none in the year whose total revenue is missing, and the
    # Total Revenue amount itself says so instead of being blank or 0.
    assert _is_refusal(ws.cell(row=rr, column=cols[f"{hs[1]} % Rev"]).value)
    assert ws.cell(row=rr, column=cols[f"{hs[0]} % Rev"]).value == pytest.approx(
        out["revenue"]["years"][0]["rooms_revenue"] / out["expense"]["years"][0]["total_revenue"]
    )
    assert _is_refusal(ws.cell(row=rows["Total Revenue"], column=cols[f"{hs[1]} Amount"]).value)

    # No cell whose input was removed turns into a zero or a blank: every PAR /
    # POR cell (no room counts at all) is a refusal or a not-applicable dash.
    for r in range(3, ws.max_row + 1):
        for h in hs:
            for sub in ("PAR", "POR"):
                v = ws.cell(row=r, column=cols[f"{h} {sub}"]).value
                assert v == "—" or _is_refusal(v), (r, h, sub, v)


@pytest.mark.asyncio
async def test_round_trip_import_ignores_every_new_row_and_column() -> None:
    from app.export.pl_workbooks import parse_projections_import

    out = await _kimpton_outputs()
    data, values, n = _build(out)
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Future P&L"]
    rows, cols = _grid(ws)
    hs = _headers(n)
    # The analyst types over every read-only addition …
    for label in ("Days", "Number of Rooms", "Available Rooms", "Occupied Rooms", "Net Cash Flow"):
        ws.cell(row=rows[label], column=cols[f"{hs[1]} Amount"]).value = 1
    for h in hs:
        for sub in ("% Rev", "PAR", "POR"):
            ws.cell(row=rows["Rooms Revenue"], column=cols[f"{h} {sub}"]).value = 0.5
    for r in range(3, ws.max_row + 1):
        ws.cell(row=r, column=cols["Exit Year"]).value = 123
    # … and edits one assumption with its note.
    wa = wb["Assumptions"]
    keys = {wa.cell(row=r, column=1).value: r for r in range(2, wa.max_row + 1)}
    wa.cell(row=keys["revpar_growth"], column=3).value = 0.05
    wa.cell(row=keys["revpar_growth"], column=9).value = "STR forecast"
    buf = io.BytesIO()
    wb.save(buf)

    cur = {a["key"]: a["value"] for a in _assumptions()}
    p = parse_projections_import(
        buf.getvalue(), deal_id="deal-1", current_values=values,
        current_assumptions=cur, needs_note=lambda k: True,
    )
    assert [(c["key"], c["new_value"]) for c in p["changes"]] == [("revpar_growth", 0.05)]
    assert p["computed_edits"] == []
    assert p["mapping_errors"] == [] and p["non_numeric"] == [] and p["rejected"] == []

    # An untouched export re-imports clean — the refusal / dash cells are not
    # read as edits.
    clean = parse_projections_import(
        data, deal_id="deal-1", current_values=values,
        current_assumptions=cur, needs_note=lambda k: True,
    )
    assert clean["computed_edits"] == [] and clean["changes"] == []

    # A typed-over engine AMOUNT is still reported against the right model year.
    wb2 = load_workbook(io.BytesIO(data))
    ws2 = wb2["Future P&L"]
    ws2.cell(row=rows["Occupancy"], column=cols[f"{hs[2]} Amount"]).value = 0.9
    buf2 = io.BytesIO()
    wb2.save(buf2)
    p2 = parse_projections_import(
        buf2.getvalue(), deal_id="deal-1", current_values=values,
        current_assumptions=cur, needs_note=lambda k: True,
    )
    assert len(p2["computed_edits"]) == 1
    ce = p2["computed_edits"][0]
    assert ce["cell_key"] == "revenue.years[2].occupancy"
    assert ce["year_index"] == 2 and ce["period_label"] == hs[2]
    assert ce["new_value"] == 0.9


def test_a_workbook_from_the_previous_layout_still_imports() -> None:
    """Files downloaded before this change (one positional column per year,
    no column-key row) keep importing exactly as before."""
    from openpyxl import Workbook

    from app.export.pl_workbooks import parse_projections_import

    wb = Workbook()
    ws = wb.active
    ws.title = "Future P&L"
    ws.append(["fondok_path", "Line item (computed — read-only)", "Base year (Year 1)", "Year 2"])
    ws.append(["revenue.occupancy", "Occupancy", 0.7, 0.75])
    rd = wb.create_sheet("README")
    rd["A3"], rd["B3"] = "format", "fondok.projections.v1"
    rd["A4"], rd["B4"] = "deal_id", "deal-1"
    wa = wb.create_sheet("Assumptions")
    wa.append(["Key", "Assumption", "Value"])
    buf = io.BytesIO()
    wb.save(buf)
    values = {"revenue.years[0].occupancy": 0.7, "revenue.years[1].occupancy": 0.72}
    p = parse_projections_import(
        buf.getvalue(), deal_id="deal-1", current_values=values,
        current_assumptions={}, needs_note=lambda k: True,
    )
    assert [(c["cell_key"], c["new_value"]) for c in p["computed_edits"]] == [
        ("revenue.years[1].occupancy", 0.75)
    ]


@pytest.mark.asyncio
async def test_export_route_carries_the_parity_rows_from_the_deals_own_run() -> None:
    """``GET /deals/{id}/exports/projections.xlsx`` on a deal whose run is
    persisted: the rows come from that run's engine outputs, and the preview
    route re-reads the downloaded file without reporting any of them."""
    import json
    from datetime import UTC, datetime

    from sqlalchemy import text

    from app.api.pl_roundtrip import export_projections, preview_projections_import
    from app.auth.context import AuthContext
    from app.database import get_session_factory
    from app.migrations import run_startup_migrations
    from app.services.engine_runner import run_all_engines

    await run_startup_migrations()
    tenant, deal_id = str(uuid4()), uuid4()
    async with get_session_factory()() as s:
        ts = datetime.now(UTC)
        await s.execute(
            text(
                "INSERT INTO deals (id, tenant_id, name, status, ai_confidence, field_overrides, "
                "created_at, updated_at) VALUES (:id,:t,'Parity','Draft',0.0,:fo,:ts,:ts)"
            ),
            {"id": str(deal_id), "t": tenant, "ts": ts, "fo": json.dumps({})},
        )
        await s.commit()
    async with get_session_factory()() as s:
        results = await run_all_engines(
            s, deal_id=str(deal_id), tenant_id=tenant, run_id=str(uuid4()),
            assumptions=kimpton_seed(),
        )
    assert all(r["status"] == "complete" for r in results.values())
    auth = AuthContext(tenant_id=UUID(tenant), user_id="u", role="member",
                       source="jwt", org_id=None, email="a@fondok.test")
    async with get_session_factory()() as s:
        resp = await export_projections(deal_id=deal_id, session=s, auth=auth)
    ws = load_workbook(io.BytesIO(resp.body))["Future P&L"]
    rows, cols = _grid(ws)
    first = ws.cell(row=1, column=3).value
    assert first.endswith(" Amount")
    year1 = first[: -len(" Amount")]
    assert ws.cell(row=rows["Days"], column=cols[f"{year1} Amount"]).value == 365
    assert ws.cell(row=rows["Number of Rooms"], column=cols[f"{year1} Amount"]).value == 132
    ncf = ws.cell(row=rows["Net Cash Flow"], column=cols[f"{year1} Amount"]).value
    assert ncf == pytest.approx(results["returns"]["outputs"]["noi_by_year"][0])
    exit_header = next(h for h in cols if isinstance(h, str) and h.startswith("Exit Year"))
    exit_row = next(r for r in range(3, ws.max_row + 1)
                    if ws.cell(row=r, column=1).value == "fondok_readonly:exit_noi")
    assert ws.cell(row=exit_row, column=cols[exit_header]).value == pytest.approx(
        results["returns"]["outputs"]["terminal_noi"]
    )

    class _Up:
        async def read(self, n: int = -1) -> bytes:
            return resp.body

    async with get_session_factory()() as s:
        preview = await preview_projections_import(
            deal_id=deal_id, session=s, auth=auth, file=_Up(),
        )
    assert preview["computed_edits"] == []
    assert preview["changes"] == [] and preview["mapping_errors"] == []


@pytest.mark.asyncio
async def test_a_pinned_noi_is_shown_and_the_sheet_ties_out() -> None:
    """``noi_override_by_year`` pins Year 1 at $1M: Net Cash Flow stays the cash
    flow's, the pinned NOI gets its own read-only row, the pinned year's Net
    Cash Flow cell says so, and every year ties out on the sheet itself. The
    import ignores the new row."""
    from app.database import get_session_factory
    from app.export.pl_workbooks import (
        CASH_FLOW_NOI_PIN_LABEL,
        parse_projections_import,
    )
    from app.services.engine_runner import run_all_engines

    base = await _kimpton_outputs()
    unpinned = [y["noi"] for y in base["expense"]["years"]]
    pin = [1_000_000.0, *unpinned[1:]]
    async with get_session_factory()() as s:
        results = await run_all_engines(
            s, deal_id=KIMPTON_FIXTURE_DEAL_ID, tenant_id=str(uuid4()),
            run_id=str(uuid4()),
            assumptions={**kimpton_seed(), "noi_override_by_year": pin},
        )
    assert all(r["status"] == "complete" for r in results.values())
    out = {k: v["outputs"] for k, v in results.items()}
    assert out["returns"]["noi_by_year"][0] == pytest.approx(1_000_000.0)

    data, values, n = _build(out)
    ws = load_workbook(io.BytesIO(data))["Future P&L"]
    rows, cols = _grid(ws)
    hs = _headers(n)
    assert CASH_FLOW_NOI_PIN_LABEL in rows
    assert rows["FF&E Reserve"] < rows[CASH_FLOW_NOI_PIN_LABEL] < rows["Net Cash Flow"]
    assert ws.cell(row=rows[CASH_FLOW_NOI_PIN_LABEL], column=1).value.startswith(
        "fondok_readonly:")

    def amt(label: str, i: int):
        return ws.cell(row=rows[label], column=cols[f"{hs[i]} Amount"]).value

    for i in range(n):
        ffe = amt("FF&E Reserve", i)
        ncf = amt("Net Cash Flow", i)
        ncf_cell = ws.cell(row=rows["Net Cash Flow"], column=cols[f"{hs[i]} Amount"])
        if i == 0:
            # Pinned: the pin row + the sheet's own FF&E line = Net Cash Flow.
            assert ncf == pytest.approx(1_000_000.0)
            assert amt(CASH_FLOW_NOI_PIN_LABEL, i) - ffe == pytest.approx(ncf)
            assert ncf_cell.comment is not None
            assert "noi_override_by_year" in ncf_cell.comment.text
            # … and it is visibly NOT the operating NOI above it.
            assert amt("NOI (before FF&E reserve)", i) - ffe != pytest.approx(ncf)
        else:
            # Unpinned years: the NOI lines above tie out on their own.
            assert amt(CASH_FLOW_NOI_PIN_LABEL, i) == "—"
            assert amt("NOI (before FF&E reserve)", i) - ffe == pytest.approx(ncf)
            assert ncf_cell.comment is None

    # Unpinned deal: no pin row at all.
    plain = load_workbook(io.BytesIO(_build(base)[0]))["Future P&L"]
    assert CASH_FLOW_NOI_PIN_LABEL not in _grid(plain)[0]

    # Typed over, the pin row is not an edit.
    wb = load_workbook(io.BytesIO(data))
    w = wb["Future P&L"]
    for i in range(n):
        w.cell(row=rows[CASH_FLOW_NOI_PIN_LABEL], column=cols[f"{hs[i]} Amount"]).value = 7
    buf = io.BytesIO()
    wb.save(buf)
    p = parse_projections_import(
        buf.getvalue(), deal_id="deal-1", current_values=values,
        current_assumptions={a["key"]: a["value"] for a in _assumptions()},
        needs_note=lambda k: True,
    )
    assert p["computed_edits"] == [] and p["changes"] == [] and p["mapping_errors"] == []


def test_ratio_message_names_zero_and_negative_denominators_apart() -> None:
    from app.export.pl_workbooks import _ratio

    assert _ratio(5.0, "X", 0.0, "Total Revenue") == (
        "not available: Total Revenue is zero for this year")
    assert _ratio(5.0, "X", -3.0, "Total Revenue") == (
        "not available: Total Revenue is negative for this year")


@pytest.mark.asyncio
async def test_with_a_benchmark_column_every_parity_row_reads_a_dash_never_blank() -> None:
    """The methodology card: a Benchmark % with no published ratio shows a dash,
    never an empty cell — including the exit NOI row."""
    from app.export.pl_workbooks import (
        FORWARD_NOI_LABEL,
        build_projections_workbook,
        projection_extras,
        projection_values,
    )

    out = await _kimpton_outputs()
    count, values = projection_values(out)
    data = build_projections_workbook(
        deal_id="deal-1", deal_name="Kimpton", year_headers=_headers(count),
        values=values, assumptions=_assumptions(),
        extras=projection_extras(out, count),
        benchmark={"expense.dept_expenses.rooms": 0.25},
    )
    ws = load_workbook(io.BytesIO(data))["Future P&L"]
    rows, cols = _grid(ws)
    bm = next(c for name, c in cols.items() if "Benchmark" in name)
    ret = out["returns"]
    fwd_label = f"{FORWARD_NOI_LABEL} — {ret['exit_noi_period_label']}"
    for label in ("Days", "Number of Rooms", "Available Rooms", "Occupied Rooms",
                  "Net Cash Flow", fwd_label):
        assert ws.cell(row=rows[label], column=bm).value == "—", label
