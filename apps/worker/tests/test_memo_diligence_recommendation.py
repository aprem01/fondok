"""FON-54a — diligence status + IC recommendation confirmation persist through
the memo override channel and every consumer reads the same source.

* ``memo_diligence`` (keyed by variance *concept*) is normalised by
  :func:`app.memo_overrides.diligence_status`; the live Excel payload's
  Variance rows carry that status (no export-side copy).
* ``memo_recommendation_confirmed`` gates the verdict: memo body and live
  export header say "Pending analyst decision" until the analyst has
  selected AND confirmed a verdict — never the model's inferred verdict.

Pure functions only — no DB, no LLM.
"""

from __future__ import annotations


def _sections() -> list[dict]:
    return [
        {"section_id": "investment_thesis", "title": "Investment Thesis", "body": "gen thesis", "citations": []},
        {"section_id": "recommendation", "title": "Recommendation", "body": "gen rationale", "citations": []},
    ]


def _by_id(sections: list[dict]) -> dict[str, dict]:
    return {s["section_id"]: s for s in sections}


# ═══════════════════ recommendation confirmation ═══════════════════


def test_recommendation_decision_requires_valid_verdict_and_true_flag() -> None:
    from app.memo_overrides import recommendation_decision

    assert recommendation_decision(None) == (None, False)
    assert recommendation_decision({}) == (None, False)
    assert recommendation_decision({"memo_recommendation_override": "Proceed"}) == ("Proceed", False)
    assert recommendation_decision(
        {"memo_recommendation_override": "Proceed", "memo_recommendation_confirmed": True}
    ) == ("Proceed", True)
    # Truthy-but-not-True values do not confirm; out-of-vocabulary verdicts never do.
    assert recommendation_decision(
        {"memo_recommendation_override": "Proceed", "memo_recommendation_confirmed": "yes"}
    ) == ("Proceed", False)
    assert recommendation_decision(
        {"memo_recommendation_override": "Maybe", "memo_recommendation_confirmed": True}
    ) == (None, False)
    assert recommendation_decision({"memo_recommendation_confirmed": True}) == (None, False)


def test_ic_recommendation_label_is_pending_until_confirmed() -> None:
    from app.memo_overrides import PENDING_DECISION, ic_recommendation_label

    assert PENDING_DECISION == "Pending analyst decision"
    assert ic_recommendation_label(None) == PENDING_DECISION
    assert ic_recommendation_label({}) == PENDING_DECISION
    # Legacy row: verdict persisted before the confirm flag existed → pending.
    assert ic_recommendation_label({"memo_recommendation_override": "Do Not Proceed"}) == PENDING_DECISION
    assert ic_recommendation_label(
        {"memo_recommendation_override": "Do Not Proceed", "memo_recommendation_confirmed": True}
    ) == "Do Not Proceed"
    assert ic_recommendation_label(
        {"memo_recommendation_override": "Do Not Proceed", "memo_recommendation_confirmed": False}
    ) == PENDING_DECISION


def test_memo_recommendation_headline_pending_when_unconfirmed() -> None:
    from app.memo_overrides import apply_memo_overrides

    out = apply_memo_overrides(_sections(), {"memo_recommendation_override": "Proceed"})
    body = _by_id(out)["recommendation"]["body"]
    assert body.startswith("IC recommendation: Pending analyst decision.")
    assert "Proceed." not in body  # the draft verdict is not printed as recorded
    assert "gen rationale" in body


def test_memo_recommendation_headline_verdict_when_confirmed() -> None:
    from app.memo_overrides import apply_memo_overrides

    out = apply_memo_overrides(
        _sections(),
        {"memo_recommendation_override": "Proceed with Conditions", "memo_recommendation_confirmed": True},
    )
    body = _by_id(out)["recommendation"]["body"]
    assert body.startswith("IC recommendation: Proceed with Conditions.")
    assert "gen rationale" in body


def test_confirm_flag_alone_is_a_noop() -> None:
    from app.memo_overrides import apply_memo_overrides, has_memo_overrides

    sections = _sections()
    assert not has_memo_overrides({"memo_recommendation_confirmed": True})
    assert apply_memo_overrides(sections, {"memo_recommendation_confirmed": True}) is sections


# ═══════════════════ diligence status ═══════════════════


def test_diligence_status_normalises_and_drops_unknown() -> None:
    from app.memo_overrides import DILIGENCE_STATUSES, diligence_status

    assert DILIGENCE_STATUSES == {"Open", "Resolved", "Accepted"}
    assert diligence_status(None) == {}
    assert diligence_status({}) == {}
    assert diligence_status({"memo_diligence": "nope"}) == {}

    got = diligence_status(
        {
            "memo_diligence": {
                "rooms_revenue": {"status": "Resolved", "note": " tied to T-12 ", "updated_at": "2026-09-10T00:00:00Z"},
                "noi": {"status": "Accepted"},
                "gop": "Resolved",  # defensive: bare string
                "adr": {"status": "Bogus"},  # out of vocabulary → dropped (= Open)
                " ": {"status": "Resolved"},  # blank concept → dropped
                "occupancy": {"status": "Open", "details": True},  # UI-only key stripped
            }
        }
    )
    assert got == {
        "rooms_revenue": {"status": "Resolved", "note": "tied to T-12", "updated_at": "2026-09-10T00:00:00Z"},
        "noi": {"status": "Accepted"},
        "gop": {"status": "Resolved"},
        "occupancy": {"status": "Open"},
    }


def test_diligence_alone_does_not_trigger_memo_layer() -> None:
    """``memo_diligence`` is workspace state, not memo prose — the memo stays
    byte-identical when it is the only memo_* key."""
    from app.memo_overrides import apply_memo_overrides, has_memo_overrides

    ov = {"memo_diligence": {"noi": {"status": "Resolved"}}}
    sections = _sections()
    assert not has_memo_overrides(ov)
    assert apply_memo_overrides(sections, ov) is sections


# ═══════════════════ live export reads the same persisted state ═══════════════════


def test_live_payload_variance_rows_carry_concept_label_and_diligence() -> None:
    from app.api.analysis import VarianceFlagOut, VarianceRawFieldOut
    from app.export.live_payload import _variance_flags_from_out
    from app.memo_overrides import diligence_status

    flags = [
        VarianceFlagOut(
            field="rooms_revenue", rule_id="BROKER_VS_T12_NOI_VARIANCE", severity="Critical",
            actual=12_300_000.0, broker=12_950_000.0, delta=-650_000.0, delta_pct=-0.0528,
            note="Rooms revenue: broker proforma $12,950,000 vs T-12 actual $12,300,000 — broker overstates the T-12 by 5.3%.",
            concept="rooms_revenue", concept_label="Rooms revenue", impact_basis="revenue",
            raw_fields=[
                VarianceRawFieldOut(field="broker_proforma.rooms_revenue_usd", severity="Warn"),
                VarianceRawFieldOut(field="broker.rooms_revenue", severity="Critical"),
            ],
        ),
        VarianceFlagOut(
            field="noi", rule_id="BROKER_VS_T12_NOI_VARIANCE", severity="Critical",
            actual=4_181_000.0, broker=5_200_000.0, delta=-1_019_000.0, delta_pct=-0.2437,
            concept="noi", concept_label="NOI", impact_basis="noi",
        ),
    ]
    overrides = {
        "memo_diligence": {"rooms_revenue": {"status": "Resolved", "note": "Tied to the T-12"}},
    }
    rows = _variance_flags_from_out(flags, diligence=diligence_status(overrides))

    assert rows[0]["metric"] == "Rooms revenue"  # concept_label, not a humanised path
    assert rows[0]["concept"] == "rooms_revenue"
    assert rows[0]["impact_basis"] == "revenue"
    assert rows[0]["raw_fields"] == ["broker_proforma.rooms_revenue_usd", "broker.rooms_revenue"]
    assert rows[0]["diligence_status"] == "Resolved"
    assert rows[0]["diligence_note"] == "Tied to the T-12"
    # Absent from the persisted map → Open (never fabricated as resolved).
    assert rows[1]["diligence_status"] == "Open"
    assert "diligence_note" not in rows[1]
    assert rows[1]["impact_basis"] == "noi"


def test_live_export_verdict_follows_confirmed_recommendation() -> None:
    from app.export.live_payload import _build_memo

    def build(overrides: dict) -> dict:
        return _build_memo(
            live_sections=[{"section_id": "recommendation", "title": "Recommendation", "body": "gen rationale", "citations": []}],
            overrides=overrides,
            property_name="Harbor House",
            deal_name="Deal",
            location=None,
            keys=None,
            deal_stage=None,
            ai_confidence=None,
            engines_run=[],
            documents_reviewed=[],
        )

    # Unconfirmed (legacy or freshly selected) → canonical pending wording,
    # never the old "In Review", never the selected-but-unconfirmed verdict.
    memo = build({"memo_recommendation_override": "Proceed"})
    assert memo["header"]["recommendation"] == "Pending analyst decision"
    assert memo["sections"][0]["body"].startswith("IC recommendation: Pending analyst decision.")
    assert build({})["header"]["recommendation"] == "Pending analyst decision"

    # Confirmed → the analyst's verdict, in the header and the memo body.
    memo = build({"memo_recommendation_override": "Proceed", "memo_recommendation_confirmed": True})
    assert memo["header"]["recommendation"] == "Proceed"
    assert memo["sections"][0]["body"].startswith("IC recommendation: Proceed.")


def test_excel_variance_sheet_renders_diligence_and_technical_detail(tmp_path) -> None:
    from openpyxl import load_workbook

    from app.export.excel import _build_variance
    from openpyxl import Workbook

    wb = Workbook()
    _build_variance(
        wb,
        {
            "variance_flags": [
                {
                    "flag_id": "VF-001", "severity": "CRITICAL", "metric": "Rooms revenue",
                    "rule_id": "BROKER_VS_T12_NOI_VARIANCE", "broker_value": 12_950_000.0,
                    "t12_value": 12_300_000.0, "variance_pct": -0.0528,
                    "recommended_action": "Rooms revenue: broker overstates the T-12 by 5.3%.",
                    "raw_fields": ["broker_proforma.rooms_revenue_usd", "broker.rooms_revenue"],
                    "diligence_status": "Resolved", "diligence_note": "Tied to the T-12",
                },
                # Fixture-shaped row (no FON-54a keys) still renders, blanks stay blank.
                {"flag_id": "VF-002", "severity": "INFO", "metric": "Occupancy"},
            ]
        },
    )
    path = tmp_path / "v.xlsx"
    wb.save(path)
    ws = load_workbook(path)["Variance"]
    headers = [c.value for c in ws[1]]
    assert headers[:7] == ["Flag ID", "Severity", "Metric", "Broker / Value", "T-12 / Threshold", "Variance", "Action"]
    assert headers[7:9] == ["Diligence", "Technical detail"]
    assert ws.cell(row=2, column=8).value == "Resolved — Tied to the T-12"
    assert ws.cell(row=2, column=9).value == (
        "rule BROKER_VS_T12_NOI_VARIANCE · fields broker_proforma.rooms_revenue_usd, broker.rooms_revenue"
    )
    assert ws.cell(row=3, column=8).value is None
    assert ws.cell(row=3, column=9).value is None
