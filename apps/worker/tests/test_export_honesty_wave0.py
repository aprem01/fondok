"""Go-live Wave 0, Task 1 — exports and the IC deck never invent a number.

Pins the behaviours the pre-launch audit asked for (audit ids in brackets):

* the deck's sensitivity heat-map is the persisted sensitivity ENGINE
  output, or a slide that says why there is none [A1]; no Kimpton "Next
  Steps" copy [A2];
* every export solves pricing only against the deal's own targets; the
  request models carry no 15% / 1.8x default and the ``/pricing/*`` routes
  refuse without a target [A3];
* the Assumptions sheet, proforma and LOI print "not available" with the
  field name — never 0 / 0% / $0 / "0 guest rooms" [A6 A7 A8];
* the Op-Ratio Provenance sheet is always emitted and says when every ratio
  is a seed [B2]; the unsourced brand recovery table is gone [B4]; exported
  debt figures carry the flat-SOFR provenance note [B9];
* ``or 0`` coalescing is gone from the builders: partnership total equity
  is stated only when both components are [D5 D6 C9];
* the ``/analyze`` stub-job route no longer exists [C8]; there is no demo
  fixture or demo slug in production modules [C1 C13].

Hermetic: per-test SQLite, no LLM, no network.
"""

from __future__ import annotations

import os
import tempfile
import zipfile
from pathlib import Path
from uuid import uuid4

import pytest
from sqlalchemy import text

_TMP_DB = Path(tempfile.gettempdir()) / "fondok-tests-export-honesty-w0.db"
if _TMP_DB.exists():
    _TMP_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_TMP_DB}"

from fondok_schemas.reasons import REFUSAL_GLYPH, ReasonCode  # noqa: E402

from app.export.labels import (  # noqa: E402
    BENCHMARK_DEFAULT_NOTE,
    OP_RATIO_ALL_SEED_NOTE,
)
from app.export.refusals import NOT_AVAILABLE_PREFIX  # noqa: E402
from tests.fixtures.kimpton_export_payload import (  # noqa: E402
    kimpton_deal,
    kimpton_memo,
    kimpton_model,
)


def _pptx_text(path: Path) -> str:
    """Every text run in the deck, concatenated."""
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
    return "\n".join(parts)


def _sheet_values(path: Path, sheet: str) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        return [str(c.value) for row in wb[sheet].iter_rows() for c in row if c.value is not None]
    finally:
        wb.close()


# ═════════════════ A1 / A2 — the IC deck ═════════════════


def test_deck_sensitivity_slide_reads_the_engine_output(tmp_path: Path) -> None:
    from app.export import build_pptx

    out = build_pptx(kimpton_deal(), kimpton_model(), kimpton_memo(), tmp_path / "deck.pptx")
    assert zipfile.is_zipfile(out)
    text_blob = _pptx_text(out)
    # The grid's title names the engine's axes and metric, and the base cell
    # value from the fixture's engine block is on the slide.
    assert "Levered IRR — Exit Cap × RevPAR Growth" in text_blob  # noqa: RUF001 — the slide's title
    assert "23.5%" in text_blob  # fixture base cell 0.2348
    assert "Source: sensitivity engine" in text_blob
    assert "Sensitivity analysis not available" not in text_blob


def test_deck_without_sensitivity_output_states_why(tmp_path: Path) -> None:
    from app.export import build_pptx

    model = kimpton_model()
    model["sensitivity_engine"] = {
        "unavailable_reason": "the sensitivity engine run is failed — division by zero"
    }
    out = build_pptx(kimpton_deal(), model, kimpton_memo(), tmp_path / "deck-failed.pptx")
    text_blob = _pptx_text(out)
    assert (
        "Sensitivity analysis not available: the sensitivity engine run is failed — "
        "division by zero"
    ) in text_blob
    # No synthesized axes from the old builder.
    assert "Exit Cap × RevPAR Growth" not in text_blob  # noqa: RUF001

    model.pop("sensitivity_engine")
    out2 = build_pptx(kimpton_deal(), model, kimpton_memo(), tmp_path / "deck-none.pptx")
    assert "Sensitivity analysis not available: no sensitivity engine output" in _pptx_text(out2)


def test_deck_carries_no_kimpton_next_steps_and_no_default_verdict(tmp_path: Path) -> None:
    from pptx import Presentation

    from app.export import build_pptx

    deal = {"id": "x", "name": "Harbor House", "city": "Tampa, FL", "brand": "", "keys": 150}
    memo = {"header": {"subject_property": "Harbor House"}, "sections": [], "appendix": {}}
    out = build_pptx(deal, {"deal_name": "Harbor House"}, memo, tmp_path / "bare.pptx")
    assert len(Presentation(str(out)).slides) == 8
    text_blob = _pptx_text(out)
    assert "Next Steps" not in text_blob
    assert "Kimpton" not in text_blob and "brand standards" not in text_blob
    assert "PROCEED TO LOI" not in text_blob
    assert "Lifestyle Boutique" not in text_blob and "Miami Beach" not in text_blob
    assert f"{NOT_AVAILABLE_PREFIX}: no IC recommendation confirmed by the analyst" in text_blob
    assert f"AI Confidence: {NOT_AVAILABLE_PREFIX}: no confidence score" in text_blob
    # No "0%" / "$0" anywhere on a deal with no figures.
    assert " 0%" not in text_blob and "$0" not in text_blob


# ═════════════════ A6 / A7 / C9 / D5 / D6 — the workbook ═════════════════


def test_assumptions_sheet_prints_not_available_with_the_field_name(tmp_path: Path) -> None:
    from app.export import build_excel

    model = {
        "deal_name": "Harbor House",
        "investment_engine": {"purchase_price_usd": 10_000_000},
        "debt_engine": {"loan_amount_usd": 6_000_000},
        "returns_engine": {"hold_years": 5},
    }
    out = build_excel("x", model, tmp_path / "assumptions.xlsx")
    vals = _sheet_values(out, "Assumptions")
    assert "10000000" in vals or "10000000.0" in vals
    # A field the deal does not carry names itself — never 0.
    assert f"{NOT_AVAILABLE_PREFIX}: no closing costs on the deal's capital engine output" in vals
    assert f"{NOT_AVAILABLE_PREFIX}: no interest rate on the deal's debt engine output" in vals
    # A whole section the run never produced is one honest row.
    assert f"{NOT_AVAILABLE_PREFIX}: no refinance modelled on the deal's canonical run" in vals
    assert "0" not in vals and "0.0" not in vals


def test_cover_and_returns_sheets_never_coalesce_to_zero(tmp_path: Path) -> None:
    from app.export import build_excel

    out = build_excel("x", {"deal_name": "Harbor House"}, tmp_path / "bare.xlsx")
    for sheet in ("Cover", "Returns", "Partnership", "Sources & Uses", "Operating Proforma",
                  "Debt Schedule"):
        vals = _sheet_values(out, sheet)
        assert "0" not in vals and "0.0" not in vals, (sheet, vals)
        assert any(v.startswith(NOT_AVAILABLE_PREFIX) for v in vals), (sheet, vals)
    assert (
        f"{NOT_AVAILABLE_PREFIX}: no monthly debt schedule on the deal's debt engine output"
        in _sheet_values(out, "Debt Schedule")
    )


def test_proforma_missing_years_are_not_zero(tmp_path: Path) -> None:
    from app.export import build_excel
    from app.export.live_payload import _build_proforma_and_cf

    def rev(n: int, total: float) -> dict:
        return {"year": n, "rooms_revenue": total * 0.75, "fb_revenue": total * 0.2,
                "other_revenue": total * 0.05, "resort_fees": 0.0, "total_revenue": total}

    def exp(n: int, total: float) -> dict:
        return {"year": n, "dept_expenses": {"total": total * 0.3}, "undistributed": {"total": total * 0.2},
                "fixed_charges": {"total": total * 0.08}, "mgmt_fee": total * 0.03,
                "ffe_reserve": total * 0.04, "noi": total * 0.35, "noi_institutional": total * 0.39}

    proforma, cf, noi_y1 = _build_proforma_and_cf(
        {"years": [rev(1, 10_000_000.0), rev(2, 10_500_000.0)]},
        {"years": [exp(1, 10_000_000.0), exp(2, 10_500_000.0)]},
        {},  # no debt engine output at all
    )
    by_label = {r["label"]: r for r in proforma["lines"]}
    # Years the run did not project are None, never 0.
    assert by_label["Total Revenue"]["y3"] is None
    assert by_label["Total Revenue"]["y1"] == 10_000
    # Debt service unknown → not 0, and cash flow after debt is not stated.
    assert by_label["Debt Service"]["y1"] is None
    assert by_label["Cash Flow After Debt"]["y1"] is None
    assert cf == {}
    assert noi_y1 == pytest.approx(3_900_000.0)

    out = build_excel("x", {"deal_name": "H", "p_and_l_engine_proforma": proforma}, tmp_path / "pf.xlsx")
    vals = _sheet_values(out, "Operating Proforma")
    assert f"{NOT_AVAILABLE_PREFIX}: Year 3 not projected on the deal's run" in vals
    assert "0" not in vals and "0.0" not in vals


def test_partnership_total_equity_is_not_summed_over_a_missing_component() -> None:
    from app.export.live_payload import _build_partnership

    out = _build_partnership({"lp": {"contributed_equity": 9_000_000.0}, "gp": {"irr": 0.3}})
    assert out["lp_equity_usd"] == 9_000_000.0
    assert "gp_equity_usd" not in out
    assert "total_equity_usd" not in out
    assert out["total_equity_missing"] == (
        "GP contributed equity not on the partnership engine output"
    )
    both = _build_partnership(
        {"lp": {"contributed_equity": 9_000_000.0}, "gp": {"contributed_equity": 1_000_000.0}},
        {"pref_rate": 0.08, "waterfall": [
            {"label": "Tier 1", "hurdle_rate": 0.12, "gp_split": 0.2, "lp_split": 0.8},
        ]},
    )
    assert both["total_equity_usd"] == 10_000_000.0
    assert both["lp_pref_pct"] == 0.08
    assert both["waterfall_tiers"] == [{"label": "Tier 1", "hurdle_rate": 0.12, "gp_split": 0.2}]


def test_partnership_sheet_names_the_missing_component(tmp_path: Path) -> None:
    from app.export import build_excel

    model = {"deal_name": "H", "partnership_engine": {
        "structure": "GP/LP waterfall", "lp_equity_usd": 9_000_000.0,
        "total_equity_missing": "GP contributed equity not on the partnership engine output",
    }}
    vals = _sheet_values(build_excel("x", model, tmp_path / "p.xlsx"), "Partnership")
    assert f"{NOT_AVAILABLE_PREFIX}: GP contributed equity not on the partnership engine output" in vals
    assert "Distribution Per Year (estimate)" not in vals


# ═════════════════ B2 / B4 / B9 ═════════════════


def test_op_ratio_sheet_and_memo_say_when_every_ratio_is_a_seed(tmp_path: Path) -> None:
    from app.export import build_excel
    from app.export.live_payload import _build_op_ratio_provenance
    from app.export.memo_pdf import _render_op_ratio_provenance

    base = {
        "rooms_dept_pct": 0.27, "admin_pct": 0.08,
        "__sources__": {"rooms_dept_pct": "seed", "admin_pct": "seed"},
    }
    op_prov = _build_op_ratio_provenance(base)
    assert op_prov["all_seed"] is True and op_prov["note"] == OP_RATIO_ALL_SEED_NOTE
    assert len(op_prov["lines"]) == 2

    out = build_excel("x", {"deal_name": "H", "op_ratio_provenance": op_prov}, tmp_path / "op.xlsx")
    vals = _sheet_values(out, "Op-Ratio Provenance")
    assert OP_RATIO_ALL_SEED_NOTE in vals
    assert "Rooms Dept Exp %" in vals

    html_str = _render_op_ratio_provenance(op_prov)
    assert "Operating Ratios" in html_str and OP_RATIO_ALL_SEED_NOTE in html_str

    # A grounded deal carries no note.
    grounded = _build_op_ratio_provenance(
        {"rooms_dept_pct": 0.25, "__sources__": {"rooms_dept_pct": "t12_actual"}}
    )
    assert grounded["all_seed"] is False and "note" not in grounded


def test_brand_recovery_multiplier_table_is_gone(tmp_path: Path) -> None:
    from app.export import build_excel

    vals = _sheet_values(build_excel("x", kimpton_model(), tmp_path / "k.xlsx"), "Renovation Plan")
    assert "BRAND RECOVERY MULTIPLIERS" not in vals
    assert not any("Marriott" in v or "Hyatt" in v or "Kimpton / Lifestyle" in v for v in vals)


def test_flat_sofr_note_travels_with_exported_debt_figures(tmp_path: Path) -> None:
    from app.export import build_excel, build_pptx
    from app.export.live_payload import _benchmark_default_note, _build_debt
    from app.export.memo_pdf import _render_html

    stack = {"tranches": [
        {"kind": "senior", "rate_type": "floating", "terms_pending": False,
         "benchmark_is_default": True, "benchmark_name": "SOFR", "benchmark_rate": 0.043},
    ]}
    note = _benchmark_default_note(stack)
    assert note == f"Benchmark SOFR 4.30% is {BENCHMARK_DEFAULT_NOTE}"
    # An analyst-entered index is not flagged; a fixed senior is not flagged.
    assert _benchmark_default_note({"tranches": [
        {"rate_type": "floating", "benchmark_is_default": False, "benchmark_rate": 0.039},
    ]}) is None
    assert _benchmark_default_note({"tranches": [{"rate_type": "fixed", "benchmark_is_default": None}]}) is None

    debt_block = _build_debt({"loan_amount": 1.0, "debt_stack": stack}, {})
    assert debt_block["benchmark_note"] == note

    model = kimpton_model()
    model["debt_engine"]["benchmark_note"] = note
    out = build_excel("x", model, tmp_path / "sofr.xlsx")
    for sheet in ("Assumptions", "Debt Schedule", "Returns", "Cover"):
        assert note in _sheet_values(out, sheet), sheet
    assert note in _render_html(kimpton_memo(), model)
    deck = build_pptx(kimpton_deal(), model, kimpton_memo(), tmp_path / "sofr.pptx")
    assert note in _pptx_text(deck)


def test_debt_schedule_sheet_prints_the_engine_schedule(tmp_path: Path) -> None:
    from app.export import build_excel

    model = kimpton_model()
    vals = _sheet_values(build_excel("x", model, tmp_path / "ds.xlsx"), "Debt Schedule")
    first = model["debt_engine"]["monthly_schedule"][0]
    assert str(first["payment"]) in vals and str(first["ending_balance"]) in vals
    # 60 engine months + header + totals.
    assert "60" in vals and "Totals" in vals


def test_pricing_sheet_renders_max_price_without_a_grid_and_states_why(tmp_path: Path) -> None:
    """EM-only payload shape: max_price present, grid absent, reason present —
    the workbook keeps the solved EM price and prints the reason; the memo
    section is the same sentence."""
    from app.export import build_excel
    from app.export.labels import IRR_GRID_NO_TARGET_NOTE
    from app.export.memo_pdf import _render_max_price_callout, _render_sensitivity_grid

    max_price = {
        "target_irr": None, "target_em": 1.9, "max_price_for_irr": None,
        "max_price_for_em": 38_500_000.0, "binding_constraint": "em",
        "final_price_per_key": 256_666.67, "max_price": 38_500_000.0,
        "irr_status": "not_requested", "em_status": "converged",
    }
    model = {"deal_name": "H", "max_price": max_price,
             "sensitivity_grid_unavailable": IRR_GRID_NO_TARGET_NOTE}
    vals = _sheet_values(build_excel("x", model, tmp_path / "em.xlsx"), "Pricing Sensitivity")
    assert "MAX-PRICE FINDINGS" in vals and ("38500000" in vals or "38500000.0" in vals)
    assert IRR_GRID_NO_TARGET_NOTE in vals
    assert f"{NOT_AVAILABLE_PREFIX}: no IRR target set on this deal" in vals
    assert not any(v.startswith("LEVERED IRR") or v == "EQUITY MULTIPLE" for v in vals)
    assert "0" not in vals and "0.0" not in vals

    assert IRR_GRID_NO_TARGET_NOTE in _render_sensitivity_grid(None, IRR_GRID_NO_TARGET_NOTE)
    assert _render_sensitivity_grid(None, None) == ""
    callout = _render_max_price_callout(max_price)
    assert "Max price for 1.90x EM:</strong> $38,500,000" in callout
    assert "IRR:</strong>" not in callout
    # A non-converged search is described, never priced at the bracket end.
    unreachable = dict(max_price, em_status="unreachable", max_price=None, max_price_for_em=None)
    assert "no price down to half the basis clears this hurdle" in _render_max_price_callout(unreachable)


def test_pricing_sheet_without_a_grid_reason_refuses_rather_than_guessing(tmp_path: Path) -> None:
    """A payload with a max price but neither a grid nor a stated reason must
    not be told "no IRR target" — it gets a typed not-available."""
    from app.export import build_excel
    from app.export.labels import IRR_GRID_NO_TARGET_NOTE

    model = {"deal_name": "H", "max_price": {
        "target_irr": 0.14, "target_em": None, "max_price_for_irr": 40_000_000.0,
        "max_price_for_em": None, "binding_constraint": "irr", "final_price_per_key": 200_000.0,
        "max_price": 40_000_000.0, "irr_status": "converged", "em_status": "not_requested",
    }}
    vals = _sheet_values(build_excel("x", model, tmp_path / "no-reason.xlsx"), "Pricing Sensitivity")
    assert f"{NOT_AVAILABLE_PREFIX}: no IRR sensitivity grid on the export payload" in vals
    assert IRR_GRID_NO_TARGET_NOTE not in vals
    assert "40000000" in vals or "40000000.0" in vals


def test_sensitivity_axis_ticks_do_not_collide_on_rate_axes() -> None:
    from app.export.presentation import _fmt_sens_axis

    rates = [0.0600, 0.0650, 0.0700, 0.0750, 0.0800]
    for variable in ("exit_cap_rate", "interest_rate", "ltv"):
        ticks = [_fmt_sens_axis(variable, v) for v in rates]
        assert len(set(ticks)) == len(ticks), (variable, ticks)
    assert _fmt_sens_axis("interest_rate", 0.0675) == "6.75%"
    assert _fmt_sens_axis("revpar_growth", 0.03) == "3.0%"
    assert _fmt_sens_axis("hold_years", 7.0) == "7"
    assert _fmt_sens_axis("purchase_price", 36_400_000.0) == "$36.4M"


def test_export_helpers_exist_once() -> None:
    """One implementation each — the builders bind precision, never re-define."""
    import inspect

    from app.export import excel, live_payload, memo_pdf, presentation, refusals

    for mod in (excel, memo_pdf, presentation, live_payload):
        src = inspect.getsource(mod)
        for name in ("def _as_float", "def _sum_strict", "def _num(", "def _fmt_usd(",
                     "def _fmt_pct(", "def _fmt_mult(", "def _fmt_usd_m("):
            if mod is presentation and name.startswith("def _fmt"):
                continue  # one-line precision wrappers over the shared helpers
            assert name not in src, (mod.__name__, name)
    assert memo_pdf._fmt_pct is refusals.fmt_pct and excel._as_float is refusals.as_float
    assert live_payload._num is refusals.as_float


# ═════════════════ A3 / A8 / C8 — the routes ═════════════════


def _h(tenant: str) -> dict[str, str]:
    return {"X-Tenant-Id": tenant}


async def _seed(session, *, tenant: str, keys: int | None, target_irr: float | None) -> str:
    deal_id = str(uuid4())
    await session.execute(
        text(
            "INSERT INTO deals (id, tenant_id, name, city, keys, target_irr, field_overrides) "
            "VALUES (:id, :tenant, 'Harbor House', 'Tampa, FL', :keys, :irr, '{}')"
        ),
        {"id": deal_id, "tenant": tenant, "keys": keys, "irr": target_irr},
    )
    await session.commit()
    return deal_id


@pytest.mark.asyncio
async def test_pricing_request_models_carry_no_default_hurdle() -> None:
    import inspect

    from app.api import analysis

    for name in ("_SensitivityRequest", "_LOIRequest", "_MaxPriceTargetsBody"):
        model = getattr(analysis, name)
        assert model.model_fields["target_irr"].default is None, name
        assert model.model_fields["target_em"].default is None, name
    assert not hasattr(analysis, "_MaxPriceRequest")
    src = inspect.getsource(analysis)
    assert "default=0.15" not in src and "default=1.8" not in src


@pytest.mark.asyncio
async def test_sensitivity_and_loi_routes_refuse_without_a_target() -> None:
    from httpx import ASGITransport, AsyncClient

    from app.api.analysis import SENSITIVITY_NO_IRR_TARGET_MESSAGE
    from app.database import get_session_factory
    from app.engines.price_solver import NO_TARGET_MESSAGE
    from app.main import app
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    tenant = str(uuid4())
    async with get_session_factory()() as session:
        deal_id = await _seed(session, tenant=tenant, keys=150, target_irr=None)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        r = await c.post(f"/analysis/{deal_id}/pricing/sensitivity", json={}, headers=_h(tenant))
        assert r.status_code == 422, r.text
        assert r.json()["detail"] == NO_TARGET_MESSAGE
        # Only an EM target → the grid still has no IRR hurdle to sweep against.
        r = await c.post(
            f"/analysis/{deal_id}/pricing/sensitivity", json={"target_em": 1.9}, headers=_h(tenant)
        )
        assert r.status_code == 422, r.text
        assert r.json()["detail"] == SENSITIVITY_NO_IRR_TARGET_MESSAGE
        r = await c.post(f"/analysis/{deal_id}/pricing/loi", json={}, headers=_h(tenant))
        assert r.status_code == 422, r.text
        assert r.json()["detail"] == NO_TARGET_MESSAGE


@pytest.mark.asyncio
async def test_loi_route_refuses_when_the_deal_has_no_key_count() -> None:
    from httpx import ASGITransport, AsyncClient

    from app.api.analysis import LOI_MISSING_FACT_MESSAGE
    from app.database import get_session_factory
    from app.main import app
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    tenant = str(uuid4())
    async with get_session_factory()() as session:
        deal_id = await _seed(session, tenant=tenant, keys=None, target_irr=0.14)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        r = await c.post(f"/analysis/{deal_id}/pricing/loi", json={}, headers=_h(tenant))
        assert r.status_code == 422, r.text
        assert r.json()["detail"] == LOI_MISSING_FACT_MESSAGE.format(what="key count")
        assert "0 guest rooms" not in r.text and "$0" not in r.text


@pytest.mark.asyncio
async def test_analyze_stub_route_is_gone() -> None:
    from httpx import ASGITransport, AsyncClient

    from app.database import get_session_factory
    from app.main import app
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    tenant = str(uuid4())
    async with get_session_factory()() as session:
        deal_id = await _seed(session, tenant=tenant, keys=150, target_irr=None)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        r = await c.post(f"/analysis/{deal_id}/analyze", json={}, headers=_h(tenant))
    assert r.status_code in (404, 405), r.text
    assert "stub-job" not in r.text
    assert not any(getattr(route, "path", "").endswith("/analyze") for route in app.routes)


# ═════════════════ C1 / C13 — no fixture, no demo path ═════════════════


def test_no_fixture_or_demo_path_remains_in_export_modules() -> None:
    import importlib.util
    import inspect

    from app.api import export as export_api
    from app.export import excel, labels, live_payload, memo_pdf, presentation

    assert importlib.util.find_spec("app.export.fixtures") is None
    for mod in (export_api, excel, live_payload, memo_pdf, presentation):
        src = inspect.getsource(mod)
        assert "load_demo_payload" not in src and "DEMO_EXPORT_SLUGS" not in src, mod.__name__
    assert not hasattr(labels, "LEGACY_NOI")
    assert "Net Operating Income" not in labels.NOI_HEADLINE_LABELS
    assert REFUSAL_GLYPH == "—" and ReasonCode.AWAITING_ANALYST.value == "awaiting_analyst"
