"""Regression baselines — list runs + field-level run compare (2026-10-16).

Sam: "preserve the current Kimpton model outputs so we can distinguish
expected changes from unintended calculation discrepancies during regression
testing." ``engine_outputs`` is append-only, so earlier runs ARE the
baselines; these tests lock:

* ``GET /deals/{id}/engines/runs`` lists one entry per run, tenant-scoped,
  with canonical / full-chain / scenario flags and the ``?before=`` default
  baseline pick;
* ``GET /deals/{id}/engines/compare`` diffs headline fields (abs + % delta,
  provenance per side), omits unchanged fields unless ``?all=1``, and 404s
  for another tenant or an unknown run;
* ``GET /deals/{id}/engines/compare/xlsx`` returns a real workbook;
* the pure differ's math and the extractor's coverage of Sam's checklist.
"""

from __future__ import annotations

import io
import json
import os
import tempfile
from pathlib import Path
from uuid import uuid4

import pytest
from sqlalchemy import text

_TMP_DB = Path(tempfile.gettempdir()) / "fondok-tests-run-compare.db"
if _TMP_DB.exists():
    _TMP_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_TMP_DB}"
os.environ.setdefault("ALLOW_TENANT_HEADER_WITHOUT_JWT", "1")


@pytest.fixture(autouse=True)
async def _reset_db() -> None:
    from app.database import get_session_factory
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    async with get_session_factory()() as session:
        for tbl in ("engine_outputs", "scenarios", "deals"):
            try:
                await session.execute(text(f"DELETE FROM {tbl}"))
            except Exception:  # noqa: BLE001
                pass
        await session.commit()
    yield


async def _seed_deal(session, *, deal_id: str, tenant_id: str) -> None:
    await session.execute(
        text(
            "INSERT INTO deals (id, tenant_id, name, field_overrides) "
            "VALUES (:id, :tenant, :name, '{}')"
        ),
        {"id": deal_id, "tenant": tenant_id, "name": "Compare Test"},
    )
    await session.commit()


async def _insert_row(
    session,
    *,
    deal_id: str,
    tenant_id: str,
    run_id: str,
    engine: str,
    outputs: dict | None,
    started: str,
    completed: str | None,
    status: str = "complete",
) -> None:
    await session.execute(
        text(
            "INSERT INTO engine_outputs (id, deal_id, tenant_id, run_id, "
            "engine_name, status, inputs, outputs, error, started_at, "
            "completed_at, runtime_ms) VALUES (:id, :deal, :tenant, :run, "
            ":engine, :status, NULL, :outputs, NULL, :started, :completed, 1)"
        ),
        {
            "id": str(uuid4()),
            "deal": deal_id,
            "tenant": tenant_id,
            "run": run_id,
            "engine": engine,
            "status": status,
            "outputs": json.dumps(outputs) if outputs is not None else None,
            "started": started,
            "completed": completed,
        },
    )


def _client():
    from httpx import ASGITransport, AsyncClient

    from app.main import app

    return AsyncClient(transport=ASGITransport(app=app), base_url="http://t")


# ───────────────────────────── pure differ ─────────────────────────────


def test_diff_math_and_unchanged_omitted() -> None:
    from app.services.run_compare import diff_fields

    def f(key, value, unit="usd"):
        return {"key": key, "engine": "returns", "group": "Returns",
                "label": key, "unit": unit, "value": value,
                "provenance": {"state": "calculated", "source": None}}

    base = [f("a", 100.0), f("b", 0.25, "pct"), f("c", 0.0), f("d", 5.0),
            f("e", None), f("g", 7.0)]
    head = [f("a", 110.0), f("b", 0.25, "pct"), f("c", 3.0), f("d", None),
            f("e", 2.0), f("g", 7.0 + 1e-12)]
    rows = {r["key"]: r for r in diff_fields(base, head)}

    assert set(rows) == {"a", "c", "d", "e"}  # b and g unchanged → omitted
    assert rows["a"]["status"] == "changed"
    assert rows["a"]["delta_abs"] == pytest.approx(10.0)
    assert rows["a"]["delta_pct"] == pytest.approx(0.10)
    assert rows["a"]["base_provenance"]["state"] == "calculated"
    # Zero base: abs delta yes, % delta refused with a reason (never invented).
    assert rows["c"]["delta_abs"] == pytest.approx(3.0)
    assert rows["c"]["delta_pct"] is None
    assert "base is 0" in rows["c"]["delta_reason"]
    assert rows["d"]["status"] == "removed" and rows["d"]["delta_abs"] is None
    assert rows["d"]["delta_reason"] == "not in head run"
    assert rows["e"]["status"] == "added"

    every = {r["key"]: r for r in diff_fields(base, head, include_unchanged=True)}
    assert set(every) == {"a", "b", "c", "d", "e", "g"}
    assert every["b"]["status"] == "unchanged"
    assert every["b"]["delta_abs"] == 0.0 and every["b"]["delta_pct"] == 0.0


def test_extractor_covers_sams_checklist_on_the_golden_baseline() -> None:
    """Every headline Sam named is present on the committed Kimpton baseline."""
    from app.services.run_compare import extract_headline_fields

    repo = Path(__file__).resolve().parents[3]
    doc = json.loads(
        (repo / "evals/baselines/kimpton-golden-2026-10-16.json").read_text()
    )
    fields = extract_headline_fields(
        {k: v["outputs"] for k, v in doc["engines"].items()}
    )
    by_key = {f["key"]: f for f in fields}
    for key in (
        "revenue.years[0].occupancy", "revenue.years[0].adr",
        "revenue.years[0].revpar", "revenue.years[0].total_revenue",
        "expense.years[0].dept_expenses.total",
        "expense.years[0].undistributed.total", "expense.years[0].mgmt_fee",
        "expense.years[0].ffe_reserve", "expense.years[0].gop",
        "expense.years[0].noi_institutional", "expense.years[0].noi",
        "expense.stabilization.stabilized_year",
        "expense.stabilization.stabilized_noi_before_reserve",
        "capital.purchase_price", "capital.total_capital",
        "capital.equity_amount", "capital.renovation_total_usd",
        "debt.loan_amount", "debt.schedule[0].dscr", "debt.schedule[4].dscr",
        "returns.levered_irr", "returns.unlevered_irr",
        "returns.equity_multiple", "returns.gross_sale_price",
        "returns.terminal_noi", "partnership.lp.irr", "partnership.gp.irr",
    ):
        assert key in by_key, key
        assert by_key[key]["value"] is not None, key
    assert by_key["capital.purchase_price"]["value"] == pytest.approx(36_400_000)
    assert by_key["capital.total_capital"]["value"] == pytest.approx(43_262_900)
    # Provenance is carried through where the engine records it.
    assert by_key["revenue.years[0].adr"]["provenance"]["formula"]
    # Groups render in engine order.
    engines = [f["engine"] for f in fields]
    assert engines.index("revenue") < engines.index("expense") < engines.index(
        "capital") < engines.index("debt") < engines.index("returns") < engines.index(
        "partnership")


# ─────────────────────────── endpoints ───────────────────────────


@pytest.mark.asyncio
async def test_runs_list_and_compare_real_chain_runs_tenant_scoped() -> None:
    from app.database import get_session_factory
    from app.services.engine_runner import run_all_engines

    deal_id = str(uuid4())
    tenant = str(uuid4())
    other = str(uuid4())
    base_run, head_run = str(uuid4()), str(uuid4())
    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        base = await run_all_engines(
            session, deal_id=deal_id, tenant_id=tenant, run_id=base_run
        )
        head = await run_all_engines(
            session, deal_id=deal_id, tenant_id=tenant, run_id=head_run,
            overrides={"starting_occupancy": 0.70},
        )

    async with _client() as c:
        r = await c.get(f"/deals/{deal_id}/engines/runs",
                        headers={"X-Tenant-Id": tenant})
        assert r.status_code == 200, r.text
        body = r.json()
        assert [x["run_id"] for x in body["runs"]] == [head_run, base_run]
        assert body["canonical_run_id"] == head_run
        assert body["runs"][0]["is_canonical"] is True
        assert body["runs"][1]["is_canonical"] is False
        assert all(x["is_full_chain"] and x["status"] == "complete"
                   for x in body["runs"])
        assert "cash_flow" in body["runs"][0]["engines"]

        # Default baseline: newest complete run completed before the head
        # run started.
        cutoff = body["runs"][0]["started_at"]
        r = await c.get(f"/deals/{deal_id}/engines/runs",
                        params={"before": cutoff},
                        headers={"X-Tenant-Id": tenant})
        assert r.json()["baseline_run_id"] == base_run

        r = await c.get(f"/deals/{deal_id}/engines/compare",
                        params={"base": base_run, "head": "latest"},
                        headers={"X-Tenant-Id": tenant})
        assert r.status_code == 200, r.text
        cmp = r.json()
        assert cmp["base"]["run_id"] == base_run
        assert cmp["head"]["run_id"] == head_run
        assert cmp["fields_changed"] == len(cmp["fields"]) > 0
        assert all(f["status"] != "unchanged" for f in cmp["fields"])
        rows = {f["key"]: f for f in cmp["fields"]}
        occ = rows["revenue.years[0].occupancy"]
        assert occ["base_value"] == pytest.approx(0.762)
        assert occ["head_value"] == pytest.approx(0.70)
        assert occ["delta_abs"] == pytest.approx(0.70 - 0.762)
        assert occ["delta_pct"] == pytest.approx((0.70 - 0.762) / 0.762)
        assert occ["base_provenance"]["formula"]
        irr = rows["returns.levered_irr"]
        assert irr["base_value"] == pytest.approx(
            base["returns"]["outputs"]["levered_irr"])
        assert irr["head_value"] == pytest.approx(
            head["returns"]["outputs"]["levered_irr"])
        # Purchase price did not move → omitted …
        assert "capital.purchase_price" not in rows
        # … but present with ?all=1.
        r = await c.get(f"/deals/{deal_id}/engines/compare",
                        params={"base": base_run, "head": head_run, "all": 1},
                        headers={"X-Tenant-Id": tenant})
        every = {f["key"]: f for f in r.json()["fields"]}
        assert every["capital.purchase_price"]["status"] == "unchanged"
        assert r.json()["fields_compared"] == len(every)

        # Same run against itself: nothing changed.
        r = await c.get(f"/deals/{deal_id}/engines/compare",
                        params={"base": base_run, "head": base_run},
                        headers={"X-Tenant-Id": tenant})
        assert r.json()["fields"] == [] and r.json()["fields_changed"] == 0

        # Another tenant: 404 on all three routes, required params present.
        for path, params in (
            ("runs", {}),
            ("compare", {"base": base_run, "head": head_run}),
            ("compare/xlsx", {"base": base_run, "head": head_run}),
        ):
            r = await c.get(f"/deals/{deal_id}/engines/{path}", params=params,
                            headers={"X-Tenant-Id": other})
            assert r.status_code == 404, (path, r.status_code, r.text)

        # Unknown run on the right tenant: 404, not an empty diff.
        r = await c.get(f"/deals/{deal_id}/engines/compare",
                        params={"base": str(uuid4())},
                        headers={"X-Tenant-Id": tenant})
        assert r.status_code == 404
        # Bad timestamp: 422.
        r = await c.get(f"/deals/{deal_id}/engines/runs",
                        params={"before": "yesterday"},
                        headers={"X-Tenant-Id": tenant})
        assert r.status_code == 422


@pytest.mark.asyncio
async def test_run_ids_of_another_tenant_never_resolve() -> None:
    """A run id from tenant B's deal cannot be compared through tenant A's deal."""
    from app.database import get_session_factory

    deal_a, deal_b = str(uuid4()), str(uuid4())
    ta, tb = str(uuid4()), str(uuid4())
    run_a, run_b = str(uuid4()), str(uuid4())
    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_a, tenant_id=ta)
        await _seed_deal(session, deal_id=deal_b, tenant_id=tb)
        await _insert_row(session, deal_id=deal_a, tenant_id=ta, run_id=run_a,
                          engine="returns", outputs={"levered_irr": 0.2},
                          started="2026-10-15T20:00:00+00:00",
                          completed="2026-10-15T20:00:05+00:00")
        await _insert_row(session, deal_id=deal_b, tenant_id=tb, run_id=run_b,
                          engine="returns", outputs={"levered_irr": 0.9},
                          started="2026-10-15T20:00:00+00:00",
                          completed="2026-10-15T20:00:05+00:00")
        await session.commit()

    async with _client() as c:
        r = await c.get(f"/deals/{deal_a}/engines/compare",
                        params={"base": run_a, "head": run_b},
                        headers={"X-Tenant-Id": ta})
        assert r.status_code == 404
        r = await c.get(f"/deals/{deal_a}/engines/runs",
                        headers={"X-Tenant-Id": ta})
        assert [x["run_id"] for x in r.json()["runs"]] == [run_a]


@pytest.mark.asyncio
async def test_latest_head_is_404_when_the_deal_has_no_canonical_run() -> None:
    """``head=latest`` means the canonical run — the one the tabs serve and
    the one the panel labels "Latest (canonical)". A deal that has never
    completed a full 8-engine chain has no such run, so ``latest`` must say
    so (404 naming the gap) instead of silently diffing against the newest
    partial run. An explicit run id still compares."""
    from app.database import get_session_factory

    deal, tenant = str(uuid4()), str(uuid4())
    partial_old, partial_new = str(uuid4()), str(uuid4())
    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal, tenant_id=tenant)
        # Two single-engine runs: neither is a full chain → no canonical run.
        await _insert_row(session, deal_id=deal, tenant_id=tenant, run_id=partial_old,
                          engine="returns", outputs={"levered_irr": 0.20},
                          started="2026-10-15T20:00:00+00:00",
                          completed="2026-10-15T20:00:05+00:00")
        await _insert_row(session, deal_id=deal, tenant_id=tenant, run_id=partial_new,
                          engine="returns", outputs={"levered_irr": 0.25},
                          started="2026-10-16T20:00:00+00:00",
                          completed="2026-10-16T20:00:05+00:00")
        await session.commit()

    async with _client() as c:
        h = {"X-Tenant-Id": tenant}
        runs = (await c.get(f"/deals/{deal}/engines/runs", headers=h)).json()
        assert runs["canonical_run_id"] is None
        assert [x["run_id"] for x in runs["runs"]] == [partial_new, partial_old]
        assert not any(x["is_full_chain"] for x in runs["runs"])

        for head in ("latest", "canonical"):
            for path in ("compare", "compare/xlsx"):
                r = await c.get(f"/deals/{deal}/engines/{path}",
                                params={"base": partial_old, "head": head},
                                headers=h)
                assert r.status_code == 404, (head, path, r.status_code, r.text)
                detail = r.json()["detail"]
                assert "canonical" in detail and partial_new not in detail
        # ``base=latest`` is held to the same rule.
        r = await c.get(f"/deals/{deal}/engines/compare",
                        params={"base": "latest", "head": partial_new}, headers=h)
        assert r.status_code == 404

        # Naming the partial run explicitly still works.
        r = await c.get(f"/deals/{deal}/engines/compare",
                        params={"base": partial_old, "head": partial_new}, headers=h)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["head"]["run_id"] == partial_new
        irr = {f["key"]: f for f in body["fields"]}["returns.levered_irr"]
        assert irr["base_value"] == pytest.approx(0.20)
        assert irr["head_value"] == pytest.approx(0.25)

        # No runs at all: the pre-existing 404 still says so.
        empty = str(uuid4())
        async with get_session_factory()() as session:
            await _seed_deal(session, deal_id=empty, tenant_id=tenant)
        r = await c.get(f"/deals/{empty}/engines/compare",
                        params={"base": "latest", "head": "latest"}, headers=h)
        assert r.status_code == 404
        assert "no engine runs" in r.json()["detail"]


@pytest.mark.asyncio
async def test_before_picks_newest_full_chain_base_run_and_skips_scenarios() -> None:
    from app.database import get_session_factory
    from app.services.engine_runner import CANONICAL_ENGINES

    deal_id, tenant = str(uuid4()), str(uuid4())
    early, late, scen, partial, after = (str(uuid4()) for _ in range(5))

    async def full(session, run_id, ts):
        for name in CANONICAL_ENGINES:
            await _insert_row(session, deal_id=deal_id, tenant_id=tenant,
                              run_id=run_id, engine=name,
                              outputs={"levered_irr": 0.1} if name == "returns" else {},
                              started=f"{ts}:00+00:00", completed=f"{ts}:30+00:00")

    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        await full(session, early, "2026-10-14T10:00")
        await full(session, late, "2026-10-15T21:00")
        await full(session, scen, "2026-10-15T21:20")      # non-base scenario
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant,  # lone re-run
                          run_id=partial, engine="returns", outputs={},
                          started="2026-10-15T21:30:00+00:00",
                          completed="2026-10-15T21:30:05+00:00")
        await full(session, after, "2026-10-16T09:00")
        await session.execute(
            text("INSERT INTO scenarios (id, deal_id, tenant_id, name, is_base, "
                 "overrides, last_run_id) VALUES (:id, :deal, :t, 'Downside', "
                 ":b, '[]', :run)"),
            {"id": str(uuid4()), "deal": deal_id, "t": tenant, "b": False, "run": scen},
        )
        await session.commit()

    async with _client() as c:
        r = await c.get(f"/deals/{deal_id}/engines/runs",
                        params={"before": "2026-10-15T21:40Z"},
                        headers={"X-Tenant-Id": tenant})
        body = r.json()
    assert body["baseline_run_id"] == late
    by_id = {x["run_id"]: x for x in body["runs"]}
    assert by_id[scen]["scenario"]["name"] == "Downside"
    assert by_id[scen]["scenario"]["is_base"] is False
    assert by_id[partial]["is_full_chain"] is False
    assert body["canonical_run_id"] == after
    assert by_id[after]["is_canonical"] is True


@pytest.mark.asyncio
async def test_compare_xlsx_export() -> None:
    from openpyxl import load_workbook

    from app.database import get_session_factory

    deal_id, tenant = str(uuid4()), str(uuid4())
    r1, r2 = str(uuid4()), str(uuid4())
    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant, run_id=r1,
                          engine="returns",
                          outputs={"levered_irr": 0.30, "equity_multiple": 3.0},
                          started="2026-10-15T20:00:00+00:00",
                          completed="2026-10-15T20:00:05+00:00")
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant, run_id=r2,
                          engine="returns",
                          outputs={"levered_irr": 0.27, "equity_multiple": 3.0},
                          started="2026-10-16T20:00:00+00:00",
                          completed="2026-10-16T20:00:05+00:00")
        await session.commit()

    async with _client() as c:
        r = await c.get(f"/deals/{deal_id}/engines/compare/xlsx",
                        params={"base": r1, "head": r2},
                        headers={"X-Tenant-Id": tenant})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert ".xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Changed fields"]
    ws = wb["Changed fields"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    assert rows[0][:8] == ["Engine", "Group", "Field", "Status", "Baseline",
                           "Head", "Δ (abs)", "Δ (%)"]
    data = [row for row in rows[1:] if row[2] == "Levered IRR"]
    assert len(data) == 1
    assert data[0][4] == pytest.approx(0.30)
    assert data[0][5] == pytest.approx(0.27)
    assert data[0][6] == pytest.approx(-0.03)
    assert data[0][7] == pytest.approx(-0.1)
    # Unchanged equity multiple is not in the changed-fields sheet.
    assert not any(row[2] == "Equity multiple" for row in rows)
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["Baseline run"] == r1 and summary["Head run"] == r2


# ──────────────── baseline pick: older scenario runs + failed views ────────────────


async def _full_chain(session, *, deal_id, tenant_id, run_id, ts, extra=()):
    from app.services.engine_runner import CANONICAL_ENGINES

    for name in CANONICAL_ENGINES:
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant_id,
                          run_id=run_id, engine=name,
                          outputs={"levered_irr": 0.1} if name == "returns" else {},
                          started=f"{ts}:00+00:00", completed=f"{ts}:30+00:00")
    for name, status in extra:
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant_id,
                          run_id=run_id, engine=name, outputs=None, status=status,
                          started=f"{ts}:31+00:00", completed=f"{ts}:32+00:00")


async def _audit_scenario_run(session, *, tenant_id, deal_id, run_id, scenario_id):
    """Exactly what ``api/scenarios.py::run_scenario`` writes per scenario run."""
    from app.audit import log_audit

    await log_audit(
        session,
        tenant_id=tenant_id,
        action="engine_run.ran_with_scenario",
        resource_type="engine_run",
        resource_id=run_id,
        output_payload={"engines": [], "engine_count": 0},
        tags=["engine_run", "scenario"],
        metadata={"deal_id": deal_id, "scenario_id": scenario_id},
    )
    await session.commit()


@pytest.mark.asyncio
async def test_older_scenario_run_is_never_the_default_baseline() -> None:
    """Sam's QA flow: base chain 10-13, Downside 10-14, Downside again 10-16.
    ``scenarios.last_run_id`` now names the 10-16 run only, yet the 10-14
    Downside run must still read as a scenario run — the default baseline
    before 10-15 21:40 is the 10-13 BASE run, not the older Downside run."""
    from app.database import get_session_factory

    deal_id, tenant, other_tenant = str(uuid4()), str(uuid4()), str(uuid4())
    base, downside_old, downside_new, foreign = (str(uuid4()) for _ in range(4))
    scenario_id = str(uuid4())

    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=base,
                          ts="2026-10-13T10:00")
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=downside_old,
                          ts="2026-10-14T10:00")
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=downside_new,
                          ts="2026-10-16T10:00")
        # A full-chain run the audit trail tags for ANOTHER tenant's scenario
        # (impossible in practice; proves the audit read is tenant-scoped).
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=foreign,
                          ts="2026-10-14T12:00")
        await session.execute(
            text("INSERT INTO scenarios (id, deal_id, tenant_id, name, is_base, "
                 "overrides, last_run_id) VALUES (:id, :deal, :t, 'Downside', "
                 ":b, '[]', :run)"),
            {"id": scenario_id, "deal": deal_id, "t": tenant, "b": False, "run": downside_new},
        )
        await session.commit()
        await _audit_scenario_run(session, tenant_id=tenant, deal_id=deal_id,
                                  run_id=downside_old, scenario_id=scenario_id)
        await _audit_scenario_run(session, tenant_id=tenant, deal_id=deal_id,
                                  run_id=downside_new, scenario_id=scenario_id)
        await _audit_scenario_run(session, tenant_id=other_tenant, deal_id=deal_id,
                                  run_id=foreign, scenario_id=str(uuid4()))

    async with _client() as c:
        r = await c.get(f"/deals/{deal_id}/engines/runs",
                        params={"before": "2026-10-15T21:40Z"},
                        headers={"X-Tenant-Id": tenant})
        assert r.status_code == 200, r.text
        body = r.json()
    by_id = {x["run_id"]: x for x in body["runs"]}
    # The older Downside run is tagged from the audit trail, by name.
    assert by_id[downside_old]["scenario"] == {
        "id": scenario_id, "name": "Downside", "is_base": False,
    }
    assert by_id[downside_new]["scenario"]["name"] == "Downside"
    # Another tenant's audit row never tags this tenant's run.
    assert by_id[foreign]["scenario"] is None
    assert by_id[base]["scenario"] is None
    # Newest qualifying run before the cutoff is the foreign-tagged-but-untagged
    # full chain (10-14 12:00), NOT the 10-14 Downside run.
    assert body["baseline_run_id"] == foreign

    # With that run gone, the pick falls back to the 10-13 base run.
    from app.services.engine_runner import CANONICAL_ENGINES
    from app.services.run_compare import list_runs, pick_baseline

    async with get_session_factory()() as session:
        await session.execute(text("DELETE FROM engine_outputs WHERE run_id = :r"), {"r": foreign})
        await session.commit()
        runs = await list_runs(session, deal_uuid=deal_id, tenant_id=tenant,
                               canonical_engines=CANONICAL_ENGINES, canonical_run_id=None)
    from app.services.run_compare import parse_ts

    picked = pick_baseline(runs, parse_ts("2026-10-15T21:40Z"))
    assert picked is not None and picked["run_id"] == base


@pytest.mark.asyncio
async def test_deleted_scenario_still_tags_its_runs_as_non_base() -> None:
    """The audit row outlives the scenario: a run of a since-deleted scenario
    stays tagged (name unknown) and is skipped by the default pick."""
    from app.database import get_session_factory

    deal_id, tenant = str(uuid4()), str(uuid4())
    base, orphan = str(uuid4()), str(uuid4())
    gone_scenario = str(uuid4())

    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=base,
                          ts="2026-10-13T10:00")
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=orphan,
                          ts="2026-10-14T10:00")
        await _audit_scenario_run(session, tenant_id=tenant, deal_id=deal_id,
                                  run_id=orphan, scenario_id=gone_scenario)

    async with _client() as c:
        body = (await c.get(f"/deals/{deal_id}/engines/runs",
                            params={"before": "2026-10-15T21:40Z"},
                            headers={"X-Tenant-Id": tenant})).json()
    from app.services.run_compare import DELETED_SCENARIO_LABEL

    by_id = {x["run_id"]: x for x in body["runs"]}
    assert by_id[orphan]["scenario"] == {
        "id": gone_scenario, "name": DELETED_SCENARIO_LABEL, "is_base": False,
    }
    assert body["baseline_run_id"] == base


@pytest.mark.asyncio
async def test_failed_cash_flow_view_does_not_block_a_complete_base_chain() -> None:
    """Status folds over the canonical engines only: all eight base engines
    complete + ``cash_flow`` failed is a complete, full-chain run and a valid
    baseline. A failed CANONICAL engine still makes the run ``partial``."""
    from app.database import get_session_factory

    deal_id, tenant = str(uuid4()), str(uuid4())
    view_failed, canon_failed, view_running = (str(uuid4()) for _ in range(3))

    async with get_session_factory()() as session:
        await _seed_deal(session, deal_id=deal_id, tenant_id=tenant)
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=view_failed,
                          ts="2026-10-14T10:00", extra=[("cash_flow", "failed")])
        # Newer, but a canonical engine failed → partial, never the baseline.
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=canon_failed,
                          ts="2026-10-15T10:00")
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant, run_id=canon_failed,
                          engine="returns", outputs=None, status="failed",
                          started="2026-10-15T10:00:40+00:00",
                          completed="2026-10-15T10:00:41+00:00")
        # Newer still: base chain complete, cash_flow row still running.
        await _full_chain(session, deal_id=deal_id, tenant_id=tenant, run_id=view_running,
                          ts="2026-10-15T12:00")
        await _insert_row(session, deal_id=deal_id, tenant_id=tenant, run_id=view_running,
                          engine="cash_flow", outputs=None, status="running",
                          started="2026-10-15T12:00:31+00:00", completed=None)
        await session.commit()

    async with _client() as c:
        body = (await c.get(f"/deals/{deal_id}/engines/runs",
                            params={"before": "2026-10-15T21:40Z"},
                            headers={"X-Tenant-Id": tenant})).json()
    by_id = {x["run_id"]: x for x in body["runs"]}
    assert by_id[view_failed]["status"] == "complete"
    assert by_id[view_failed]["is_full_chain"] is True
    assert "cash_flow" in by_id[view_failed]["engines"]
    assert by_id[canon_failed]["status"] == "partial"
    assert by_id[view_running]["status"] == "complete"
    assert by_id[view_running]["completed_at"] is not None
    assert body["baseline_run_id"] == view_running


def test_golden_compare_script_exits_nonzero_when_it_cannot_run() -> None:
    """``scripts/compare_golden_baseline.py`` is a report (exit 0 with or
    without differences) but a comparison that could NOT run must not look
    like "no headline differences" to a caller checking the exit code."""
    import subprocess
    import sys
    from pathlib import Path

    script = Path(__file__).resolve().parents[3] / "scripts" / "compare_golden_baseline.py"
    proc = subprocess.run(
        [sys.executable, str(script), "--baseline", "/nonexistent/kimpton-golden.json"],
        capture_output=True, text=True, timeout=300,
    )
    assert proc.returncode == 2, (proc.returncode, proc.stdout[-500:], proc.stderr[-500:])
    assert "could not complete the comparison" in proc.stdout


def test_run_status_folds_over_canonical_rows_only() -> None:
    from app.services.run_compare import _run_status, parse_ts

    canon = {"revenue", "returns"}
    done = parse_ts("2026-10-14T10:00:30Z")
    assert _run_status([("revenue", "complete", done), ("returns", "complete", done),
                        ("cash_flow", "failed", done)], canon) == "complete"
    assert _run_status([("revenue", "complete", done), ("returns", "failed", done)],
                       canon) == "partial"
    assert _run_status([("revenue", "failed", done), ("returns", "failed", done),
                        ("cash_flow", "complete", done)], canon) == "failed"
    assert _run_status([("revenue", "complete", done), ("returns", "running", None)],
                       canon) == "running"
    # No canonical rows at all (a lone view re-run): fold over what exists.
    assert _run_status([("cash_flow", "failed", done)], canon) == "failed"
    assert _run_status([("cash_flow", "complete", done)], canon) == "complete"
