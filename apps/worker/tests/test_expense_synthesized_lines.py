"""B3 (go-live audit) — every Year-1 expense line synthesized at a hotel-type
ratio carries a provenance entry (``seed``, "no T-12 line for <name>;
hotel-type default <ratio>") that reaches the Excel Assumptions sheet and the
memo's op-ratio provenance block."""

from __future__ import annotations

from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.engines.expense import ExpenseEngine, ExpenseEngineInput
from app.engines.fb_revenue import FBRevenueOutput, FBRevenueYear


def _input(**kw) -> ExpenseEngineInput:
    years = [
        FBRevenueYear(
            year=y, rooms_revenue=7_400_000.0, fb_revenue=2_200_000.0,
            other_revenue=400_000.0, total_revenue=10_000_000.0,
        )
        for y in range(1, 4)
    ]
    return ExpenseEngineInput(
        deal_id=uuid4(),
        revenue=FBRevenueOutput(
            deal_id=uuid4(), years=years, fb_ratio_used=0.3, other_ratio_used=0.05
        ),
        hotel_type="select",
        **kw,
    )


def test_no_t12_every_line_is_a_badged_hotel_type_default() -> None:
    out = ExpenseEngine().run(_input())
    by_line = {s.line: s for s in out.synthesized_lines}
    assert set(by_line) == {
        "rooms_dept_expense", "fb_dept_expense", "other_dept_expense",
        "administrative_general", "information_telecom", "sales_marketing",
        "property_operations", "utilities", "property_taxes", "insurance",
        "other_fixed",
    }
    rooms = by_line["rooms_dept_expense"]
    assert rooms.source == "seed"
    assert rooms.ratio == pytest.approx(0.27)
    assert rooms.value == pytest.approx(out.years[0].dept_expenses.rooms)
    assert rooms.reason.startswith(
        "no T-12 line for Rooms department expense; hotel-type default 27.00%"
    )
    utilities = by_line["utilities"]
    assert utilities.ratio == pytest.approx(0.22 * 0.25)
    assert utilities.value == pytest.approx(out.years[0].undistributed.utilities)
    # The Year-1 value's provenance trace carries the same reason, as seed.
    trace = out.provenance["years[0].undistributed.utilities"]
    assert trace.note == utilities.reason
    assert trace.state == "assumption"


def test_t12_lines_are_not_synthesized() -> None:
    actuals = {
        "rooms_dept_expense": 2_000_000.0,
        "administrative_general": 800_000.0,
        "sales_marketing": 600_000.0,
        "property_operations": 500_000.0,
        "utilities": 400_000.0,
        "property_taxes": 300_000.0,
        "insurance": 200_000.0,
    }
    out = ExpenseEngine().run(_input(t12_actuals=actuals))
    lines = {s.line for s in out.synthesized_lines}
    # Actuals are not listed; omitted lines in a grounded category fold to 0
    # (not synthesized); departmental lines the T-12 omits ARE synthesized.
    assert lines == {"fb_dept_expense", "other_dept_expense"}


def test_ratio_override_is_not_called_a_hotel_type_default() -> None:
    out = ExpenseEngine().run(_input(overrides={"rooms_dept_pct": 0.25}))
    rooms = next(s for s in out.synthesized_lines if s.line == "rooms_dept_expense")
    assert rooms.source == "ratio_override"
    assert "hotel-type default" not in rooms.reason
    assert "ratio override 25.00%" in rooms.reason


def test_synthesized_lines_reach_payload_excel_assumptions_and_memo() -> None:
    import json

    from app.export.excel import _build_assumptions, _build_op_ratio_provenance
    from app.export.live_payload import _build_op_ratio_provenance as build_prov
    from app.export.memo_pdf import _render_op_ratio_provenance

    expense = json.loads(ExpenseEngine().run(_input()).model_dump_json())
    prov = build_prov({"__sources__": {}}, expense)
    synth = prov["synthesized_lines"]
    assert len(synth) == 11
    assert all(s["source"] == "seed" for s in synth)
    assert synth[0]["reason"].startswith("no T-12 line for ")

    wb = Workbook()
    _build_assumptions(wb, {"op_ratio_provenance": prov})
    cells = [str(c.value) for r in wb["Assumptions"].iter_rows() for c in r if c.value]
    assert "Year-1 Expense Lines Synthesized from Ratios" in cells
    assert any("no T-12 line for Utilities; hotel-type default" in c for c in cells)

    assert _build_op_ratio_provenance(wb, prov)
    cells = [str(c.value) for r in wb["Op-Ratio Provenance"].iter_rows() for c in r if c.value]
    assert any("no T-12 line for Insurance" in c for c in cells)

    html_out = _render_op_ratio_provenance(prov)
    assert "Year-1 expense lines synthesized from ratios" in html_out
    assert "no T-12 line for Rooms department expense; hotel-type default" in html_out
