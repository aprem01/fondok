"""R-030 (upload formats) + R-067 (P&L benchmark → expense categories).

R-030
  * ``.pptx`` decks parse slide-by-slide (text frames, tables, speaker
    notes) and are accepted at the upload boundary.
  * ``.csv`` exports parse (the allowlist accepted them but no parser was
    registered, so every CSV failed at parse time).
  * Known-but-unparsed formats (.ppt, .key, .doc, .ods, Google Drive
    ``.gsheet`` / ``.gslides`` shortcuts) are rejected up front with a
    "convert it to …" instruction, not the generic allowlist copy.

R-067
  * A CBRE P&L benchmark filename routes to PNL_BENCHMARK, not
    CBRE_HORIZONS.
  * ``services.pnl_benchmark_map`` maps the comp-set column onto the model's
    expense categories on the model's own ratio basis, never inventing a
    value the report did not publish.
  * The market-data ``pnl_benchmark`` block carries those categories.
"""

from __future__ import annotations

import io
import os
import shutil
import tempfile
from pathlib import Path
from uuid import uuid4

import pytest

_TMP_DB = Path(tempfile.gettempdir()) / "fondok-tests-upload-formats.db"
if _TMP_DB.exists():
    _TMP_DB.unlink()
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_TMP_DB}"
os.environ["DOCUMENT_STORAGE_ROOT"] = str(
    Path(tempfile.gettempdir()) / "fondok-tests-upload-formats-storage"
)
os.environ.setdefault("EVALS_MOCK", "true")

_STORAGE_ROOT = Path(os.environ["DOCUMENT_STORAGE_ROOT"])
if _STORAGE_ROOT.exists():
    shutil.rmtree(_STORAGE_ROOT)


# ───────────────────────────── fixtures ─────────────────────────────


def _build_tiny_pptx() -> bytes:
    """A two-slide OM teaser: title + bullet, a table, and speaker notes."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Harbor View Hotel — Offering Memorandum"
    s1.placeholders[1].text = "142 keys · Upper Upscale · Built 1998"
    s1.notes_slide.notes_text_frame.text = "Seller expects pricing near $38M."

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Operating Summary"
    table = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, (label, value) in enumerate(
        [("Metric", "T-12"), ("Occupancy", "78.4%"), ("ADR", "$212.50")]
    ):
        table.cell(r, 0).text = label
        table.cell(r, 1).text = value

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
async def deal_id() -> str:
    from sqlalchemy import text

    from app.config import get_settings
    from app.database import dispose_engine, get_session_factory
    from app.migrations import run_startup_migrations

    await run_startup_migrations()
    settings = get_settings()
    factory = get_session_factory()
    new_id = uuid4()
    async with factory() as session:
        await session.execute(
            text(
                """
                INSERT INTO deals (id, tenant_id, name, status, created_at, updated_at)
                VALUES (:id, :tenant, :name, 'Draft', :ts, :ts)
                """
            ),
            {
                "id": str(new_id),
                "tenant": settings.DEFAULT_TENANT_ID,
                "name": "Formats Hotel",
                "ts": "2026-10-08 00:00:00",
            },
        )
        await session.commit()
    yield str(new_id)
    await dispose_engine()


async def _upload(deal_id: str, name: str, body: bytes, mime: str) -> tuple[int, dict]:
    from httpx import ASGITransport, AsyncClient

    from app.main import app
    from app.storage import reset_raw_store_cache

    reset_raw_store_cache()
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        r = await client.post(
            f"/deals/{deal_id}/documents/upload",
            files=[("files", (name, body, mime))],
        )
    return r.status_code, r.json()[0]


# ───────────────────────────── R-030 parsing ─────────────────────────────


@pytest.mark.asyncio
async def test_pptx_parses_slides_tables_and_notes() -> None:
    from app.extraction import parse_document

    parsed = await parse_document(_build_tiny_pptx(), "Harbor View OM.pptx")
    assert parsed.parser == "python-pptx"
    assert parsed.total_pages == 2
    first, second = parsed.pages
    assert "Harbor View Hotel" in first.text
    assert "142 keys" in first.text
    assert "[Speaker notes] Seller expects pricing near $38M." in first.text
    assert first.metadata["slide_index"] == 1
    assert second.tables == [
        [["Metric", "T-12"], ["Occupancy", "78.4%"], ["ADR", "$212.50"]]
    ]


@pytest.mark.asyncio
async def test_csv_parses_comma_and_semicolon_exports() -> None:
    from app.extraction import parse_document

    comma = "Line,2024,2025\nRooms Revenue,\"1,200,000\",1300000\n,,\nTotal Revenue,1500000,1650000\n"
    parsed = await parse_document(comma.encode("utf-8-sig"), "pnl export.csv")
    assert parsed.parser == "csv"
    assert parsed.total_pages == 1
    page = parsed.pages[0]
    # Blank rows are dropped, quoted thousands stay one cell.
    assert page.tables == [
        [
            ["Line", "2024", "2025"],
            ["Rooms Revenue", "1,200,000", "1300000"],
            ["Total Revenue", "1500000", "1650000"],
        ]
    ]
    assert page.text.splitlines()[1] == "Rooms Revenue\t1,200,000\t1300000"

    semicolon = "Line;2024\nRooms Revenue;1200000\nTotal Revenue;1500000\n"
    parsed2 = await parse_document(semicolon.encode("cp1252"), "eu.csv")
    assert parsed2.pages[0].tables[0][1] == ["Rooms Revenue", "1200000"]


@pytest.mark.parametrize(
    ("filename", "needle"),
    [
        ("deck.ppt", ".pptx"),
        ("deck.key", "Keynote"),
        ("memo.doc", ".docx"),
        ("model.ods", ".xlsx"),
        ("pitch.odp", ".pptx"),
        ("budget.numbers", ".xlsx"),
        ("T12.gsheet", "File → Download → Microsoft Excel"),
        ("OM.gslides", "Microsoft PowerPoint (.pptx)"),
    ],
)
def test_conversion_hint_names_the_fix(filename: str, needle: str) -> None:
    from app.extraction.parser import conversion_hint

    hint = conversion_hint(filename)
    assert hint is not None
    assert needle in hint


def test_supported_formats_have_no_conversion_hint() -> None:
    from app.extraction.parser import conversion_hint
    from app.extraction.registry import get_parser

    for ext in ("pdf", "xls", "xlsx", "xlsm", "csv", "docx", "pptx"):
        assert conversion_hint(f"file.{ext}") is None
        assert get_parser(ext) is not None, ext


@pytest.mark.asyncio
async def test_parse_document_backstops_legacy_ppt() -> None:
    from app.extraction import ParseError, parse_document

    with pytest.raises(ParseError, match=r"re-save the deck as \.pptx"):
        await parse_document(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1junk", "old.ppt")


# ───────────────────────────── R-030 upload boundary ─────────────────────────────


@pytest.mark.asyncio
async def test_pptx_upload_is_accepted(deal_id: str) -> None:
    code, rec = await _upload(
        deal_id,
        "Harbor View OM.pptx",
        _build_tiny_pptx(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
    assert code == 201, rec
    assert rec["status"] != "FAILED"
    assert rec.get("error_kind") != "unsupported_type"


@pytest.mark.asyncio
async def test_pptx_upload_with_stripped_mime_is_accepted(deal_id: str) -> None:
    """Extension allowlist now names .pptx — a broker-stripped MIME still lands."""
    code, rec = await _upload(
        deal_id, "teaser.pptx", _build_tiny_pptx(), "application/octet-stream"
    )
    assert code == 201, rec
    assert rec["status"] != "FAILED"


@pytest.mark.asyncio
async def test_renamed_non_zip_pptx_is_rejected(deal_id: str) -> None:
    code, rec = await _upload(
        deal_id, "fake.pptx", b"MZ\x90\x00not a deck" * 20, "application/octet-stream"
    )
    assert code == 422
    assert rec["error_kind"] == "unsupported_type"


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "mime", "needle"),
    [
        ("OM.key", "application/vnd.apple.keynote", "Keynote"),
        ("OM.ppt", "application/vnd.ms-powerpoint", ".pptx"),
        ("memo.doc", "application/msword", ".docx"),
        ("T12.gsheet", "application/octet-stream", "Google Sheets"),
    ],
)
async def test_unparsed_formats_rejected_with_conversion_hint(
    deal_id: str, filename: str, mime: str, needle: str
) -> None:
    code, rec = await _upload(deal_id, filename, b"\x00\x01\x02\x03" * 64, mime)
    assert code == 422
    assert rec["status"] == "FAILED"
    assert rec["error_kind"] == "unsupported_type"
    assert needle in rec["error_message"]


@pytest.mark.asyncio
async def test_generic_rejection_lists_powerpoint(deal_id: str) -> None:
    code, rec = await _upload(deal_id, "photo.heic", b"\x00\x00\x00\x18ftypheic" * 8, "image/heic")
    assert code == 422
    assert "PowerPoint (.pptx)" in rec["error_message"]
    assert "PDF" in rec["error_message"]


# ───────────────────────────── R-067 routing ─────────────────────────────


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("CBRE_PnL_Benchmark_Downtown_Upscale.pdf", "PNL_BENCHMARK"),
        ("CBRE Benchmarker 2025.xlsx", "PNL_BENCHMARK"),
        ("CBRE Trends Hotel Industry PnL.pdf", "PNL_BENCHMARK"),
        ("HotStats_benchmark.pdf", "PNL_BENCHMARK"),
        ("CBRE_Horizons_Q2_2026.pdf", "CBRE_HORIZONS"),
        ("cbre_forecast.pdf", "CBRE_HORIZONS"),
    ],
)
def test_cbre_benchmark_filenames_route_to_pnl_benchmark(filename: str, expected: str) -> None:
    from app.api.documents import _guess_doc_type

    assert _guess_doc_type(filename) == expected


# ───────────────────────────── R-067 mapping ─────────────────────────────


def _flat(**fields: object) -> dict[str, object]:
    return {f"pnl_benchmark.{k.replace('__', '.')}": v for k, v in fields.items()}


def test_map_uses_department_basis_for_departmental_lines() -> None:
    from app.services.pnl_benchmark_map import map_benchmark_to_categories

    flat = _flat(
        peer__rooms_revenue__total_usd=10_000_000,
        peer__fb_revenue__total_usd=4_000_000,
        peer__total_revenue__total_usd=15_000_000,
        peer__rooms_dept_expense__total_usd=2_500_000,
        peer__rooms_dept_expense__par_usd=12_000,
        peer__rooms_dept_expense__por_usd=45.5,
        peer__fb_dept_expense__total_usd=2_800_000,
        peer__a_and_g__total_usd=1_200_000,
        # Subject column is the property itself — never a benchmark.
        subject__utilities__total_usd=600_000,
    )
    cats = {c["key"]: c for c in map_benchmark_to_categories(flat)}
    assert set(cats) == {"rooms", "food_beverage", "administrative_general"}

    rooms = cats["rooms"]
    assert rooms["ratio"] == pytest.approx(0.25)  # 2.5M / 10M rooms revenue
    assert rooms["ratio_basis"] == "department_revenue"
    assert rooms["ratio_source"] == "computed_from_totals"
    assert rooms["par_usd"] == 12_000
    assert rooms["por_usd"] == 45.5

    assert cats["food_beverage"]["ratio"] == pytest.approx(0.70)  # 2.8M / 4M F&B
    ag = cats["administrative_general"]
    assert ag["ratio"] == pytest.approx(0.08)  # 1.2M / 15M total
    assert ag["ratio_basis"] == "total_revenue"


def test_map_falls_back_par_then_reported_then_legacy() -> None:
    from app.services.pnl_benchmark_map import map_benchmark_to_categories

    flat = _flat(
        # PAR basis — no totals on the report.
        peer__total_revenue__par_usd=50_000,
        peer__sales_marketing__par_usd=4_000,
        # Printed ratio only (percent units).
        peer__utilities__ratio_pct=3.4,
        # POR only — category kept, ratio stays unknown (no guessing).
        peer__insurance__por_usd=6.25,
        # Legacy summary aliases (decimals) for older extractions.
        property_taxes_pct=0.031,
        fb_dept_margin=0.30,
    )
    cats = {c["key"]: c for c in map_benchmark_to_categories(flat)}

    assert cats["sales_marketing"]["ratio"] == pytest.approx(0.08)
    assert cats["sales_marketing"]["ratio_source"] == "computed_from_par"
    assert cats["utilities"]["ratio"] == pytest.approx(0.034)
    assert cats["utilities"]["ratio_source"] == "reported"
    assert cats["insurance"]["ratio"] is None
    assert cats["insurance"]["ratio_source"] is None
    assert cats["insurance"]["por_usd"] == 6.25
    assert cats["property_taxes"]["ratio"] == pytest.approx(0.031)
    assert cats["property_taxes"]["ratio_source"] == "legacy_summary"
    # Margin alias flips to a cost ratio.
    assert cats["food_beverage"]["ratio"] == pytest.approx(0.70)
    # Lines the report never published are absent, not zero-filled.
    assert "rent" not in cats
    assert "information_telecom" not in cats


def test_map_returns_nothing_for_an_empty_report() -> None:
    from app.services.pnl_benchmark_map import map_benchmark_to_categories

    assert map_benchmark_to_categories({}) == []
    assert map_benchmark_to_categories({"pnl_benchmark.peer_set_size": 12}) == []


def test_market_data_block_carries_categories() -> None:
    from app.api.documents import _build_pnl_block

    flat = _flat(
        peer_set_size=9,
        peer_set_avg_keys=210,
        peer__total_revenue__total_usd=20_000_000,
        peer__property_taxes__total_usd=700_000,
    )
    block = _build_pnl_block(flat)
    assert block is not None
    assert block.peer_set_size == 9
    assert block.peer_set_avg_keys == 210
    assert [c.key for c in block.categories] == ["property_taxes"]
    assert block.categories[0].ratio == pytest.approx(0.035)
    dumped = block.model_dump()
    assert dumped["categories"][0]["ratio_basis"] == "total_revenue"
